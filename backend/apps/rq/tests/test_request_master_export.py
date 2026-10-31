"""
SYSPCC-030 · GET /api/v1/request-master/export/ (Excel).

Hojas y encabezados, valores, mismos filtros y alcance que el listado, fechas reales,
neutralización de fórmulas, tope de filas con aviso, ausencia de datos de trabajadores y de
precios, y errores.
"""
import datetime
import io
import logging
from decimal import Decimal
from unittest import mock

import pytest
from openpyxl import load_workbook

from apps.rq.services import request_master_export as export
from apps.rq.tests.catalog_fixtures import capture_logs
from apps.rq.tests.master_fixtures import (  # noqa: F401  (fixtures)
    central_warehouse, client_for, logistics_chief, logistics_supervisor, make_approval,
    make_delivery, make_item, make_request, world,
)

URL = '/api/v1/request-master/export/'


def download(user, **params):
    return client_for(user).get(URL, params)


def workbook(response):
    assert response.status_code == 200, response.content
    return load_workbook(io.BytesIO(response.content))


def table(ws):
    """(encabezado, filas) de una hoja sin fila de advertencia."""
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    return values[0], values[1:]


def cell_text(wb):
    return '\n'.join(
        str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r if v is not None
    )


class TestWorkbookShape:
    def test_response_headers_and_name(self, world, logistics_coordinator):
        response = download(logistics_coordinator)
        assert response['Content-Type'] == export.XLSX_CONTENT_TYPE
        assert response['Content-Disposition'].startswith('attachment; filename="Master_RQ_TODOS_')
        assert response['Content-Disposition'].endswith('.xlsx"')
        assert response['X-Export-Truncated'] == 'false'
        assert response['Cache-Control'] == 'no-store'
        assert response.content[:2] == b'PK'

    def test_three_sheets_in_order_with_the_original_headers(self, world, logistics_coordinator):
        wb = workbook(download(logistics_coordinator))
        assert wb.sheetnames == ['MASTER', 'RESUMEN POR RQ', 'ENTREGAS']
        assert table(wb['MASTER'])[0] == export.MASTER_HEADERS
        assert table(wb['RESUMEN POR RQ'])[0] == export.SUMMARY_HEADERS
        assert table(wb['ENTREGAS'])[0] == export.DELIVERIES_HEADERS
        for header in ('ID RQ', 'DESCRIPCIÓN', 'ÁREA', 'UNIDAD', 'CANT. SOLICITADA', 'CANT. ATENDIDA',
                       'STOCK ALMACÉN CENTRAL', 'CANT. FALTANTE', 'FUENTE', 'PRIORIDAD',
                       'FECHA DE APROBACIÓN', 'FECHA REQUERIDA', 'FECHA ATENDIDA', 'ESTADO DEL ÍTEM',
                       'N° DE GUÍA', 'COMENTARIO', 'PRESUPUESTO / ADICIONAL', 'RESPONSABLE DE ATENCIÓN',
                       'APROBACIÓN DE GERENCIA', '% ATENCIÓN DEL RQ'):
            assert header in export.MASTER_HEADERS

    def test_master_rows_match_the_listing(self, world, logistics_coordinator):
        wb = workbook(download(logistics_coordinator))
        header, data = table(wb['MASTER'])
        assert len(data) == 6
        col = {name: i for i, name in enumerate(header)}
        first = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0101' and r[col['LÍNEA']] == 1)
        assert first[col['DESCRIPCIÓN']] == 'Casco de seguridad'
        assert first[col['ÁREA']].startswith('Operaciones')
        assert (first[col['CANT. SOLICITADA']], first[col['CANT. ATENDIDA']], first[col['CANT. FALTANTE']]) == (
            10, 10, 0)
        assert first[col['STOCK ALMACÉN CENTRAL']] == 10
        assert first[col['FUENTE']] == 'EPP' and first[col['PRIORIDAD']] == 'Alta'
        assert first[col['ESTADO DEL ÍTEM']] == 'Atendido'
        assert first[col['N° DE GUÍA']] == 'GR-A1 / GR-A2'
        assert first[col['COMENTARIO']] == 'Entrega urgente'
        assert first[col['PRESUPUESTO / ADICIONAL']] == 'P'
        assert first[col['RESPONSABLE DE ATENCIÓN']] == 'Clerk_Two Test'
        assert first[col['APROBACIÓN DE GERENCIA']] == 'Aprobado (Adicional / fuera de plan)'
        assert first[col['% ATENCIÓN DEL RQ']] == 75
        assert first[col['SEGUIMIENTO']] == 'Vinculado' and first[col['ETAPA DEL RQ']] == 'Entregado'

    def test_dates_are_real_dates_in_lima(self, world, logistics_coordinator):
        header, data = table(workbook(download(logistics_coordinator))['MASTER'])
        col = {name: i for i, name in enumerate(header)}
        first = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0101' and r[col['LÍNEA']] == 1)
        for name, expected in (('FECHA DE APROBACIÓN', datetime.date(2026, 10, 1)),
                               ('FECHA REQUERIDA', datetime.date(2026, 11, 1)),
                               ('FECHA ATENDIDA', datetime.date(2026, 10, 8))):
            value = first[col[name]]
            assert isinstance(value, datetime.datetime) and value.date() == expected, name

    def test_unlinked_missing_and_empty_percent_read_as_dashes(self, world, logistics_coordinator):
        header, data = table(workbook(download(logistics_coordinator))['MASTER'])
        col = {name: i for i, name in enumerate(header)}
        unlinked = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0101' and r[col['LÍNEA']] == 4)
        assert unlinked[col['CANT. FALTANTE']] == '—' and unlinked[col['ESTADO DEL ÍTEM']] == 'Sin vincular'
        assert 'Ítem sin vincular' in unlinked[col['ALERTAS']]
        voided = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0103')
        assert voided[col['% ATENCIÓN DEL RQ']] == '—'

    def test_summary_sheet(self, world, logistics_coordinator):
        header, data = table(workbook(download(logistics_coordinator))['RESUMEN POR RQ'])
        assert len(data) == 3
        col = {name: i for i, name in enumerate(header)}
        a = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0101')
        assert (a[col['% ATENCIÓN']], a[col['ÍTEMS']], a[col['ATENDIDOS']], a[col['PARCIALES']],
                a[col['SIN VINCULAR']]) == (75, 4, 2, 1, 1)
        assert a[col['ETAPA']] == 'Entregado' and a[col['APROBACIÓN DE GERENCIA']].startswith('Aprobado')

    def test_deliveries_sheet_has_one_row_per_delivery_with_the_system_user(self, world, logistics_coordinator):
        header, data = table(workbook(download(logistics_coordinator))['ENTREGAS'])
        assert len(data) == 4                                  # A1×2, A2×1, C1×1
        col = {name: i for i, name in enumerate(header)}
        a1 = [r for r in data if r[col['ID RQ']] == 'RQ-2026-0101' and r[col['LÍNEA']] == 1]
        assert [(r[col['TIPO']], r[col['N° DE GUÍA']], r[col['CANTIDAD']], r[col['RESPONSABLE']]) for r in a1] == [
            ('ENTREGA', 'GR-A1', 6, 'Clerk_One Test'), ('ENTREGA', 'GR-A2', 4, 'Clerk_Two Test')]
        assert a1[0][col['FECHA']].date() == datetime.date(2026, 10, 5)

    def test_a_reversal_is_a_negative_row_of_type_reversa(self, world, logistics_coordinator):
        from apps.warehouse.models import WarehouseDispatch
        original = WarehouseDispatch.objects.get(dispatch_guide_number='GR-A2')
        make_delivery(world.a, world.a1, 4, guide='REV-1', on=datetime.date(2026, 10, 9),
                      by=world.clerk1, reverses=original)
        header, data = table(workbook(download(logistics_coordinator))['ENTREGAS'])
        col = {name: i for i, name in enumerate(header)}
        rev = [r for r in data if r[col['TIPO']] == 'REVERSA']
        assert len(rev) == 1 and rev[0][col['CANTIDAD']] == -4

    def test_override_note_goes_into_the_comment(self, world, logistics_coordinator):
        world.b1.attention_override, world.b1.attention_note = 'STAND_BY', 'Esperando al cliente'
        world.b1.save()
        header, data = table(workbook(download(logistics_coordinator))['MASTER'])
        col = {name: i for i, name in enumerate(header)}
        row = next(r for r in data if r[col['ID RQ']] == 'RQ-2026-0102')
        assert row[col['COMENTARIO']] == 'Stand by: Esperando al cliente'
        assert row[col['ESTADO DEL ÍTEM']] == 'Stand by'


class TestScopeAndFilters:
    def ids(self, response, sheet='MASTER'):
        _header, data = table(workbook(response)[sheet])
        return sorted({r[0] for r in data})

    def test_same_visibility_as_the_listing(self, world, requester, other_requester, project_resident,
                                            general_manager):
        assert self.ids(download(requester)) == ['RQ-2026-0101']
        assert self.ids(download(other_requester)) == ['RQ-2026-0102']
        assert self.ids(download(project_resident)) == ['RQ-2026-0101']
        assert len(self.ids(download(general_manager))) == 3

    def test_a_scoped_user_gets_no_foreign_row_in_any_sheet(self, world, requester):
        text = cell_text(workbook(download(requester)))
        assert 'RQ-2026-0102' not in text and 'RQ-2026-0103' not in text
        assert 'GR-C1' not in text and 'Papel bond' not in text

    def test_same_filters_as_the_listing(self, world, logistics_coordinator, project):
        assert self.ids(download(logistics_coordinator, project=project.pk)) == ['RQ-2026-0101']
        assert self.ids(download(logistics_coordinator, item_status='ANULADO')) == ['RQ-2026-0103']
        assert self.ids(download(logistics_coordinator, search='GR-C1')) == ['RQ-2026-0103']
        assert self.ids(download(logistics_coordinator, with_alerts='true')) == ['RQ-2026-0101', 'RQ-2026-0103']

    def test_resumen_keeps_the_whole_rq_percent_under_item_filters(self, world, logistics_coordinator):
        response = download(logistics_coordinator, item_status='PARCIAL')
        _header, master_rows = table(workbook(response)['MASTER'])
        _header, summary = table(workbook(response)['RESUMEN POR RQ'])
        assert len(master_rows) == 1 and len(summary) == 1
        assert summary[0][10] == 75 and summary[0][11] == 4

    def test_filename_uses_the_project_code_only_when_visible(self, world, logistics_coordinator, requester,
                                                              project, other_project):
        response = download(logistics_coordinator, project=project.pk)
        assert f'Master_RQ_{project.code}_' in response['Content-Disposition']
        # Un solicitante sin RQs en ese proyecto no averigua su código.
        hidden = download(requester, project=other_project.pk)
        assert other_project.code not in hidden['Content-Disposition']
        assert f'Master_RQ_PROYECTO{other_project.pk}_' in hidden['Content-Disposition']

    def test_empty_result_is_still_a_valid_workbook(self, world, hr_manager):
        wb = workbook(download(hr_manager))
        assert wb.sheetnames == ['MASTER', 'RESUMEN POR RQ', 'ENTREGAS']
        assert len(table(wb['MASTER'])[1]) == 0

    def test_requires_authentication_and_valid_filters(self, world, api_client, logistics_coordinator):
        assert api_client.get(URL).status_code == 401
        assert download(logistics_coordinator, flow='NOPE').status_code == 400
        assert download(logistics_coordinator, search='\x00').status_code == 400


class TestHostileContent:
    HOSTILE = ['=cmd|"/c calc"!A1', '+SUM(1+1)', '-2+3', '@SUM(A1)', '\t=1+1', '\r=1+1']

    def test_formula_like_text_is_neutralized_in_every_sheet(self, world, logistics_coordinator):
        for n, payload in enumerate(self.HOSTILE, start=10):
            world.a.rq_number = world.a.rq_number            # sin cambio: el RQ es auto-generado
            item = make_item(world.b, n, description=payload, quantity=1, comentarios=payload,
                             unit=payload[:30])
            make_delivery(world.b, item, 1, guide=payload[:50], by=world.clerk1,
                          on=datetime.date(2026, 10, 14))
        world.b.requested_by.first_name = '=HYPERLINK("http://evil")'
        world.b.requested_by.save()
        wb = workbook(download(logistics_coordinator))
        checked = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        checked += 1
                        assert cell.data_type != 'f', (ws.title, cell.coordinate, cell.value)
                        assert not cell.value.startswith(('=', '+', '-', '@', '\t', '\r')) or cell.value == '—' \
                            or cell.value[0] == "'", (ws.title, cell.coordinate, cell.value)
        assert checked > 50

    def test_override_note_cannot_inject_a_formula(self, world, logistics_coordinator):
        world.b1.attention_override = 'STAND_BY'
        world.b1.attention_note = '=1+1'
        world.b1.save()
        wb = workbook(download(logistics_coordinator))
        assert all(c.data_type != 'f' for ws in wb.worksheets for r in ws.iter_rows() for c in r)

    def test_control_characters_do_not_break_the_workbook(self, world, logistics_coordinator):
        world.a4.description = 'Cinta\x07 señal\x1fizadora'
        world.a4.save()
        wb = workbook(download(logistics_coordinator))
        assert 'Cinta señalizadora' in cell_text(wb)

    def test_very_long_text_is_cut_to_the_cell_limit(self, world, logistics_coordinator):
        world.a4.comentarios = 'x' * 40000
        world.a4.save()
        wb = workbook(download(logistics_coordinator))
        assert max(len(str(c.value)) for ws in wb.worksheets for r in ws.iter_rows() for c in r
                   if c.value is not None) <= 32001


class TestNoSensitiveData:
    def test_no_worker_names_dni_vouchers_or_prices(self, world, logistics_coordinator):
        from apps.core.models import Personal
        from apps.warehouse.models import MovementGroup

        worker = Personal.objects.create(dni='45678912', apellidos_nombres='Quispe Mamani, Hipólito',
                                         estado='ACTIVO', puesto='Operario de prueba')
        for group in MovementGroup.objects.exclude(warehouse_dispatch=None):
            group.destination_detail = 'Hipólito Quispe Mamani DNI 45678912'
            group.voucher_number = 'V-0099'
            group.destination_personal = worker
            group.save()
        for item in (world.a1, world.b1):
            item.unit_price = Decimal('12.5')
            item.save()
        text = cell_text(workbook(download(logistics_coordinator)))
        for secret in ('Hipólito', 'Quispe', '45678912', 'V-0099'):
            assert secret not in text, secret
        # Los precios (125 = 10 × 12.5, 62.5 = 5 × 12.5) no aparecen en ninguna celda numérica.
        numbers = {c.value for ws in workbook(download(logistics_coordinator)).worksheets
                   for r in ws.iter_rows() for c in r if isinstance(c.value, (int, float))}
        assert not numbers & {12.5, 125, 62.5}
        for forbidden in ('PRECIO', 'COSTO', 'DNI', 'VALE'):
            assert not any(forbidden in h for h in export.MASTER_HEADERS + export.DELIVERIES_HEADERS
                           + export.SUMMARY_HEADERS)


class TestRowCap:
    def test_truncated_export_warns_in_the_header_and_in_the_sheet(self, world, logistics_coordinator):
        with mock.patch.object(export, 'MAX_MASTER_ROWS', 3):
            response = download(logistics_coordinator)
        assert response['X-Export-Truncated'] == 'true'
        wb = workbook(response)
        rows = list(wb['MASTER'].iter_rows(values_only=True))
        assert rows[0][0] == export.TRUNCATED_WARNING
        assert list(rows[1]) == export.MASTER_HEADERS
        assert len(rows) == 2 + 3
        # La hoja de resumen sigue el mismo conjunto de filtros (no se trunca).
        assert len(table(wb['RESUMEN POR RQ'])[1]) == 3

    def test_deliveries_are_capped_with_their_own_warning(self, world, logistics_coordinator):
        with mock.patch.object(export, 'MAX_DELIVERY_ROWS', 2):
            response = download(logistics_coordinator)
        assert response['X-Export-Truncated'] == 'true'
        rows = list(workbook(response)['ENTREGAS'].iter_rows(values_only=True))
        assert rows[0][0] == export.TRUNCATED_WARNING and len(rows) == 2 + 2

    def test_exact_cap_is_not_truncated(self, world, logistics_coordinator):
        with mock.patch.object(export, 'MAX_MASTER_ROWS', 6):
            response = download(logistics_coordinator)
        assert response['X-Export-Truncated'] == 'false'

    def test_chunked_writing_keeps_every_row_and_the_order(self, world, logistics_coordinator):
        with mock.patch('apps.rq.services.request_master.CHUNK_SIZE', 2):
            wb = workbook(download(logistics_coordinator, ordering='rq_number'))
        _header, data = table(wb['MASTER'])
        assert [(r[0], r[22]) for r in data] == [
            ('RQ-2026-0101', 1), ('RQ-2026-0101', 2), ('RQ-2026-0101', 3), ('RQ-2026-0101', 4),
            ('RQ-2026-0102', 1), ('RQ-2026-0103', 1)]


class TestFailures:
    def test_an_unexpected_error_is_a_generic_500_and_is_logged(self, world, logistics_coordinator, caplog):
        with capture_logs(caplog, 'apps.rq.views.master', level=logging.ERROR), \
                mock.patch('apps.rq.views.master.build_master_workbook', side_effect=RuntimeError('boom')):
            response = download(logistics_coordinator)
        assert response.status_code == 500
        body = response.json()
        assert body['code'] == 'INTERNAL_ERROR' and 'boom' not in response.content.decode()
        assert any(r.message == 'rq.master.export.failed' for r in caplog.records)

    def test_the_export_has_its_own_throttle_declared_in_settings(self):
        from django.conf import settings
        from apps.rq.throttling import RequestMasterExportThrottle
        assert RequestMasterExportThrottle.scope == 'rq_master_export'
        assert settings.REST_FRAMEWORK['DEFAULT_THROTTLE_RATES']['rq_master_export'] == '10/min'
        assert RequestMasterExportThrottle.THROTTLE_RATES['rq_master_export'] == '10/min'

    def test_the_third_download_in_a_minute_answers_429(self, world, logistics_coordinator, monkeypatch):
        from apps.rq.throttling import RequestMasterExportThrottle
        monkeypatch.setitem(RequestMasterExportThrottle.THROTTLE_RATES, 'rq_master_export', '2/min')
        statuses = [download(logistics_coordinator).status_code for _ in range(3)]
        assert statuses == [200, 200, 429]

    def test_the_export_limit_does_not_throttle_the_listing(self, world, logistics_coordinator, monkeypatch):
        from apps.rq.throttling import RequestMasterExportThrottle
        monkeypatch.setitem(RequestMasterExportThrottle.THROTTLE_RATES, 'rq_master_export', '1/min')
        assert download(logistics_coordinator).status_code == 200
        assert download(logistics_coordinator).status_code == 429
        assert client_for(logistics_coordinator).get('/api/v1/request-master/').status_code == 200
