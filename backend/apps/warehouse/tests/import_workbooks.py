"""
Synthetic workbooks for the importer tests (SYSPCC-025).

Everything here is INVENTED. Never copy names, destinations or values from a real
workbook into tests, fixtures or logs. Worker names below are made-up placeholders.
"""
import io
import re
import zipfile
from datetime import date

from openpyxl import Workbook

from apps.warehouse.services.imports import parser
from apps.warehouse.services.imports.reporting import TEMPLATE_HEADERS

# Invented workers (the importer matches them against core.Personal by name).
WORKER_A = 'Quispe Mamani, Rosa Test'
WORKER_A_REORDERED = 'ROSA TEST QUISPE MAMANI'
WORKER_B = 'Ccori Huaman, Pedro Test'
WORKER_C = 'Apaza Condori, Luis Test'

D1 = date(2026, 10, 5)
D2 = date(2026, 10, 6)
D3 = date(2026, 10, 7)


# -- row builders (aligned with the template headers) ----------------------------

def catalog_row(code, description, unit='EA', order_number=None, oc=None, packing=None, shipment=None):
    return [order_number, code, description, unit, oc, packing, shipment]


def entry_row(code, qty, when=D1, tag=None, location='PISO 2', po=None, guide=None, transfer=None, obs=None):
    return [None, when, tag, code, None, qty, None, location, po, guide, transfer, obs]


def exit_row(code, qty, destination, voucher=None, when=D2, location='PISO 2', transfer=None, obs=None):
    return [None, voucher, when, code, None, qty, location, transfer, destination, obs]


def summary_row(code, description, unit='EA', initial=None, stock=None, location='PISO 2',
                observation=None, project='ON PLANT', entries=None, exits=None):
    return [None, code, description, unit, initial, entries, exits, stock, location, observation, project]


# -- workbook builder ---------------------------------------------------------------

def build_workbook(*, catalog=None, entries=None, exits=None, summary=None, sheet_names=None,
                   headers=None, header_offset=0, ghost_rows=False) -> bytes:
    """
    xlsx bytes. Each argument is a list of row lists (see builders) or None to omit the sheet.
    ``sheet_names`` overrides the sheet titles, ``headers`` the header rows (per canonical
    sheet name), ``header_offset`` puts N blank/warning rows before the header and
    ``ghost_rows`` inflates the declared dimension to ~1 M rows like the real workbook.
    """
    sheet_names = sheet_names or {}
    headers = headers or {}
    wb = Workbook()
    wb.remove(wb.active)
    data = {
        parser.SHEET_CATALOG: catalog,
        parser.SHEET_ENTRIES: entries,
        parser.SHEET_EXITS: exits,
        parser.SHEET_SUMMARY: summary,
    }
    for name in parser.SHEET_ORDER:
        rows = data[name]
        if rows is None:
            continue
        ws = wb.create_sheet(sheet_names.get(name, name))
        for i in range(header_offset):
            ws.append([f'ADVERTENCIA: fila {i + 1}'])
        ws.append(headers.get(name, TEMPLATE_HEADERS[name]))
        for n, row in enumerate(rows, start=1):
            if name != parser.SHEET_CATALOG and row and row[0] is None:
                row = [n] + list(row[1:])
            ws.append(list(row))
        if ghost_rows:
            ws.cell(row=1_000_000, column=1).number_format = '0'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def standard_workbook(**overrides) -> bytes:
    """
    A small, fully consistent workbook: 3 catalog items, 2 entries, 3 exits (one vale with
    two lines), a summary that reconciles with the movements and one project label.
    """
    parts = dict(
        catalog=[
            catalog_row('EPP000001', 'GUANTE DE CUERO T: L', oc=1.02, packing='24003-13', order_number='ORD-1'),
            catalog_row('EPP000002', 'CASCO BLANCO', oc=2.03),
            catalog_row('EPP000003', 'LENTE CLARO'),
        ],
        entries=[
            entry_row('EPP000001', 10, D1, po='PO-77', guide='GR-1'),
            entry_row('EPP000002', 5, D1, po='PO-77', guide='GR-1'),
        ],
        exits=[
            exit_row('EPP000001', 2, WORKER_A, voucher=120, when=D2),
            exit_row('EPP000002', 1, WORKER_A, voucher='120', when=D2),
            exit_row('EPP000001', 3, WORKER_B, voucher=121, when=D3),
        ],
        summary=[
            summary_row('EPP000001', 'GUANTE DE CUERO T: L', entries=10, exits=5, stock=5),
            summary_row('EPP000002', 'CASCO BLANCO', entries=5, exits=1, stock=4),
            summary_row('EPP000003', 'LENTE CLARO', entries=0, exits=0, stock=0),
        ],
    )
    parts.update(overrides)
    return build_workbook(**parts)


# -- hostile / special files -------------------------------------------------------------

def rewrite_zip(data: bytes, *, replace=None, add=None) -> bytes:
    """Copy an xlsx replacing/adding parts: ``replace``/``add`` map part name -> bytes."""
    replace, add = replace or {}, add or {}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, replace.get(info.filename, src.read(info.filename)))
        for name, content in add.items():
            dst.writestr(name, content)
    return out.getvalue()


def inject_cached_values(data: bytes, sheet_part: str, values: dict) -> bytes:
    """Give formula cells a cached value, like a workbook saved by Excel: ``{'F2': 5}``."""
    xml = zipfile.ZipFile(io.BytesIO(data)).read(sheet_part).decode()
    for cell, value in values.items():
        pattern = rf'(<c r="{cell}"[^>]*><f>[^<]*</f>)<v></v>'
        xml, count = re.subn(pattern, rf'\g<1><v>{value}</v>', xml)
        assert count == 1, f'formula cell {cell} not found in {sheet_part}'
    return rewrite_zip(data, replace={sheet_part: xml.encode()})
