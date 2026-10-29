"""
Excel exports in the layout of the real workbook "INVENTARIO PCC - ALM" (SYSPCC-025).

Sheets and headers (written as values, never formulas):

* ``RESUMEN DE MOVIMIENTO``: N° | ITEM | DESCRIPCION | UOM | INV. INICIAL | INGRESOS |
  EGRESOS | STOCK ACTUAL | LOCACIÓN | OBSERVACIÓN | PROYECTO
* ``INGRESOS``: N° | FECHA | TAG | ITEM | DESCRIPCIÓN | CANTIDAD INGRESADA | UOM |
  LOCACIÓN | PO | GUÍA DE REMISIÓN | TRANSFERENCIA INTERNA | OBSERVACIÓN
* ``EGRESOS``: N° | Nª VALE | FECHA | ITEM | DESCRIPCIÓN | CANTIDAD A EGRESAR | LOCACIÓN |
  TRANSFERENCIA INTERNA | USUARIO/ DESTINO | OBSERVACIÓN

The exported workbook can be re-imported by the importer. Every free-text value
goes through ``sanitize_excel_value`` (formula-injection defence). When a sheet is
cut at ``MAX_EXPORT_ROWS`` a warning row is written as the first row, following the
precedent of the other warehouse exports (header detection must therefore scan the
first rows instead of assuming row 1).
"""
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core.utils.excel import sanitize_excel_value, truncate_for_export
from apps.warehouse.services import kardex

SUMMARY_SHEET = 'RESUMEN DE MOVIMIENTO'
ENTRIES_SHEET = 'INGRESOS'
EXITS_SHEET = 'EGRESOS'

SUMMARY_HEADERS = [
    'N°', 'ITEM', 'DESCRIPCION', 'UOM', 'INV. INICIAL', 'INGRESOS', 'EGRESOS',
    'STOCK ACTUAL', 'LOCACIÓN', 'OBSERVACIÓN', 'PROYECTO',
]
ENTRIES_HEADERS = [
    'N°', 'FECHA', 'TAG', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD INGRESADA', 'UOM',
    'LOCACIÓN', 'PO', 'GUÍA DE REMISIÓN', 'TRANSFERENCIA INTERNA', 'OBSERVACIÓN',
]
EXITS_HEADERS = [
    'N°', 'Nª VALE', 'FECHA', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD A EGRESAR', 'LOCACIÓN',
    'TRANSFERENCIA INTERNA', 'USUARIO/ DESTINO', 'OBSERVACIÓN',
]

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
DATE_FORMAT = 'dd/mm/yyyy'

# One wording for every truncated warehouse export (first row of the sheet, before the header).
TRUNCATED_EXPORT_WARNING = (
    'ADVERTENCIA: la exportación se truncó. Aplique filtros adicionales para ver el resto.'
)


def _clean(values):
    return [sanitize_excel_value(v) for v in values]


def _start_sheet(ws, headers, truncated: bool):
    if truncated:
        ws.append([TRUNCATED_EXPORT_WARNING])
    ws.append(headers)
    header_row = 2 if truncated else 1
    bold = Font(bold=True)
    for cell in ws[header_row]:
        cell.font = bold
    ws.freeze_panes = f'A{header_row + 1}'  # a string: ws.cell() would create a blank row


def _movement_note(mv) -> str:
    if mv.notes:
        return mv.notes
    return mv.group.notes if mv.group_id and mv.group else ''


def build_summary_sheet(ws, items, *, truncated: bool = False):
    ws.title = SUMMARY_SHEET
    _start_sheet(ws, SUMMARY_HEADERS, truncated)
    for n, item in enumerate(items, start=1):
        row = kardex.summary_row(item)
        ws.append(_clean([
            n,
            row['product_code'],
            row['description'],
            row['unit'],
            row['initial_stock'],
            row['entries'],
            row['exits'],
            row['closing_stock'],
            row['location'],
            row['observation'],
            row['project_name'],
        ]))


def build_entries_sheet(ws, movements, *, truncated: bool = False):
    ws.title = ENTRIES_SHEET
    _start_sheet(ws, ENTRIES_HEADERS, truncated)
    for n, mv in enumerate(movements, start=1):
        ws.append(_clean([
            n,
            mv.movement_date,
            mv.tag,
            mv.inventory.product_code,
            mv.inventory.description,
            float(mv.quantity),
            mv.inventory.unit,
            mv.inventory.location_display,
            mv.po_number,
            mv.guide_number,
            mv.internal_transfer,
            _movement_note(mv),
        ]))
        ws.cell(row=ws.max_row, column=2).number_format = DATE_FORMAT


def build_exits_sheet(ws, movements, *, truncated: bool = False):
    ws.title = EXITS_SHEET
    _start_sheet(ws, EXITS_HEADERS, truncated)
    for n, mv in enumerate(movements, start=1):
        ws.append(_clean([
            n,
            mv.voucher_number,
            mv.movement_date,
            mv.inventory.product_code,
            mv.inventory.description,
            float(mv.quantity),
            mv.inventory.location_display,
            mv.internal_transfer,
            mv.destination_detail,
            _movement_note(mv),
        ]))
        ws.cell(row=ws.max_row, column=3).number_format = DATE_FORMAT


def split_movements(qs, *, initial_loads_in_entries: bool = False):
    """
    (entries, exits) querysets from a movement queryset.

    One rule for the initial load: it is written to INGRESOS unless the INV. INICIAL column
    of the same workbook already carries it (``initial_loads_in_entries=False``). The caller
    passes ``kardex.initial_loads_count_as_entries(date_from)`` when the workbook has a
    summary sheet, and ``True`` when it has none (movements-only export), so the movement
    sheets plus INV. INICIAL always add up to the stock and the file can be re-imported
    without losing units.
    """
    base = qs.select_related('inventory', 'group')
    entries = base.filter(kardex.IN_Q)
    if not initial_loads_in_entries:
        entries = entries.exclude(kardex.INITIAL_Q)
    exits = base.filter(kardex.OUT_Q)
    order = ('movement_date', 'pk')
    return entries.order_by(*order), exits.order_by(*order)


def build_summary_workbook(items_qs, *, movements_qs=None, date_from=None) -> Workbook:
    """
    Workbook with ``RESUMEN DE MOVIMIENTO`` and, when ``movements_qs`` is given,
    the ``INGRESOS`` and ``EGRESOS`` sheets for those movements.

    ``date_from`` is the window the summary was computed with: it decides whether the
    initial loads are INGRESOS of the period or part of INV. INICIAL (see
    ``kardex.initial_loads_count_as_entries``).
    """
    wb = Workbook()
    items, truncated = truncate_for_export(items_qs)
    build_summary_sheet(wb.active, items, truncated=truncated)
    if movements_qs is not None:
        entries_qs, exits_qs = split_movements(
            movements_qs, initial_loads_in_entries=kardex.initial_loads_count_as_entries(date_from),
        )
        entries, e_trunc = truncate_for_export(entries_qs)
        exits, x_trunc = truncate_for_export(exits_qs)
        build_entries_sheet(wb.create_sheet(), entries, truncated=e_trunc)
        build_exits_sheet(wb.create_sheet(), exits, truncated=x_trunc)
    return wb


def build_movements_workbook(movements_qs) -> Workbook:
    """
    Workbook with only ``INGRESOS`` and ``EGRESOS`` (movements export, layout excel_pcc).
    There is no INV. INICIAL column here, so the initial loads are written as entries.
    """
    wb = Workbook()
    entries_qs, exits_qs = split_movements(movements_qs, initial_loads_in_entries=True)
    entries, e_trunc = truncate_for_export(entries_qs)
    exits, x_trunc = truncate_for_export(exits_qs)
    build_entries_sheet(wb.active, entries, truncated=e_trunc)
    build_exits_sheet(wb.create_sheet(), exits, truncated=x_trunc)
    return wb
