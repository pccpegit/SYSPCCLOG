"""Excel files produced by the importer: the blank template and the problems report (SYSPCC-025)."""
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core.utils.excel import sanitize_excel_value
from apps.warehouse.services.imports import parser
from apps.warehouse.services.imports.problems import SEVERITY_LABELS, code_label

TEMPLATE_HEADERS = {
    parser.SHEET_CATALOG: [
        'Order Number', 'Item', 'Item Description', 'Unit of Measure', 'OC Nª',
        'Supplier Packing List Number', 'Shipment Quantity',
    ],
    parser.SHEET_ENTRIES: [
        'N°', 'FECHA', 'TAG', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD INGRESADA', 'UOM', 'LOCACIÓN', 'PO',
        'GUÍA DE REMISIÓN', 'TRANSFERENCIA INTERNA', 'OBSERVACIÓN',
    ],
    parser.SHEET_EXITS: [
        'N°', 'Nª VALE', 'FECHA', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD A EGRESAR', 'LOCACIÓN',
        'TRANSFERENCIA INTERNA', 'USUARIO/ DESTINO', 'OBSERVACIÓN',
    ],
    parser.SHEET_SUMMARY: [
        'N°', 'ITEM', 'DESCRIPCION', 'UOM', 'INV. INICIAL', 'INGRESOS', 'EGRESOS', 'STOCK ACTUAL',
        'LOCACIÓN', 'OBSERVACIÓN', 'PROYECTO',
    ],
}

# Same words as the screen: Gravedad / Tipo de problema (a readable label, not the raw code).
PROBLEM_HEADERS = ['Gravedad', 'Hoja', 'Fila', 'Tipo de problema', 'Mensaje']


def _header(ws, headers):
    ws.append(headers)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    ws.freeze_panes = 'A2'


def build_template_workbook() -> Workbook:
    """Blank workbook with the four sheets and their headers (empty rows)."""
    wb = Workbook()
    first = True
    for name in parser.SHEET_ORDER:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        _header(ws, TEMPLATE_HEADERS[name])
    return wb


def build_problems_workbook(problems: list) -> Workbook:
    """
    Problems report. It carries sheet, row, type of problem and the (name-free) message, so it can
    be shared without exposing the workers listed in the workbook.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'PROBLEMAS'
    _header(ws, PROBLEM_HEADERS)
    for item in problems:
        ws.append([
            sanitize_excel_value(value) for value in (
                SEVERITY_LABELS.get(item.get('severity'), item.get('severity', '')),
                item.get('sheet', ''),
                item.get('row') or '',
                code_label(item.get('code', '')),
                item.get('message', ''),
            )
        ])
    return wb
