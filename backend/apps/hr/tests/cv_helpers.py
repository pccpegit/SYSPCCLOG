"""Shared helpers for the CV analysis tests (SYSPCC-023): fake Anthropic client
and generators of tiny test files. No network anywhere."""

import io
import re
from types import SimpleNamespace

import anthropic
import httpx2 as httpx
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.hr.services.cv_schemas import PostingExtraction, PostingRequirement, RequirementResult  # noqa: E402

REQ = httpx.Request('POST', 'https://api.anthropic.com/v1/messages')


def api_error(cls, status, headers=None):
    resp = httpx.Response(status, request=REQ, headers=headers or {})
    return cls('boom with PII Juan Perez', response=resp, body=None)


def rate_limit(retry_after=None):
    return api_error(anthropic.RateLimitError, 429, {'retry-after': str(retry_after)} if retry_after else None)


def default_report(output_format, name='Ana Pérez Gómez', **overrides):
    """A valid report for the requested schema (with or without evaluation)."""
    data = {
        'candidate': {'full_name': name, 'email': 'ana@example.com', 'phone': '999111222', 'location': 'Lima'},
        'headline': 'Residente de obra',
        'summary': 'Ingeniera civil con experiencia en obras.',
        'experience': [
            {'company': 'ACME', 'role': 'Residente', 'start': '2018-01', 'end': '2021-12', 'is_current': False,
             'description': 'Obra vial'},
            {'company': 'BETA', 'role': 'Supervisora', 'start': '2020-01', 'end': None, 'is_current': True,
             'description': ''},
        ],
        'education': [{'institution': 'UNI', 'degree': 'Ingeniería Civil', 'status': 'completed', 'end_year': 2017}],
        'skills': ['AutoCAD', 'MS Project'],
        'certifications': [{'name': 'PMP', 'issuer': 'PMI', 'year': 2020}],
        'languages': [{'language': 'Inglés', 'level': 'B2'}],
        'readability': 'good', 'injection_suspected': False,
    }
    data.update(overrides)
    if 'evaluation' in output_format.model_fields and 'evaluation' not in data:
        data['evaluation'] = {'requirement_results': [], 'strengths': ['Experiencia'], 'gaps': []}
    return output_format.model_validate(data)


class FakeCVMessages:
    def __init__(self, owner, beta=False):
        self.owner, self.beta = owner, beta

    def parse(self, **kwargs):
        owner = self.owner
        owner.calls.append({'beta': self.beta, **kwargs})
        step = owner.script.pop(0) if owner.script else None
        if isinstance(step, Exception):
            raise step
        if step is not None:
            return step(kwargs) if callable(step) else step
        return owner.ok(kwargs)


class FakeCVClient:
    """`script`: list consumed per call; items are an Exception to raise, a
    response, or a callable(kwargs)->response. Empty script -> happy path."""

    def __init__(self, script=None, tokens=(1000, 500), statuses=None):
        self.calls, self.script, self.tokens = [], list(script or []), tokens
        self.messages = FakeCVMessages(self)
        self.beta = SimpleNamespace(messages=FakeCVMessages(self, beta=True))

    def ok(self, kwargs, **overrides):
        report = default_report(kwargs['output_format'], **overrides)
        if 'evaluation' in kwargs['output_format'].model_fields:
            text = kwargs['messages'][0]['content'][-1]['text']
            ids = [int(i) for i in re.findall(r'^(\d+)\. ', text, re.M)]
            report.evaluation.requirement_results = [
                RequirementResult(requirement_id=i, status='met', evidence='ok') for i in ids]
        return self.response(report)

    def response(self, parsed, stop_reason='end_turn', tokens=None):
        t = tokens or self.tokens
        return SimpleNamespace(stop_reason=stop_reason, parsed_output=parsed,
                               usage=SimpleNamespace(input_tokens=t[0], output_tokens=t[1]))


def response_with(client, **overrides):
    """Script step building a happy response with report overrides."""
    return lambda kwargs: client.response(default_report(kwargs['output_format'], **overrides))


def make_pdf(text='Curriculum Vitae Ana Perez', pages=1) -> bytes:
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages):
        c.drawString(72, 750, f'{text} p{i}')
        c.showPage()
    c.save()
    return buf.getvalue()


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter()
    for p in PdfReader(io.BytesIO(make_pdf())).pages:
        w.add_page(p)
    w.encrypt('secret')
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def make_docx(text='Ana Pérez Gómez — Ingeniera Civil con diez años de experiencia en obras viales y edificaciones.',
              table=True) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = 'Idioma', 'Inglés B2'
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_png(size=(40, 40), color=(200, 30, 30)) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new('RGB', size, color).save(out, format='PNG')
    return out.getvalue()


def upload(name: str, data: bytes, content_type='application/octet-stream'):
    return SimpleUploadedFile(name, data, content_type=content_type)


def posting_response(client, requirements=(('Ingeniero civil colegiado', True), ('5 años en obras viales', False)),
                     title='Residente de obra', omitted=False, tokens=(1830, 412)):
    """Script step: a happy `PostingExtraction` response from the fake client."""
    parsed = PostingExtraction(
        position_title=title, omitted_protected=omitted,
        requirements=[PostingRequirement(text=t, mandatory=m) for t, m in requirements])
    return lambda kwargs: client.response(parsed, tokens=tokens)


POSTING_TEXT = (
    'CONVOCATORIA: Residente de obra\n'
    'Somos una empresa constructora líder en el sector.\n\n'
    'Requisitos:\n'
    '- Ingeniero civil colegiado y habilitado (obligatorio)\n'
    '- Experiencia mínima de 5 años en obras viales\n'
    '- Manejo de AutoCAD y MS Project\n'
    '- Deseable: curso de seguridad en obra\n\n'
    'Funciones:\n'
    '- Supervisar el avance de la obra\n'
    '- Elaborar valorizaciones mensuales\n\n'
    'Ofrecemos:\n'
    '- Planilla completa y seguro de salud\n'
)


def make_blank_pdf(pages=1) -> bytes:
    """A PDF with no text layer (like a scanned one)."""
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for _ in range(pages):
        c.rect(50, 50, 100, 100)
        c.showPage()
    c.save()
    return buf.getvalue()
