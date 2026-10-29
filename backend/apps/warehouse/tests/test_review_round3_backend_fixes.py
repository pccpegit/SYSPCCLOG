"""
SYSPCC-025 · django-backend, fixes de la revisión final (code-review + localization-reviewer).

Todo es inventado: ningún nombre ni valor sale del libro real.

* Texto de usuario: Ingreso/Egreso (PDF, enum, export), "Jefe Logístico", "N° de vale (físico)",
  "vale" en lugar de "grupo", cantidades sin ceros finales, encabezados con tildes.
* PostgreSQL: un byte NUL en un parámetro de consulta es 400 (nunca 500).
* Movimientos ordenados por la fecha operativa (los sin fecha al final).
* Stock real vs. stock que suman los movimientos (``totals.current``, ``untracked_stock``).
"""
import io
from datetime import date
from decimal import Decimal
from unittest import mock

import openpyxl
import pytest
from django.db import IntegrityError
from pypdf import PdfReader

from apps.core.enums import MovementTypeChoices, RoleChoices
from apps.warehouse import permissions as wh_permissions
from apps.warehouse import views as wh_views
from apps.warehouse.models import InventoryStock, MovementGroup
from apps.warehouse.services import movements as svc
from apps.warehouse.services.imports import problems, reporting
from apps.warehouse.services.imports.validators import EntryRowSerializer, validate_row
from apps.warehouse.services.pdf_generator import generate_group_pdf, generate_movement_pdf

from .conftest import client_for, make_user

pytestmark = pytest.mark.django_db

BASE = '/api/v1/warehouse/'
INVENTORY = BASE + 'inventory/'
MOVEMENTS = BASE + 'movements/'
SUMMARY = INVENTORY + 'movement-summary/'


@pytest.fixture(autouse=True)
def _no_pdf_task():
    with mock.patch('apps.warehouse.views._enqueue_group_pdf_upload'):
        yield


def _entry(item, qty, user, when, **extra):
    kwargs = dict(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        source_type=extra.pop('source_type', 'PURCHASE'), supplier_name=extra.pop('supplier_name', ''),
        invoice_number=extra.pop('invoice_number', ''), notes='', registered_by=user,
    )
    if when is None:
        kwargs['undated'] = True
    else:
        kwargs['movement_date'] = when
    return svc.register_entry_batch(**kwargs, **extra)


def _exit(item, qty, user, when, **extra):
    kwargs = dict(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        destination_type='EMPLOYEE', destination_detail='Persona Inventada', requested_by_id=None,
        authorized_by_id=None, notes='', registered_by=user,
    )
    if when is None:
        kwargs['undated'] = True
    else:
        kwargs['movement_date'] = when
    return svc.register_exit_batch(**kwargs, **extra)


def _pdf_text(pdf: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf))
    return ' '.join(' '.join(page.extract_text() or '' for page in reader.pages).split())


# ---------------------------------------------------------------------------
# Texto de usuario
# ---------------------------------------------------------------------------

class TestRoleLabelsInMessages:
    """El rol LOGISTICS_CHIEF se llama «Jefe Logístico» (RoleChoices), nunca «Jefatura de Logística»."""

    @pytest.mark.parametrize('permission', [
        wh_permissions.CanImportInventory, wh_permissions.CanAssignLocations,
        wh_permissions.CanLinkPersonal, wh_permissions.CanLookupPersonal,
    ])
    def test_permission_messages_use_the_official_label(self, permission):
        assert str(RoleChoices.LOGISTICS_CHIEF.label) == 'Jefe Logístico'
        assert 'Jefe Logístico' in permission.message
        assert 'Jefatura' not in permission.message

    def test_storage_denied_message_uses_the_official_label(self):
        assert 'Jefe Logístico' in wh_views.STORAGE_DENIED and 'Jefatura' not in wh_views.STORAGE_DENIED

    def test_the_403_a_site_warehouse_user_gets_names_the_official_label(self, epp_item):
        user = make_user('site_wh', RoleChoices.SITE_WAREHOUSE)
        resp = client_for(user).patch(f'{INVENTORY}{epp_item.pk}/', {'location': 'F3-2'}, format='json')
        assert resp.status_code == 403
        assert 'Jefe Logístico' in resp.data['detail']


class TestIngresoEgresoVocabulary:
    def test_the_enum_labels_are_ingreso_and_egreso(self):
        assert str(MovementTypeChoices.ENTRY.label) == 'Ingreso'
        assert str(MovementTypeChoices.EXIT.label) == 'Egreso'

    def test_the_movements_export_writes_ingreso_egreso_in_the_type_column(self, epp_item, central_warehouse):
        _entry(epp_item, 2, central_warehouse, date(2026, 10, 5))
        _exit(epp_item, 1, central_warehouse, date(2026, 10, 6))
        resp = client_for(central_warehouse).get(MOVEMENTS + 'export/')
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        header = [c.value for c in sheet[1]]
        assert header[:3] == ['N° de movimiento', 'Tipo', 'Fecha']
        assert {sheet.cell(row=r, column=2).value for r in (2, 3)} == {'Ingreso', 'Egreso'}

    def test_the_duplicate_voucher_conflict_says_egreso(self, epp_item, central_warehouse):
        client = client_for(central_warehouse)
        body = {'warehouse': 'CENTRAL', 'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona',
                'voucher_number': 'V-77', 'items': [{'inventory': epp_item.pk, 'quantity': '1'}]}
        assert client.post(MOVEMENTS + 'batch-exit/', body, format='json').status_code == 201
        dup = client.post(MOVEMENTS + 'batch-exit/', body, format='json')
        assert dup.status_code == 409
        assert dup.data['detail'] == 'Ya existe un egreso con ese N° de vale en este almacén.'

    def test_the_epp_delivery_duplicate_voucher_says_egreso(self, epp_item, central_warehouse):
        client = client_for(central_warehouse)
        body = {'warehouse': 'CENTRAL', 'voucher_number': 'V-78', 'movement_date': '2026-10-16',
                'recipient': {'name': 'Persona Libre'}, 'items': [{'inventory': epp_item.pk, 'quantity': '1'}]}
        assert client.post(MOVEMENTS + 'epp-delivery/', body, format='json').status_code == 201
        dup = client.post(MOVEMENTS + 'epp-delivery/', body, format='json')
        assert dup.status_code == 409
        assert 'Ya existe un egreso' in dup.data['detail'] and 'salida' not in dup.data['detail']


class TestVoucherVocabularyInErrors:
    """El usuario ve «vale», no «grupo»."""

    def test_an_unknown_vale_is_vale_no_encontrado(self, central_warehouse):
        client = client_for(central_warehouse)
        for suffix in ('voucher', 'pdf'):
            resp = client.get(f'{MOVEMENTS}groups/999999/{suffix}/')
            assert resp.status_code == 404
            assert resp.data['detail'] == 'Vale no encontrado.'

    @pytest.mark.parametrize('path,body', [
        ('batch-entry/', {'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}),
        ('batch-exit/', {'warehouse': 'CENTRAL', 'destination_type': 'EMPLOYEE', 'destination_detail': 'P'}),
        ('epp-delivery/', {'warehouse': 'CENTRAL', 'voucher_number': 'V-1', 'movement_date': '2026-10-16',
                           'recipient': {'name': 'P'}}),
    ])
    def test_the_numbering_conflict_speaks_of_the_internal_vale(self, epp_item, central_warehouse, path, body):
        body = {**body, 'items': [{'inventory': epp_item.pk, 'quantity': '1'}]}
        target = {
            'batch-entry/': 'register_entry_batch', 'batch-exit/': 'register_exit_batch',
            'epp-delivery/': 'register_exit_batch',
        }[path]
        with mock.patch.object(svc, target, side_effect=IntegrityError('duplicate group_number')):
            resp = client_for(central_warehouse).post(MOVEMENTS + path, body, format='json')
        assert resp.status_code == 409
        assert resp.data['detail'] == 'No se pudo generar el N° de vale interno. Intente nuevamente.'
        assert 'grupo' not in str(resp.data).lower()


class TestStockMessages:
    def test_stock_errors_do_not_print_trailing_zeros(self, epp_item, central_warehouse):
        with pytest.raises(ValueError) as exc:
            _exit(epp_item, 11, central_warehouse, date(2026, 10, 5))
        assert 'Disponible: 10, solicitado: 11' in str(exc.value)
        assert '10.000' not in str(exc.value)

    def test_fractional_quantities_keep_their_decimals(self, make_item, central_warehouse):
        item = make_item('TMT000001', 'TELA', 'M2', stock='2.5')
        with pytest.raises(ValueError) as exc:
            _exit(item, '3.250', central_warehouse, date(2026, 10, 5))
        assert 'Disponible: 2.5, solicitado: 3.25' in str(exc.value)

    def test_the_negative_adjustment_message_is_spelled_with_its_accent(self, epp_item, central_warehouse):
        with pytest.raises(ValueError) as exc:
            svc.register_adjustment(
                inventory=epp_item, quantity=Decimal('-11'), warehouse='CENTRAL', project_id=None,
                notes='', registered_by=central_warehouse,
            )
        message = str(exc.value)
        assert 'resultaría en stock negativo' in message and 'resultaria' not in message
        assert 'Stock actual: 10, ajuste: -11' in message

    @pytest.mark.parametrize('value,expected', [
        ('5.000', '5'), ('2.500', '2.5'), ('0.001', '0.001'), ('100', '100'), (Decimal('1E+2'), '100'),
        (Decimal('0.000'), '0'), (3, '3'),
    ])
    def test_format_quantity(self, value, expected):
        assert svc.format_quantity(value) == expected


class TestPdfWording:
    @pytest.fixture
    def entry_group(self, epp_item, central_warehouse):
        return _entry(
            epp_item, 3, central_warehouse, date(2026, 10, 5), supplier_name='Proveedor Inventado',
            invoice_number='F001-55', po_number='OC-9', guide_number='T001-9', voucher_number='V-5',
        )

    @pytest.fixture
    def exit_group(self, epp_item, central_warehouse):
        return _exit(epp_item, 2, central_warehouse, date(2026, 10, 6), voucher_number='V-6')

    def test_the_entry_vale_is_titled_vale_de_ingreso(self, entry_group):
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=entry_group.pk)))
        assert 'VALE DE INGRESO' in text
        assert 'ENTRADA' not in text and 'SALIDA' not in text

    def test_the_exit_vale_is_titled_vale_de_egreso(self, exit_group):
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=exit_group.pk)))
        assert 'VALE DE EGRESO' in text
        assert 'ENTRADA' not in text and 'SALIDA' not in text

    def test_the_physical_vale_number_label_is_spelled_correctly(self, entry_group):
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=entry_group.pk)))
        assert 'N° de vale (físico)' in text and 'V-5' in text
        assert 'N. Vale' not in text and 'fisico' not in text

    def test_labels_carry_their_accents(self, exit_group):
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=exit_group.pk)))
        for expected in ('Almacén Central', 'Código', 'Descripción', 'Entregó', 'Recibió', 'Autorizó',
                         'automáticamente', 'Tipo de destino', 'Colaborador'):
            assert expected in text, expected
        for wrong in ('Almacen', 'Codigo', 'Descripcion', 'Entrego', 'Recibio', 'automaticamente'):
            assert wrong not in text, wrong

    def test_the_entry_vale_prints_invoice_delivery_note_and_purchase_order(self, entry_group):
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=entry_group.pk)))
        assert 'N° de factura' in text and 'F001-55' in text
        assert 'Guía de remisión' in text and 'T001-9' in text
        assert 'OC / PO' in text and 'OC-9' in text
        assert 'Factura / Guia' not in text

    def test_references_that_are_empty_are_not_printed(self, epp_item, central_warehouse):
        group = _entry(epp_item, 1, central_warehouse, date(2026, 10, 5), supplier_name='Proveedor Inventado')
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=group.pk)))
        assert 'N° de factura' in text and 'Proveedor Inventado' in text
        assert 'Guía de remisión' not in text and 'OC / PO' not in text

    def test_the_internal_transfer_source_is_called_transferencia_entrante(self, epp_item, central_warehouse):
        group = _entry(epp_item, 1, central_warehouse, date(2026, 10, 5), source_type='TRANSFER_IN')
        text = _pdf_text(generate_group_pdf(MovementGroup.objects.get(pk=group.pk)))
        assert 'Transferencia entrante' in text and 'Traslado' not in text

    def test_the_single_movement_pdf_follows_the_same_wording(self, epp_item, central_warehouse):
        movement = svc.register_entry(
            inventory=epp_item, quantity=Decimal('2'), warehouse='CENTRAL', project_id=None,
            source_type='PURCHASE', supplier_name='Proveedor Inventado', invoice_number='F001-56',
            notes='', registered_by=central_warehouse, movement_date=date(2026, 10, 5),
            voucher_number='V-9', guide_number='T001-10', po_number='OC-10',
        )
        text = _pdf_text(generate_movement_pdf(movement))
        assert 'VALE DE INGRESO' in text
        for expected in ('N° de factura', 'F001-56', 'Guía de remisión', 'T001-10', 'OC / PO', 'OC-10',
                         'N° de vale (físico)', 'V-9', 'Código', 'Descripción', 'Almacén Central'):
            assert expected in text, expected


class TestExcelWording:
    def test_the_inventory_export_has_accented_headers(self, epp_item, central_warehouse):
        resp = client_for(central_warehouse).get(INVENTORY + 'export/')
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        assert [c.value for c in sheet[1]] == [
            'Código', 'Descripción', 'Unidad', 'Categoría', 'Tipo',
            'Marca', 'Modelo', 'Ubicación', 'Stock total', 'Stock mínimo',
        ]

    def test_a_floor_only_item_exports_its_floor_not_a_blank(self, make_item, central_warehouse):
        make_item('EPP000010', 'CASCO', 'EA', stock=1, floor=2)
        make_item('EPP000011', 'LENTE', 'EA', stock=1, floor=2, location='F3-2')
        make_item('EPP000012', 'BOTA', 'EA', stock=1, location='Estante antiguo 3')
        resp = client_for(central_warehouse).get(INVENTORY + 'export/')
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        by_code = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=8).value
                   for r in range(2, sheet.max_row + 1)}
        assert by_code == {'EPP000010': 'PISO 2', 'EPP000011': 'F3-2', 'EPP000012': 'Estante antiguo 3'}

    def test_every_truncated_warehouse_export_uses_one_warning(self, epp_item, make_item, central_warehouse):
        make_item('EPP000002', 'CASCO', 'EA', stock=1)
        _entry(epp_item, 1, central_warehouse, date(2026, 10, 5))
        _entry(epp_item, 1, central_warehouse, date(2026, 10, 6))
        from apps.warehouse.services.exports import TRUNCATED_EXPORT_WARNING

        client = client_for(central_warehouse)
        with mock.patch('apps.core.utils.excel.MAX_EXPORT_ROWS', 1):
            for url in (INVENTORY + 'export/', MOVEMENTS + 'export/'):
                sheet = openpyxl.load_workbook(io.BytesIO(client.get(url).content)).active
                assert sheet.cell(row=1, column=1).value == TRUNCATED_EXPORT_WARNING, url
        assert 'exportación se truncó' in TRUNCATED_EXPORT_WARNING


class TestImportWording:
    def test_the_problems_report_uses_the_same_words_as_the_screen(self):
        assert reporting.PROBLEM_HEADERS == ['Gravedad', 'Hoja', 'Fila', 'Tipo de problema', 'Mensaje']

    def test_the_report_writes_the_readable_type_not_the_raw_code(self):
        book = reporting.build_problems_workbook([
            {'severity': 'warning', 'sheet': 'EGRESOS', 'row': 4, 'code': 'VALE_VACIO', 'message': 'x'},
            {'severity': 'error', 'sheet': 'X', 'row': None, 'code': 'CODIGO_NUEVO_SIN_ROTULO', 'message': 'y'},
        ])
        rows = list(book.active.iter_rows(values_only=True))
        assert rows[1][:5] == ('Advertencia', 'EGRESOS', 4, 'Egreso sin N° de vale', 'x')
        assert rows[2][3] == 'CODIGO_NUEVO_SIN_ROTULO'                      # unknown code: shown as is

    def test_every_code_of_the_catalog_has_a_readable_label(self):
        assert set(problems.CODE_LABELS) == set(problems.CATALOG)
        assert all(label and label != code and label[0].isupper() for code, label in problems.CODE_LABELS.items())

    def test_a_long_location_cell_cites_the_excel_column(self):
        _data, errors = validate_row(EntryRowSerializer, {
            'product_code': 'EPP000001', 'quantity': 1, 'location': 'x' * 101,
        })
        text = ' '.join(str(e) for e in errors)
        assert 'La columna LOCACIÓN supera los 100 caracteres.' in text
        assert 'locación' not in text


class TestModelWording:
    def test_the_internal_and_physical_vale_numbers_have_different_names(self):
        from apps.warehouse.models import InventoryMovement

        group_field = MovementGroup._meta.get_field('group_number')
        assert str(group_field.verbose_name) == 'N° de vale interno'
        assert str(MovementGroup._meta.get_field('voucher_number').verbose_name) == 'N° de vale (físico)'
        assert str(InventoryMovement._meta.get_field('voucher_number').verbose_name) == 'N° de vale (físico)'
        assert str(MovementGroup._meta.verbose_name_plural) == 'vales de movimiento'

    def test_the_location_help_text_states_the_official_format(self):
        from apps.warehouse.models import Inventory

        field = Inventory._meta.get_field('location')
        assert 'Pabellón' in str(field.help_text) and 'B3-2' in str(field.help_text)
        assert 'Estante' not in str(field.help_text)


# ---------------------------------------------------------------------------
# NUL en los parámetros de consulta (PostgreSQL no admite 0x00 en texto)
# ---------------------------------------------------------------------------

class TestNulQueryParameters:
    ENDPOINTS = [
        ('inventory/', 'search'), ('inventory/', 'category'), ('inventory/', 'item_type'),
        ('inventory/', 'warehouse'), ('inventory/export/', 'search'),
        ('inventory/movement-summary/', 'search'), ('inventory/movement-summary/', 'category'),
        ('inventory/movement-summary/export/', 'search'),
        ('inventory/check-stock/', 'descriptions'), ('inventory/alerts/', 'x'),
        ('movements/', 'search'), ('movements/', 'supplier'), ('movements/', 'voucher_number'),
        ('movements/', 'destination'), ('movements/', 'movement_type'), ('movements/', 'warehouse'),
        ('movements/export/', 'search'), ('movements/export/', 'supplier'),
        ('personal-lookup/', 'search'), ('imports/', 'status'), ('pending-requests/list/', 'search'),
    ]

    @pytest.mark.parametrize('path,name', ENDPOINTS)
    @pytest.mark.parametrize('payload', ['\x00', 'abc\x00def'])
    def test_a_nul_byte_in_a_value_is_a_400_with_the_project_envelope(self, central_warehouse, path, name, payload):
        client = client_for(central_warehouse)
        client.raise_request_exception = False
        resp = client.get(BASE + path, {name: payload})
        assert resp.status_code == 400, (path, name, resp.status_code)
        body = resp.json()
        assert body['status_code'] == 400 and body['error'] is True
        assert '\x00' not in str(body)

    def test_a_nul_byte_in_the_parameter_name_is_a_400_and_is_not_echoed(self, central_warehouse):
        resp = client_for(central_warehouse).get(INVENTORY + '?sea%00rch=x')
        assert resp.status_code == 400
        assert '\x00' not in str(resp.json())
        assert 'search' in str(resp.json())

    def test_the_kardex_of_one_item_rejects_nul_too(self, epp_item, central_warehouse):
        resp = client_for(central_warehouse).get(f'{INVENTORY}{epp_item.pk}/kardex/', {'movement_type': '\x00'})
        assert resp.status_code == 400

    def test_clean_parameters_keep_working(self, epp_item, central_warehouse):
        client = client_for(central_warehouse)
        resp = client.get(INVENTORY, {'search': 'guante'})
        assert resp.status_code == 200 and resp.data['count'] == 1
        assert client.get(MOVEMENTS, {'search': 'ñandú', 'supplier': 'A&B'}).status_code == 200

    def test_an_anonymous_caller_still_gets_401_not_400(self):
        from rest_framework.test import APIClient

        resp = APIClient().get(INVENTORY, {'search': '\x00'})
        assert resp.status_code == 401

    def test_a_role_without_access_still_gets_403_not_400(self):
        user = make_user('requester_x', RoleChoices.REQUESTER)
        assert client_for(user).get(INVENTORY, {'search': '\x00'}).status_code == 403

    def test_the_guard_only_looks_at_text_it_can_find(self):
        from django.http import QueryDict

        from apps.warehouse.query_guard import find_nul_parameter

        assert find_nul_parameter(QueryDict('a=1&b=2')) is None
        assert find_nul_parameter(QueryDict('a=1&b=%00')) == 'b'
        assert find_nul_parameter(QueryDict('a=1&a=%00')) == 'a'


# ---------------------------------------------------------------------------
# Orden de movimientos
# ---------------------------------------------------------------------------

class TestMovementOrdering:
    @pytest.fixture
    def book(self, epp_item, central_warehouse):
        """Registered out of chronological order, the way an imported book is."""
        _entry(epp_item, 5, central_warehouse, date(2026, 10, 10))
        _entry(epp_item, 1, central_warehouse, None)                                 # no date
        _entry(epp_item, 2, central_warehouse, date(2026, 10, 3))
        _exit(epp_item, 1, central_warehouse, date(2026, 10, 16))
        _entry(epp_item, 4, central_warehouse, date(2026, 10, 3))
        return epp_item

    def test_the_list_is_newest_operation_first_and_undated_rows_last(self, book, central_warehouse):
        rows = client_for(central_warehouse).get(MOVEMENTS, {'page_size': 50}).data['results']
        dates = [r['movement_date'] for r in rows]
        assert dates == ['2026-10-16', '2026-10-10', '2026-10-03', '2026-10-03', None]

    def test_the_same_day_keeps_the_latest_registration_first(self, book, central_warehouse):
        rows = client_for(central_warehouse).get(MOVEMENTS, {'page_size': 50}).data['results']
        same_day = [r for r in rows if r['movement_date'] == '2026-10-03']
        assert [float(r['quantity']) for r in same_day] == [4.0, 2.0]

    def test_pages_do_not_repeat_or_skip_rows(self, book, central_warehouse):
        client = client_for(central_warehouse)
        seen = []
        for page in (1, 2, 3):
            seen += [r['id'] for r in client.get(MOVEMENTS, {'page_size': 2, 'page': page}).data['results']]
        assert len(seen) == 5 and len(set(seen)) == 5

    def test_the_export_follows_the_same_order(self, book, central_warehouse):
        resp = client_for(central_warehouse).get(MOVEMENTS + 'export/')
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        dates = [sheet.cell(row=r, column=3).value for r in range(2, sheet.max_row + 1)]
        assert dates == ['16/10/2026', '10/10/2026', '03/10/2026', '03/10/2026', 'Sin fecha']

    def test_recent_movements_of_an_item_follow_the_operational_date(self, book, central_warehouse):
        rows = client_for(central_warehouse).get(f'{INVENTORY}{book.pk}/recent-movements/').data
        assert [r['movement_date'] for r in rows][:3] == ['2026-10-16', '2026-10-10', '2026-10-03']


# ---------------------------------------------------------------------------
# Stock real vs. movimientos
# ---------------------------------------------------------------------------

class TestRealStockVersusMovements:
    @pytest.fixture
    def mixed(self, make_item, central_warehouse):
        """A: 6 units seeded with no movement. B: 10 units, all backed by movements."""
        seeded = make_item('EPP000001', 'GUANTE T: L', 'EA', stock=6)
        tracked = make_item('EPP000002', 'CASCO', 'EA', stock=0)
        _entry(tracked, 12, central_warehouse, date(2026, 10, 5))
        _exit(tracked, 2, central_warehouse, date(2026, 10, 6))
        return seeded, tracked

    def test_totals_carry_the_real_stock_next_to_the_movement_totals(self, mixed, central_warehouse):
        body = client_for(central_warehouse).get(SUMMARY).data
        totals = body['meta']['totals']
        assert totals['closing'] == 10          # what the movements add up to
        assert totals['current'] == 16          # InventoryStock: what Inventario shows
        assert totals['current'] == float(sum(s.quantity for s in InventoryStock.objects.all()))

    def test_the_stock_without_movements_is_reported_per_row_and_in_total(self, mixed, central_warehouse):
        body = client_for(central_warehouse).get(SUMMARY).data
        rows = {r['product_code']: r for r in body['results']}
        assert rows['EPP000001']['untracked_stock'] == 6 and rows['EPP000001']['closing_stock'] == 0
        assert rows['EPP000001']['current_stock'] == 6
        assert rows['EPP000002']['untracked_stock'] == 0
        assert body['meta']['untracked_stock'] == {'items': 1, 'units': 6.0}

    def test_with_a_date_filter_the_difference_is_not_reported(self, mixed, central_warehouse):
        body = client_for(central_warehouse).get(SUMMARY, {'date_from': '2026-10-01'}).data
        assert body['meta']['untracked_stock'] is None
        assert all(r['untracked_stock'] is None for r in body['results'])
        assert body['meta']['totals']['current'] == 16

    def test_an_item_whose_movements_explain_its_stock_is_not_flagged(self, make_item, central_warehouse):
        item = make_item('EPP000003', 'LENTE', 'EA', stock=0)
        _entry(item, 5, central_warehouse, date(2026, 10, 5))
        _entry(item, 1, central_warehouse, None)                                    # undated still counts
        body = client_for(central_warehouse).get(SUMMARY).data
        assert body['meta']['untracked_stock'] == {'items': 0, 'units': 0.0}

    def test_no_page_responses_carry_the_same_fields(self, mixed, central_warehouse):
        body = client_for(central_warehouse).get(SUMMARY, {'no_page': 'true'}).data
        assert body['meta']['totals']['current'] == 16
        assert {r['product_code']: r['untracked_stock'] for r in body['results']} == {
            'EPP000001': 6, 'EPP000002': 0,
        }

    def test_the_dashboard_summary_exposes_the_real_units_in_stock(self, mixed, central_warehouse):
        body = client_for(central_warehouse).get(INVENTORY + 'summary/').data
        assert body['total_units'] == 16.0

    def test_the_dashboard_summary_with_no_stock_is_zero(self, central_warehouse):
        assert client_for(central_warehouse).get(INVENTORY + 'summary/').data['total_units'] == 0.0


# ---------------------------------------------------------------------------
# Importación: filas importadas que el libro ya no trae
# ---------------------------------------------------------------------------

class TestMissingImportedRowsInSummary:
    def test_the_summary_counts_them_and_the_warning_explains_the_consequence(self):
        message = problems.make_problem('FILAS_IMPORTADAS_AUSENTES', count=3).message
        assert '3 movimiento(s)' in message
        assert 'no se actualizan' in message and 'duplicarse' in message and 'ajuste' in message

    def test_a_first_import_reports_zero_missing_rows(self, central_warehouse):
        from .import_workbooks import standard_workbook
        from .qa_support import upload

        client = client_for(central_warehouse)
        body = upload(client, standard_workbook()).json()
        assert body['summary']['imported_rows_missing'] == 0
