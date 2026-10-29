"""
SYSPCC-025 · QA ronda 2 — ramas de la API que la medición de cobertura dejó sin probar: listados sin
paginar, exportaciones con filtros y truncadas, topes del mapeo de confirmación, identificadores que no
son UUID, el piso inválido del ítem y la numeración cuando el último número guardado no es un entero.

Sin libros reales ni nombres reales: todo es inventado.
"""
import io
import re
from unittest import mock

import openpyxl
import pytest
from django.utils import timezone

from apps.warehouse.models import InventoryMovement, MovementGroup
from apps.warehouse.serializers_imports import MAX_DESTINATION_MAP, MAX_PROJECT_MAP
from apps.warehouse.services import movements as svc

from .import_workbooks import D1, standard_workbook
from .qa_support import IMPORTS, INVENTORY, MOVEMENTS, confirm, snapshot_database, tolerant_client, upload

pytestmark = pytest.mark.django_db

SUMMARY = INVENTORY + 'movement-summary/'


@pytest.fixture(autouse=True)
def _no_pdf_task():
    with mock.patch('apps.warehouse.views._enqueue_group_pdf_upload'):
        yield


@pytest.fixture
def client(central_warehouse):
    return tolerant_client(central_warehouse)


def _entry(item, qty, user, *, warehouse='CENTRAL', when=None):
    return svc.register_entry_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
        source_type='PURCHASE', supplier_name='', invoice_number='', notes='', registered_by=user,
        movement_date=when,
    )


def _rows(book, sheet):
    return list(book[sheet].iter_rows(min_row=2, values_only=True))


# ---------------------------------------------------------------------------
# Listados sin paginar y exportaciones
# ---------------------------------------------------------------------------

class TestUnpaginatedLists:
    def test_inventory_no_page_returns_every_item_as_a_plain_list(self, client, make_item):
        for n in range(60):                                          # more than one page (25)
            make_item(f'EPP{n:06d}', f'ITEM DE PRUEBA {n}', 'EA')
        body = client.get(INVENTORY, {'no_page': 'true'}).json()
        assert isinstance(body, list) and len(body) == 60
        assert {'product_code', 'description', 'total_stock'} <= set(body[0])

    def test_inventory_without_no_page_is_paginated(self, client, make_item):
        for n in range(30):
            make_item(f'EPP{n:06d}', f'ITEM DE PRUEBA {n}', 'EA')
        body = client.get(INVENTORY).json()
        assert body['count'] == 30 and len(body['results']) == 25

    def test_movements_no_page_returns_every_movement_as_a_plain_list(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        for _ in range(30):                                          # one movement per call, past a page
            _entry(item, 1, central_warehouse)
        body = client.get(MOVEMENTS, {'no_page': 'true'}).json()
        assert isinstance(body, list) and len(body) == 30

    def test_movements_no_page_still_honours_the_filters(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 1, central_warehouse, warehouse='CENTRAL')
        _entry(item, 1, central_warehouse, warehouse='SITE')
        body = client.get(MOVEMENTS, {'no_page': 'true', 'warehouse': 'SITE'}).json()
        assert [m['warehouse'] for m in body] == ['SITE']


class TestSummaryExportWithMovements:
    def test_the_warehouse_filter_also_limits_the_movement_sheets(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 7, central_warehouse, warehouse='CENTRAL', when=D1)
        _entry(item, 5, central_warehouse, warehouse='SITE', when=D1)
        response = client.get(SUMMARY + 'export/', {'include_movements': 'true', 'warehouse': 'CENTRAL'})
        assert response.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert [row[5] for row in _rows(book, 'INGRESOS')] == [7]

    def test_without_the_flag_the_movement_sheets_are_not_written(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 7, central_warehouse, when=D1)
        book = openpyxl.load_workbook(io.BytesIO(client.get(SUMMARY + 'export/').content))
        assert book.sheetnames == ['RESUMEN DE MOVIMIENTO']

    def test_with_the_flag_all_three_sheets_are_written(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 7, central_warehouse, when=D1)
        book = openpyxl.load_workbook(io.BytesIO(client.get(SUMMARY + 'export/', {'include_movements': 'true'}).content))
        assert book.sheetnames == ['RESUMEN DE MOVIMIENTO', 'INGRESOS', 'EGRESOS']


class TestPlainMovementsExportTruncation:
    def test_a_truncated_export_starts_with_a_visible_warning_and_then_the_header(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        for _ in range(3):
            _entry(item, 1, central_warehouse, when=D1)
        with mock.patch('apps.core.utils.excel.MAX_EXPORT_ROWS', 1):
            response = client.get(MOVEMENTS + 'export/')
        assert response.status_code == 200
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        assert 'ADVERTENCIA' in sheet.cell(row=1, column=1).value
        assert sheet.cell(row=2, column=1).value == 'N° de movimiento'
        assert sheet.max_row == 3                                    # warning + header + the one row kept

    def test_an_export_under_the_cap_has_no_warning(self, client, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 1, central_warehouse, when=D1)
        sheet = openpyxl.load_workbook(io.BytesIO(client.get(MOVEMENTS + 'export/').content)).active
        assert sheet.cell(row=1, column=1).value == 'N° de movimiento'


# ---------------------------------------------------------------------------
# Confirmación: topes del mapeo, identificadores que no son UUID
# ---------------------------------------------------------------------------

class TestConfirmBodyLimits:
    @pytest.fixture
    def preview(self, client):
        return upload(client, standard_workbook()).json()

    def test_more_project_labels_than_the_cap_is_a_400_and_nothing_is_written(self, client, preview):
        before = snapshot_database()
        response = confirm(client, preview, project_map={f'ETIQUETA {n}': None for n in range(MAX_PROJECT_MAP + 1)})
        assert response.status_code == 400, response.content[:200]
        assert 'project_map' in response.content.decode()
        assert snapshot_database() == before

    def test_more_destination_links_than_the_cap_is_a_400_and_nothing_is_written(self, client, preview):
        before = snapshot_database()
        response = confirm(client, preview, destination_map={f'clave-{n}': 1 for n in range(MAX_DESTINATION_MAP + 1)})
        assert response.status_code == 400, response.content[:200]
        assert 'destination_map' in response.content.decode()
        assert snapshot_database() == before

    def test_exactly_the_cap_of_project_labels_passes_this_validation(self, client, preview):
        response = confirm(client, preview, project_map={'ON PLANT': None, **{f'E{n}': None for n in range(MAX_PROJECT_MAP - 1)}})
        assert response.status_code != 400 or 'project_map' not in response.content.decode()

    @pytest.mark.parametrize('suffix', ['', 'confirm/', 'problems/', 'problems/export/'])
    @pytest.mark.parametrize('identifier', ['no-es-un-uuid', '12345', 'zzzzzzzz-zzzz-zzzz-zzzz-zzzzzzzzzzzz'])
    def test_an_identifier_that_is_not_a_uuid_is_a_404_not_a_crash(self, client, identifier, suffix):
        method = client.post if suffix == 'confirm/' else client.get
        response = method(f'{IMPORTS}{identifier}/{suffix}')
        assert response.status_code == 404, (response.status_code, response.content[:160])

    def test_cancelling_an_identifier_that_is_not_a_uuid_is_a_404(self, client):
        assert client.delete(f'{IMPORTS}no-es-un-uuid/').status_code == 404


# ---------------------------------------------------------------------------
# Piso y ubicación exacta del ítem
# ---------------------------------------------------------------------------

class TestFloorFollowsTheExactLocation:
    @pytest.mark.parametrize('floor', [0, 3, -1, 'dos', 99])
    def test_a_floor_that_is_not_a_valid_choice_is_a_400_and_the_item_keeps_its_floor(self, client, make_item, floor):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA', location='F3-2', floor=2)
        response = client.patch(f'{INVENTORY}{item.pk}/', {'floor': floor}, format='json')
        assert response.status_code == 400, response.content[:200]
        item.refresh_from_db()
        assert (item.location, item.floor) == ('F3-2', 2)

    def test_a_floor_that_contradicts_the_location_is_a_400(self, client, make_item):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        response = client.patch(f'{INVENTORY}{item.pk}/', {'location': 'F3-2', 'floor': 1}, format='json')
        assert response.status_code == 400
        item.refresh_from_db()
        assert item.location == '' and item.floor is None

    def test_changing_the_location_moves_the_floor_with_it(self, client, make_item):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA', location='F3-2', floor=2)
        response = client.patch(f'{INVENTORY}{item.pk}/', {'location': 'A1-1'}, format='json')
        assert response.status_code == 200, response.content[:200]
        item.refresh_from_db()
        assert (item.location, item.floor) == ('A1-1', 1)


# ---------------------------------------------------------------------------
# Numeración cuando el último número guardado no termina en un entero
# ---------------------------------------------------------------------------

def _adjust(item, qty, user):
    return svc.register_adjustment(
        inventory=item, quantity=qty, warehouse='CENTRAL', project_id=None, notes='', registered_by=user,
    )


class TestNumberingFallback:
    """
    ``AJU-AAAA-####`` (ajustes, sin vale) y ``ENT/SAL-AAAA-G####`` (vales). Los números salen siempre bien
    formados, pero un dato corrupto no debe dejar al almacén sin poder registrar.
    """

    def test_an_adjustment_number_that_does_not_end_in_an_integer_does_not_stop_the_next_adjustment(
        self, make_item, central_warehouse,
    ):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA', stock=10)
        _adjust(item, 1, central_warehouse)
        year = timezone.now().year
        InventoryMovement.objects.filter(movement_type='ADJUSTMENT').update(movement_number=f'AJU-{year}-ZZZZ')

        _adjust(item, 2, central_warehouse)
        numbers = list(InventoryMovement.objects.filter(movement_type='ADJUSTMENT').values_list('movement_number', flat=True))
        assert len(numbers) == len(set(numbers)) == 2
        assert any(re.fullmatch(rf'AJU-{year}-\d{{4}}', n) for n in numbers)

    def test_a_group_number_that_does_not_end_in_an_integer_does_not_stop_the_next_vale(
        self, make_item, central_warehouse,
    ):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 1, central_warehouse)
        year = timezone.now().year
        MovementGroup.objects.update(group_number=f'ENT-{year}-GZZZ')
        InventoryMovement.objects.update(movement_number=f'ENT-{year}-GZZZ-01')

        _entry(item, 2, central_warehouse)
        numbers = list(MovementGroup.objects.values_list('group_number', flat=True))
        assert len(numbers) == len(set(numbers)) == 2
        assert any(re.fullmatch(rf'ENT-{year}-G\d{{4}}', n) for n in numbers)

    def test_adjustment_numbers_keep_counting_past_9999(self, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA', stock=10)
        _adjust(item, 1, central_warehouse)
        year = timezone.now().year
        InventoryMovement.objects.update(movement_number=f'AJU-{year}-9999')

        _adjust(item, 1, central_warehouse)
        assert set(InventoryMovement.objects.values_list('movement_number', flat=True)) == {
            f'AJU-{year}-9999', f'AJU-{year}-10000',
        }

    def test_vale_numbers_keep_counting_past_9999(self, make_item, central_warehouse):
        item = make_item('EPP000001', 'GUANTE T: L', 'EA')
        _entry(item, 1, central_warehouse)
        year = timezone.now().year
        MovementGroup.objects.update(group_number=f'ENT-{year}-G9999')
        InventoryMovement.objects.update(movement_number=f'ENT-{year}-G9999-01')

        _entry(item, 1, central_warehouse)
        assert set(MovementGroup.objects.values_list('group_number', flat=True)) == {
            f'ENT-{year}-G9999', f'ENT-{year}-G10000',
        }
