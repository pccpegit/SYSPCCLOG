"""
SYSPCC-025 · QA ronda 2 — ramas de error y de borde que la medición de cobertura dejó sin probar, y
la conciliación de los exports con ventana de fechas.

* "DEFECTO QA-n": comportamiento correcto de algo que hoy está roto; falla a propósito y pasa solo
  cuando se corrige la producción (misma convención que la ronda 1).
* El resto son pruebas de ramas de error: conflicto de numeración (409), errores de servicio (400),
  destinatario de la entrega de EPP, filtros del catálogo y del listado, vales de movimientos sueltos.
"""
import io
from datetime import date
from decimal import Decimal
from unittest import mock

import openpyxl
import pytest
from django.db import IntegrityError
from django.utils import timezone

from apps.warehouse import views
from apps.warehouse.models import Inventory, InventoryMovement, InventoryStock
from apps.warehouse.services import catalog, epp, kardex
from apps.warehouse.services import movements as svc

from .qa_support import INVENTORY, MOVEMENTS, snapshot_database, tolerant_client

pytestmark = pytest.mark.django_db

SUMMARY = INVENTORY + 'movement-summary/'


@pytest.fixture(autouse=True)
def _no_pdf_task():
    with mock.patch('apps.warehouse.views._enqueue_group_pdf_upload') as enqueue:
        yield enqueue


@pytest.fixture
def client(central_warehouse):
    return tolerant_client(central_warehouse)


@pytest.fixture
def item(make_item):
    return make_item('EPP000001', 'GUANTE T: L', 'EA', stock=10)


def _entry(item, qty, user, when, source='PURCHASE'):
    return svc.register_entry_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        source_type=source, supplier_name='', invoice_number='', notes='', registered_by=user, movement_date=when)


def _exit(item, qty, user, when):
    return svc.register_exit_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        destination_type='EMPLOYEE', destination_detail='Persona Prueba', requested_by_id=None,
        authorized_by_id=None, notes='', registered_by=user, movement_date=when)


# ---------------------------------------------------------------------------
# Export con ventana de fechas: las hojas deben sumar lo mismo que el resumen
# ---------------------------------------------------------------------------

LOAD_DAY, ENTRY_DAY, EXIT_DAY = date(2026, 10, 5), date(2026, 10, 8), date(2026, 10, 12)
WINDOWS = {
    'sin ventana': {},
    'desde antes de la carga inicial': {'date_from': '2026-10-01'},
    'desde el día de la carga inicial': {'date_from': '2026-10-05'},
    'desde después de la carga inicial': {'date_from': '2026-10-06'},
    'hasta antes del egreso': {'date_to': '2026-10-10'},
    'ventana cerrada que contiene todo': {'date_from': '2026-10-01', 'date_to': '2026-10-31'},
}


@pytest.fixture
def history(item, central_warehouse):
    item.stocks.all().delete()
    _entry(item, 10, central_warehouse, LOAD_DAY, source='INITIAL_LOAD')
    _entry(item, 4, central_warehouse, ENTRY_DAY)
    _exit(item, 3, central_warehouse, EXIT_DAY)
    return item


def _sheet_sum(workbook, name, column):
    return sum(row[column] or 0 for row in workbook[name].iter_rows(min_row=2, values_only=True))


class TestWindowedExportReconciles:
    @pytest.mark.parametrize('params', WINDOWS.values(), ids=WINDOWS.keys())
    def test_defect_qa21_the_movement_sheets_add_up_to_the_summary_columns_of_the_same_window(
        self, history, client, params,
    ):
        """
        DEFECTO QA-21: con ``date_from`` en o antes de la fecha de la carga inicial, el resumen cuenta la
        carga como INGRESO del periodo (14), pero la hoja INGRESOS del export la deja fuera (suma 4) porque
        solo "viaja en INV. INICIAL", que en esa ventana vale 0: el libro exportado queda en desacuerdo
        consigo mismo y, reimportado, pierde 10 unidades.
        """
        row = client.get(SUMMARY, {'no_page': 'true', **params}).json()['results'][0]
        response = client.get(SUMMARY + 'export/', {'include_movements': 'true', **params})
        assert response.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(response.content))

        entries_in_sheet = _sheet_sum(book, 'INGRESOS', 5)
        exits_in_sheet = _sheet_sum(book, 'EGRESOS', 5)
        # what the summary row says must be recoverable from the movement sheets plus INV. INICIAL
        assert row['initial_stock'] + entries_in_sheet - exits_in_sheet == pytest.approx(row['closing_stock'])
        assert exits_in_sheet == pytest.approx(row['exits'])

    @pytest.mark.parametrize('params', [{'date_from': '2026-10-06', 'date_to': '2026-10-10'}, {'date_from': '2026-10-09'}])
    def test_only_the_movements_inside_the_window_are_written(self, history, client, params):
        book = openpyxl.load_workbook(io.BytesIO(
            client.get(SUMMARY + 'export/', {'include_movements': 'true', **params}).content,
        ))
        dates = [r[1] for sheet in ('INGRESOS',) for r in book[sheet].iter_rows(min_row=2, values_only=True)]
        start = date.fromisoformat(params['date_from'])
        end = date.fromisoformat(params['date_to']) if 'date_to' in params else None
        assert all(start <= d.date() and (end is None or d.date() <= end) for d in dates)


# ---------------------------------------------------------------------------
# Conflicto de numeración y errores de servicio
# ---------------------------------------------------------------------------

def _body(kind, item):
    today = timezone.localdate().isoformat()
    if kind == 'entry':
        return MOVEMENTS + 'entry/', {'inventory': item.pk, 'quantity': 1, 'warehouse': 'CENTRAL',
                                      'source_type': 'PURCHASE'}
    if kind == 'exit':
        return MOVEMENTS + 'exit/', {'inventory': item.pk, 'quantity': 1, 'warehouse': 'CENTRAL',
                                     'destination_type': 'PROJECT', 'destination_detail': 'Obra'}
    if kind == 'batch-entry':
        return MOVEMENTS + 'batch-entry/', {'items': [{'inventory': item.pk, 'quantity': 1}], 'warehouse': 'CENTRAL',
                                            'source_type': 'PURCHASE'}
    if kind == 'batch-exit':
        return MOVEMENTS + 'batch-exit/', {'items': [{'inventory': item.pk, 'quantity': 1}], 'warehouse': 'CENTRAL',
                                           'destination_type': 'PROJECT', 'destination_detail': 'Obra'}
    return MOVEMENTS + 'epp-delivery/', {'warehouse': 'CENTRAL', 'voucher_number': 'V-1', 'movement_date': today,
                                         'recipient': {'name': 'Persona Prueba'},
                                         'items': [{'inventory': item.pk, 'quantity': 1}]}


SERVICE_FOR = {
    'entry': (views.warehouse_svc, 'register_entry'),
    'exit': (views.warehouse_svc, 'register_exit'),
    'batch-entry': (views.warehouse_svc, 'register_entry_batch'),
    'batch-exit': (views.warehouse_svc, 'register_exit_batch'),
    'epp-delivery': (views.epp_svc, 'register_epp_delivery'),
}


@pytest.mark.parametrize('kind', list(SERVICE_FOR))
class TestWriteEndpointsMapServiceErrors:
    def test_a_numbering_collision_is_a_409_the_client_can_retry_and_nothing_is_written(self, client, item, kind):
        url, body = _body(kind, item)
        before = snapshot_database()
        module, name = SERVICE_FOR[kind]
        with mock.patch.object(module, name, side_effect=IntegrityError('duplicate movement_number')):
            response = client.post(url, body, format='json')
        assert response.status_code == 409, (response.status_code, response.content[:200])
        assert 'Intente nuevamente' in response.json()['detail']
        assert snapshot_database() == before

    def test_a_business_error_from_the_service_is_a_400_with_its_message(self, client, item, kind):
        url, body = _body(kind, item)
        module, name = SERVICE_FOR[kind]
        with mock.patch.object(module, name, side_effect=ValueError('Stock insuficiente para el ítem.')):
            response = client.post(url, body, format='json')
        assert response.status_code == 400
        assert response.json()['detail'] == 'Stock insuficiente para el ítem.'

    def test_an_unexpected_error_is_not_swallowed(self, client, item, kind):
        """Anything that is neither a business error nor a numbering race must surface as a 500 (and be logged)."""
        url, body = _body(kind, item)
        module, name = SERVICE_FOR[kind]
        with mock.patch.object(module, name, side_effect=RuntimeError('disco lleno')):
            response = client.post(url, body, format='json')
        assert response.status_code == 500


# ---------------------------------------------------------------------------
# Servicio de entrega de EPP, directo
# ---------------------------------------------------------------------------

class TestEppServiceRules:
    def _call(self, item, user, **overrides):
        args = dict(items=[{'inventory_id': item.pk, 'quantity': 1}], warehouse='CENTRAL', voucher_number='V-1',
                    movement_date=timezone.localdate(), registered_by=user, recipient_name='Persona Prueba')
        args.update(overrides)
        return epp.register_epp_delivery(**args)

    def test_no_recipient_is_refused(self, item, central_warehouse):
        with pytest.raises(epp.RecipientError):
            self._call(item, central_warehouse, recipient_name='')
        assert InventoryMovement.objects.count() == 0

    def test_a_blank_recipient_is_no_recipient(self, item, central_warehouse):
        with pytest.raises(epp.RecipientError):
            self._call(item, central_warehouse, recipient_name='   ', internal_destination='  ')

    def test_two_recipients_are_refused(self, item, central_warehouse, make_personal):
        with pytest.raises(epp.RecipientError):
            self._call(item, central_warehouse, personal=make_personal(), recipient_name='Otra Persona')
        with pytest.raises(epp.RecipientError):
            self._call(item, central_warehouse, internal_destination='ALM - PRIMER PISO')
        assert InventoryMovement.objects.count() == 0

    @pytest.mark.parametrize('voucher', ['', '   ', None])
    def test_a_blank_voucher_is_refused(self, item, central_warehouse, voucher):
        with pytest.raises(ValueError, match='vale'):
            self._call(item, central_warehouse, voucher_number=voucher)
        assert InventoryMovement.objects.count() == 0

    def test_the_same_voucher_twice_is_refused_unless_forced_and_the_stock_moves_once_per_accepted_delivery(
        self, item, central_warehouse,
    ):
        self._call(item, central_warehouse)
        with pytest.raises(epp.DuplicateVoucherError):
            self._call(item, central_warehouse)
        assert InventoryStock.objects.get(inventory=item).quantity == 9
        self._call(item, central_warehouse, allow_duplicate_voucher=True)
        assert InventoryStock.objects.get(inventory=item).quantity == 8

    def test_an_internal_destination_is_a_transfer_out_not_an_employee(self, item, central_warehouse):
        group = self._call(item, central_warehouse, recipient_name='', internal_destination='ALM - PRIMER PISO',
                           internal_transfer='2do piso → 1er piso')
        movement = group.movements.get()
        assert (movement.destination_type, movement.destination_detail) == ('TRANSFER_OUT', 'ALM - PRIMER PISO')
        assert movement.destination_personal_id is None

    def test_a_linked_worker_keeps_the_name_as_a_snapshot(self, item, central_warehouse, make_personal):
        worker = make_personal('Mamani Quispe, Luis Test')
        movement = self._call(item, central_warehouse, recipient_name='', personal=worker).movements.get()
        assert movement.destination_personal_id == worker.pk and movement.destination_detail == 'Mamani Quispe, Luis Test'

    def test_the_delivery_logs_ids_and_counters_never_the_recipient(self, item, central_warehouse, app_logs):
        self._call(item, central_warehouse, recipient_name='Nombre Secreto Prueba')
        assert 'Secreto' not in app_logs.dump()


# ---------------------------------------------------------------------------
# Ayudantes puros del catálogo y del kardex
# ---------------------------------------------------------------------------

class TestPureHelpers:
    @pytest.mark.parametrize('raw,expected', [
        (None, ''), ('', ''), ('  epp000001 ', 'EPP000001'), (100234, '100234'), (100234.0, '100234'),
        (12.5, '12.5'), ('ab-12', 'AB-12'),
    ])
    def test_normalize_product_code(self, raw, expected):
        assert catalog.normalize_product_code(raw) == expected

    @pytest.mark.parametrize('description,size', [
        ('GUANTE DE CUERO T: L', 'L'), ('BOTA T:42', '42'), ('lente t: xl ', 'XL'), ('CASCO BLANCO', ''),
        ('', ''), (None, ''), ('GUANTE T: ' + 'X' * 30, ''),
    ])
    def test_derive_size(self, description, size):
        assert catalog.derive_size(description) == size

    def test_signed_quantity_of_every_movement_kind(self):
        def mv(kind, source='', destination='', qty='3'):
            return InventoryMovement(movement_type=kind, source_type=source, destination_type=destination,
                                     quantity=Decimal(qty))
        assert kardex.signed_quantity(mv('ENTRY')) == 3
        assert kardex.signed_quantity(mv('EXIT')) == -3
        assert kardex.signed_quantity(mv('ADJUSTMENT', source='ADJUSTMENT')) == 3
        assert kardex.signed_quantity(mv('ADJUSTMENT', destination='ADJUSTMENT')) == -3
        assert kardex.signed_quantity(mv('ADJUSTMENT')) == 0               # an adjustment with no direction moves nothing
        assert kardex.signed_quantity(mv('TRANSFER')) == 0                 # unused type: ignored, never guessed

    def test_undated_movements_are_counted_per_warehouse(self, item, central_warehouse):
        for warehouse, qty in (('CENTRAL', 2), ('SITE', 3)):
            svc.register_entry_batch(
                items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
                source_type='PURCHASE', supplier_name='', invoice_number='', notes='',
                registered_by=central_warehouse, undated=True)
        ids = Inventory.objects.values('pk')
        assert kardex.count_undated_movements(ids) == 2
        assert kardex.count_undated_movements(ids, 'SITE') == 1
        assert kardex.count_undated_movements(ids, 'OFFICE') == 0


# ---------------------------------------------------------------------------
# Filtros del catálogo y del listado de movimientos
# ---------------------------------------------------------------------------

class TestCatalogFilters:
    @pytest.fixture
    def catalog_items(self, make_item):
        return {
            'low': make_item('EPP1', 'BAJO MINIMO', 'EA', stock=2, min_stock=5),
            'ok': make_item('EPP2', 'SUFICIENTE', 'EA', stock=20, min_stock=5),
            'zero': make_item('EPP3', 'AGOTADO', 'EA', stock=0, min_stock=5),
            'none': make_item('EPP4', 'SIN FILA DE STOCK', 'EA', min_stock=0),
        }

    def codes(self, client, **params):
        response = client.get(INVENTORY, {'no_page': 'true', **params})
        assert response.status_code == 200, response.content
        return sorted(row['product_code'] for row in response.json())

    def test_low_stock_includes_the_depleted_ones(self, client, catalog_items):
        assert self.codes(client, low_stock='true') == ['EPP1', 'EPP3', 'EPP4']

    def test_low_stock_with_stock_is_below_minimum_but_not_depleted(self, client, catalog_items):
        assert self.codes(client, low_stock='true', has_stock='true') == ['EPP1']

    def test_has_stock_true_and_false_partition_the_catalog(self, client, catalog_items):
        with_stock = self.codes(client, has_stock='true')
        without = self.codes(client, has_stock='false')
        assert with_stock == ['EPP1', 'EPP2']
        assert without == ['EPP3', 'EPP4']
        assert sorted(with_stock + without) == ['EPP1', 'EPP2', 'EPP3', 'EPP4']

    def test_the_unpaginated_catalog_returns_a_plain_list(self, client, catalog_items):
        response = client.get(INVENTORY, {'no_page': 'true'})
        assert isinstance(response.json(), list) and len(response.json()) == 4

    def test_an_item_without_any_stock_row_reads_zero(self, client, catalog_items):
        rows = {r['product_code']: r for r in client.get(INVENTORY, {'no_page': 'true'}).json()}
        assert float(rows['EPP4']['total_stock']) == 0


class TestMovementListFilters:
    @pytest.fixture
    def moves(self, item, central_warehouse, project):
        svc.register_entry_batch(
            items_data=[{'inventory_id': item.pk, 'quantity': 2}], warehouse='CENTRAL', project_id=project.pk,
            source_type='PURCHASE', supplier_name='Proveedor Prueba', invoice_number='F-1', notes='',
            registered_by=central_warehouse, movement_date=date(2026, 10, 6))
        svc.register_entry_batch(
            items_data=[{'inventory_id': item.pk, 'quantity': 7}], warehouse='CENTRAL', project_id=None,
            source_type='PURCHASE', supplier_name='Otro Proveedor', invoice_number='F-2', notes='',
            registered_by=central_warehouse, movement_date=date(2026, 10, 9))
        return project

    def numbers(self, client, **params):
        response = client.get(MOVEMENTS, params)
        assert response.status_code == 200, response.content
        return sorted(Decimal(str(r['quantity'])) for r in response.json()['results'])

    def test_project_filter(self, client, moves):
        assert self.numbers(client, project=moves.pk) == [2]

    def test_supplier_filter_is_a_case_insensitive_contains(self, client, moves):
        assert self.numbers(client, supplier='otro') == [7]

    def test_quantity_range(self, client, moves):
        assert self.numbers(client, qty_min='3') == [7]
        assert self.numbers(client, qty_max='3') == [2]
        assert self.numbers(client, qty_min='2', qty_max='7') == [2, 7]
        assert self.numbers(client, qty_min='8') == []

    def test_the_unpaginated_listing_is_a_plain_list(self, client, moves):
        response = client.get(MOVEMENTS, {'no_page': 'true'})
        assert isinstance(response.json(), list) and len(response.json()) == 2

    def test_date_filters_are_inclusive(self, client, moves):
        assert self.numbers(client, date_from='2026-10-06', date_to='2026-10-06') == [2]
        assert self.numbers(client, date_from='2026-10-09', date_to='2026-10-09') == [7]


class TestVoucherOfALegacyMovementWithoutGroup:
    def test_it_lists_itself_as_the_only_line(self, client, item, central_warehouse):
        """Movements created before vales existed have no group: the voucher still renders."""
        movement = InventoryMovement.objects.create(
            movement_number='ENT-2020-0001', movement_type='ENTRY', inventory=item, quantity=Decimal('3'),
            warehouse='CENTRAL', source_type='PURCHASE', registered_by=central_warehouse,
            movement_date=date(2026, 10, 6),
        )
        response = client.get(f'{MOVEMENTS}{movement.pk}/voucher/')
        assert response.status_code == 200
        data = response.json()
        assert data['group_number'] is None
        assert [(i['product_code'], i['quantity']) for i in data['items']] == [(item.product_code, 3.0)]
        assert client.get(f'{MOVEMENTS}{movement.pk}/pdf/').status_code == 200
