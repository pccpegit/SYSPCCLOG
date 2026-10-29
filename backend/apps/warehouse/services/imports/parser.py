"""
Safe parsing of the inventory workbook (SYSPCC-025). No database access here.

Threat model: the .xlsx is hostile input. Before openpyxl sees it we check the
extension, content type, size, ZIP signature and ZIP structure (entry count, path
traversal, uncompressed size and compression ratio = zip bomb, macros, DTD/entity
declarations = XXE / billion laughs). Then it is read with ``read_only=True`` and
``data_only=True``: formulas are NEVER evaluated (cached values are used) and
external links are not loaded.

Facts about the real workbook that shape the reader (verified against it):

* ``INGRESOS`` / ``EGRESOS`` declare ~1 M rows in their dimension tag but hold a few
  hundred. ``max_row`` is therefore ignored: reading stops after
  ``MAX_CONSECUTIVE_BLANK_ROWS`` blank rows and a hard cap of ``MAX_DATA_ROWS``.
* DESCRIPCIÓN / UOM / LOCACIÓN of the movement sheets are VLOOKUP formulas and the
  summary sheet is made of SUMIFS: only input columns decide whether a row is blank.
* The summary export of SYSPCC can put a warning row BEFORE the header (truncated
  export), so the header row is searched in the first ``HEADER_SCAN_ROWS`` rows.
"""
import io
import os
import re
import struct
import unicodedata
import zipfile
import zlib
from dataclasses import dataclass, field
from xml.etree import ElementTree

import logging

from openpyxl import load_workbook

from apps.warehouse.services.imports.errors import ImportFileError
from apps.warehouse.services.imports.problems import make_problem

logger = logging.getLogger(__name__)

MB = 1024 * 1024
MAX_UPLOAD_BYTES = 10 * MB
MAX_UNCOMPRESSED_BYTES = 25 * MB       # contract allows 100 MB; stricter on purpose (1 GB host, 1 worker)
MAX_COMPRESSION_RATIO = 100
# openpyxl keeps the whole shared-string table in memory (twice: values + formulas pass), and nothing
# else bounds it: 9 MB of zip held 1.8 M strings = +226 MB RSS. The real workbook has ~27 KB.
MAX_SHARED_STRINGS_BYTES = 8 * MB
MAX_SHARED_STRINGS = 200_000
RATIO_CHECK_MIN_BYTES = 5 * MB        # tiny files may legitimately expand a lot
MAX_ZIP_ENTRIES = 500
MAX_DATA_ROWS = 5000                  # per sheet
MAX_CONSECUTIVE_BLANK_ROWS = 100
HEADER_SCAN_ROWS = 10
MAX_FILE_NAME_LENGTH = 200

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
ALLOWED_CONTENT_TYPES = {
    XLSX_CONTENT_TYPE,
    'application/octet-stream',      # some browsers / proxies send this for any binary
    'application/zip',
    'application/x-zip-compressed',
}
ZIP_SIGNATURE = b'PK\x03\x04'
_XML_PART_SUFFIXES = ('.xml', '.rels', '.vml')
# Everything the zip/deflate layer can raise while reading a damaged part. ``ImportFileError`` is not a
# subclass of any of these, so the deliberate rejections raised inside the same block keep their message.
_CORRUPT_ARCHIVE_ERRORS = (
    zipfile.BadZipFile, zlib.error, struct.error, EOFError, OSError, RuntimeError, NotImplementedError, ValueError,
)
MAX_CONTENT_TYPES_BYTES = 1 * MB
_CONTENT_TYPES_PART = '[Content_Types].xml'
_SHARED_STRINGS_TYPE_SUFFIX = 'sharedstrings+xml'
_DTD_MARKERS = (b'<!DOCTYPE', b'<!ENTITY')
_UTF16_BOMS = (b'\xff\xfe', b'\xfe\xff')
_SHARED_STRING_RE = re.compile(rb'<(?:\w+:)?si[\s>/]')
_FORMULA_RE = re.compile(rb'<(?:\w+:)?f[\s>/]')

# Canonical sheet names.
SHEET_CATALOG = 'MATRIZ DE RECURSOS'
SHEET_ENTRIES = 'INGRESOS'
SHEET_EXITS = 'EGRESOS'
SHEET_SUMMARY = 'RESUMEN DE MOVIMIENTO'
SHEET_ORDER = (SHEET_CATALOG, SHEET_ENTRIES, SHEET_EXITS, SHEET_SUMMARY)


def normalize_label(value) -> str:
    """
    Accent-, case- and punctuation-insensitive label: ``'Nª VALE'`` -> ``'NA VALE'``,
    ``'USUARIO/ DESTINO'`` -> ``'USUARIO DESTINO'``, ``'INV. INICIAL'`` -> ``'INV INICIAL'``.
    """
    if not isinstance(value, str):
        return ''
    decomposed = unicodedata.normalize('NFKD', value)
    stripped = ''.join(ch for ch in decomposed if not unicodedata.combining(ch))
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', stripped.upper()).split())


@dataclass(frozen=True)
class SheetSpec:
    name: str
    aliases: tuple[str, ...]
    fields: dict
    required: tuple[str, ...]
    inputs: tuple[str, ...]        # columns that decide whether a row is blank


SPECS: dict[str, SheetSpec] = {
    SHEET_CATALOG: SheetSpec(
        name=SHEET_CATALOG,
        aliases=('MATRIZ DE RECURSOS', 'MATRIZ RECURSOS', 'MATRIZ'),
        fields={
            'order_number': ('ORDER NUMBER', 'NUMERO DE ORDEN', 'N DE ORDEN', 'ORDEN'),
            'product_code': ('ITEM', 'CODIGO', 'CODIGO DE PRODUCTO', 'CODIGO ITEM'),
            'description': ('ITEM DESCRIPTION', 'DESCRIPCION', 'DESCRIPCION DEL ITEM'),
            'unit': ('UNIT OF MEASURE', 'UOM', 'UNIDAD', 'UNIDAD DE MEDIDA'),
            'purchase_order_ref': ('OC NA', 'OC N', 'OC', 'ORDEN DE COMPRA', 'N OC'),
            'packing_list': ('SUPPLIER PACKING LIST NUMBER', 'PACKING LIST', 'LISTA DE EMPAQUE'),
            'shipment_quantity': ('SHIPMENT QUANTITY',),
        },
        required=('product_code', 'description', 'unit'),
        inputs=('order_number', 'product_code', 'description', 'unit', 'purchase_order_ref',
                'packing_list', 'shipment_quantity'),
    ),
    SHEET_ENTRIES: SheetSpec(
        name=SHEET_ENTRIES,
        aliases=('INGRESOS', 'INGRESO'),
        fields={
            'date': ('FECHA',),
            'tag': ('TAG',),
            'product_code': ('ITEM', 'CODIGO'),
            'quantity': ('CANTIDAD INGRESADA', 'CANTIDAD'),
            'location': ('LOCACION', 'UBICACION'),
            'po_number': ('PO', 'OC', 'ORDEN DE COMPRA'),
            'guide_number': ('GUIA DE REMISION', 'GUIA'),
            'internal_transfer': ('TRANSFERENCIA INTERNA',),
            'observation': ('OBSERVACION', 'OBSERVACIONES'),
        },
        required=('product_code', 'quantity'),
        inputs=('date', 'tag', 'product_code', 'quantity', 'po_number', 'guide_number',
                'internal_transfer', 'observation'),
    ),
    SHEET_EXITS: SheetSpec(
        name=SHEET_EXITS,
        aliases=('EGRESOS', 'EGRESO'),
        fields={
            'voucher': ('NA VALE', 'N VALE', 'NRO VALE', 'N DE VALE', 'NUMERO DE VALE', 'VALE'),
            'date': ('FECHA',),
            'product_code': ('ITEM', 'CODIGO'),
            'quantity': ('CANTIDAD A EGRESAR', 'CANTIDAD EGRESADA', 'CANTIDAD'),
            'location': ('LOCACION', 'UBICACION'),
            'internal_transfer': ('TRANSFERENCIA INTERNA',),
            'destination': ('USUARIO DESTINO', 'DESTINO', 'USUARIO'),
            'observation': ('OBSERVACION', 'OBSERVACIONES'),
        },
        required=('product_code', 'quantity', 'destination'),
        inputs=('voucher', 'date', 'product_code', 'quantity', 'internal_transfer', 'destination',
                'observation'),
    ),
    SHEET_SUMMARY: SheetSpec(
        name=SHEET_SUMMARY,
        aliases=('RESUMEN DE MOVIMIENTO', 'RESUMEN DE MOVIMIENTOS', 'RESUMEN'),
        fields={
            'product_code': ('ITEM', 'CODIGO'),
            'description': ('DESCRIPCION', 'DESCRIPCION DEL ITEM'),
            'unit': ('UOM', 'UNIDAD'),
            'initial': ('INV INICIAL', 'INVENTARIO INICIAL', 'STOCK INICIAL'),
            'entries': ('INGRESOS',),
            'exits': ('EGRESOS',),
            'stock': ('STOCK ACTUAL', 'STOCK'),
            'location': ('LOCACION', 'UBICACION'),
            'observation': ('OBSERVACION', 'OBSERVACIONES'),
            'project': ('PROYECTO',),
        },
        required=('product_code',),
        inputs=('product_code', 'description', 'unit', 'initial', 'stock', 'location', 'observation',
                'project'),
    ),
}


# ---------------------------------------------------------------------------
# Result types
# ---------------------------------------------------------------------------

@dataclass
class RawRow:
    row: int                       # Excel row number (1-based)
    values: dict                   # field -> raw cell value (strings stripped, '' -> None)


@dataclass
class SheetData:
    name: str                      # canonical name
    title: str                     # title in the workbook
    header_row: int
    columns: dict                  # field -> column index
    rows: list = field(default_factory=list)
    last_row: int = 0


@dataclass
class ParsedWorkbook:
    sheets: dict = field(default_factory=dict)       # canonical name -> SheetData
    problems: list = field(default_factory=list)     # structural problems
    catalog_sheet: str | None = None                 # SHEET_CATALOG or SHEET_SUMMARY


# ---------------------------------------------------------------------------
# Upload validation
# ---------------------------------------------------------------------------

def sanitize_file_name(name) -> str:
    """Base name only, no control characters, bounded length (the extension survives the cut)."""
    base = os.path.basename(str(name or '').replace('\\', '/'))
    base = ''.join(ch for ch in base if ch.isprintable()).strip()
    if len(base) > MAX_FILE_NAME_LENGTH:
        stem, extension = os.path.splitext(base)
        extension = extension[:16]
        base = stem[:MAX_FILE_NAME_LENGTH - len(extension)] + extension
    return base or 'libro.xlsx'


def _invalid_file(message='El archivo no es un libro de Excel (.xlsx) válido.') -> ImportFileError:
    return ImportFileError('ARCHIVO_INVALIDO', message)


def inspect_upload(uploaded_file) -> tuple[bytes, str]:
    """
    Validate an uploaded file before parsing and return ``(bytes, safe_name)``.
    Raises ``ImportFileError`` (400). Nothing is stored here.
    """
    name = sanitize_file_name(getattr(uploaded_file, 'name', ''))
    if not name.lower().endswith('.xlsx'):
        raise ImportFileError(
            'extension_not_allowed',
            'Solo se aceptan libros de Excel con extensión .xlsx (no se permiten .xlsm ni .xls).',
        )

    content_type = (getattr(uploaded_file, 'content_type', '') or '').split(';')[0].strip().lower()
    if content_type and content_type not in ALLOWED_CONTENT_TYPES:
        raise _invalid_file()

    limit_mb = MAX_UPLOAD_BYTES // MB
    too_big = ImportFileError('file_too_large', f'El archivo supera el máximo de {limit_mb} MB.')
    size = getattr(uploaded_file, 'size', None)
    if size and size > MAX_UPLOAD_BYTES:
        raise too_big
    uploaded_file.seek(0)
    data = uploaded_file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise too_big
    if not data:
        raise _invalid_file('El archivo está vacío.')
    if not data.startswith(ZIP_SIGNATURE):
        raise _invalid_file()

    inspect_zip(data)
    return data, name


def _contains_dtd(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> bool:
    """True when an XML part declares a DTD/entities (XXE, billion laughs) or is UTF-16."""
    carry = b''
    first = True
    with zf.open(info) as part:
        while True:
            chunk = part.read(64 * 1024)
            if not chunk:
                return False
            if first:
                first = False
                if chunk.startswith(_UTF16_BOMS):
                    return True
            window = carry + chunk
            if any(marker in window for marker in _DTD_MARKERS):
                return True
            carry = window[-16:]


def _declared_content_types(zf: zipfile.ZipFile) -> tuple[dict, dict]:
    """
    ``(overrides, defaults)`` of ``[Content_Types].xml``: part name (no leading slash) -> content
    type, and extension -> content type, all lower-case. openpyxl finds the shared-string table and
    the sheets through THIS manifest, not through file names, so the limits below must too.
    """
    info = zf.getinfo(_CONTENT_TYPES_PART)
    if info.file_size > MAX_CONTENT_TYPES_BYTES:
        raise _invalid_file()
    if _contains_dtd(zf, info):                 # the stdlib parser is only safe on a DTD-free document
        raise _invalid_file('El libro contiene declaraciones XML no permitidas.')
    try:
        root = ElementTree.fromstring(zf.read(info))
    except ElementTree.ParseError:
        raise _invalid_file() from None
    overrides: dict = {}
    defaults: dict = {}
    for element in root:
        kind = element.tag.rsplit('}', 1)[-1]
        content_type = (element.get('ContentType') or '').strip().lower()
        if kind == 'Override':
            overrides[(element.get('PartName') or '').strip().lstrip('/').lower()] = content_type
        elif kind == 'Default':
            defaults[(element.get('Extension') or '').strip().lstrip('.').lower()] = content_type
    return overrides, defaults


def _content_type_of(name: str, overrides: dict, defaults: dict) -> str:
    lowered = name.lower()
    if lowered in overrides:
        return overrides[lowered]
    leaf = lowered.rsplit('/', 1)[-1]
    return defaults.get(leaf.rsplit('.', 1)[-1], '') if '.' in leaf else ''


def _is_xml_content_type(content_type: str) -> bool:
    """XML parts of the package, whatever they are called (SVG images and the like are not scanned)."""
    return content_type in ('application/xml', 'text/xml') or (
        content_type.startswith('application/') and content_type.endswith('+xml')
    )


def _check_shared_strings(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> None:
    """Bound the shared-string table (bytes and number of entries). Raises ``ImportFileError``."""
    too_much_text = _invalid_file('El libro contiene demasiado texto para procesarlo.')
    if info.file_size > MAX_SHARED_STRINGS_BYTES:
        raise too_much_text
    with zf.open(info) as part:
        # ZipExtFile never yields more than the declared size, so this read is bounded too.
        content = part.read(MAX_SHARED_STRINGS_BYTES + 1)
    if len(content) > MAX_SHARED_STRINGS_BYTES:
        raise too_much_text
    # Count the real <si> elements: the ``count``/``uniqueCount`` attributes are attacker-controlled.
    if sum(1 for _ in _SHARED_STRING_RE.finditer(content)) > MAX_SHARED_STRINGS:
        raise too_much_text


def inspect_zip(data: bytes) -> None:
    """Structural checks of the .xlsx container. Raises ``ImportFileError``."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except _CORRUPT_ARCHIVE_ERRORS:
        raise _invalid_file() from None

    with zf:
        infos = zf.infolist()
        if not infos or len(infos) > MAX_ZIP_ENTRIES:
            raise _invalid_file()
        names = [i.filename for i in infos]
        if len(set(names)) != len(names):
            raise _invalid_file()
        for name in names:
            if name.startswith('/') or '\\' in name or '..' in name.split('/'):
                raise _invalid_file()
        if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise _invalid_file()
        if any(n.lower().endswith('vbaproject.bin') for n in names):
            raise _invalid_file('El libro contiene macros; no se permite.')

        total = sum(i.file_size for i in infos)
        too_large = _invalid_file('El libro es sospechosamente grande al descomprimirse.')
        if total > MAX_UNCOMPRESSED_BYTES:
            raise too_large
        if total > RATIO_CHECK_MIN_BYTES and total > max(len(data), 1) * MAX_COMPRESSION_RATIO:
            raise too_large

        try:
            overrides, defaults = _declared_content_types(zf)
            for info in infos:
                lowered = info.filename.lower()
                content_type = _content_type_of(info.filename, overrides, defaults)
                # By name AND by declared content type: a part renamed to ``.dat`` or ``tablaDeCadenas.xml``
                # is still read by openpyxl as a sheet / the shared-string table.
                if (lowered.endswith(_XML_PART_SUFFIXES) or _is_xml_content_type(content_type)) \
                        and _contains_dtd(zf, info):
                    raise _invalid_file('El libro contiene declaraciones XML no permitidas.')
                if lowered.endswith('sharedstrings.xml') or content_type.endswith(_SHARED_STRINGS_TYPE_SUFFIX):
                    _check_shared_strings(zf, info)
        except _CORRUPT_ARCHIVE_ERRORS:
            # A truncated or damaged download is a normal case: the deflate stream raises
            # ``zlib.error`` (not an OSError) and a bad header can raise ``struct.error``/``ValueError``.
            raise _invalid_file() from None


# ---------------------------------------------------------------------------
# Workbook parsing
# ---------------------------------------------------------------------------

def _cell(row, col):
    value = row[col] if col < len(row) else None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _is_formula(value) -> bool:
    if isinstance(value, str):
        return value.startswith('=')
    return type(value).__name__ in ('ArrayFormula', 'DataTableFormula')


def _find_header(head_rows: list, spec: SheetSpec) -> tuple[int, dict, list]:
    """``(header_row, field -> column, missing required fields)`` from the first rows."""
    best = None
    for index, row in enumerate(head_rows, start=1):
        mapping: dict = {}
        for col, cell in enumerate(row):
            label = normalize_label(cell)
            if not label:
                continue
            for field_name, aliases in spec.fields.items():
                if field_name not in mapping and label in aliases:
                    mapping[field_name] = col
                    break
        missing = [f for f in spec.required if f not in mapping]
        if not missing:
            return index, mapping, []
        if best is None or len(mapping) > len(best[1]):
            best = (index, mapping, missing)
    if best is None or not best[1]:
        return 1, {}, list(spec.required)
    return best


def _read_sheet(ws, spec: SheetSpec, canonical: str, title: str, problems: list):
    head_rows = list(ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True))
    header_row, mapping, missing = _find_header(head_rows, spec)
    if missing:
        problems.append(make_problem(
            'CABECERA_FALTANTE', sheet=canonical, columns=', '.join(spec.fields[f][0] for f in missing),
        ))
        return None

    sheet = SheetData(name=canonical, title=title, header_row=header_row, columns=mapping, last_row=header_row)
    max_col = max(mapping.values()) + 1
    input_fields = [f for f in spec.inputs if f in mapping]
    blank_run = 0
    data_rows = ws.iter_rows(min_row=header_row + 1, max_col=max_col, values_only=True)
    for offset, raw in enumerate(data_rows):
        row_number = header_row + 1 + offset
        values = {f: _cell(raw, col) for f, col in mapping.items()}
        if all(values[f] is None for f in input_fields):
            blank_run += 1
            if blank_run >= MAX_CONSECUTIVE_BLANK_ROWS:
                break
            continue
        blank_run = 0
        if len(sheet.rows) >= MAX_DATA_ROWS:
            raise ImportFileError(
                'too_many_rows',
                f'La hoja {canonical} supera el máximo de {MAX_DATA_ROWS} filas de datos por archivo.',
            )
        sheet.rows.append(RawRow(row_number, values))
        sheet.last_row = row_number
    return sheet


def _count_uncached_formulas(data_ws, formula_ws, sheet: SheetData) -> tuple[int, int]:
    """``(formula cells, formula cells without cached value)`` in the columns we read."""
    columns = sorted(sheet.columns.values())
    max_col = columns[-1] + 1
    total = uncached = 0
    pairs = zip(
        data_ws.iter_rows(min_row=1, max_row=sheet.last_row, max_col=max_col, values_only=True),
        formula_ws.iter_rows(min_row=1, max_row=sheet.last_row, max_col=max_col, values_only=True),
    )
    for values_row, formulas_row in pairs:
        for col in columns:
            formula = formulas_row[col] if col < len(formulas_row) else None
            if not _is_formula(formula):
                continue
            total += 1
            cached = values_row[col] if col < len(values_row) else None
            if cached is None:
                uncached += 1
    return total, uncached


def parse_workbook(data: bytes) -> ParsedWorkbook:
    """
    Parse the workbook into raw rows. Raises ``ImportFileError`` when the file is
    unreadable or exceeds the row limit; structural gaps (missing sheet/header) are
    returned as blocker problems so the user sees them in the preview.
    """
    try:
        return _parse(data)
    except ImportFileError:
        raise
    except Exception as exc:  # openpyxl/lxml/zip raise many types on hostile input
        # Only the exception TYPE is logged: messages can quote cell content.
        logger.warning('warehouse.import.parse_failed', extra={'error_type': type(exc).__name__})
        raise _invalid_file() from None


def _has_formulas(data: bytes) -> bool:
    """True when any worksheet part holds a formula element (streamed, never fully in memory)."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            name = info.filename.lower()
            if not (name.startswith('xl/worksheets/') and name.endswith('.xml')):
                continue
            carry = b''
            with zf.open(info) as part:
                while True:
                    chunk = part.read(64 * 1024)
                    if not chunk:
                        break
                    if _FORMULA_RE.search(carry + chunk):
                        return True
                    carry = chunk[-8:]
    return False


def _parse(data: bytes) -> ParsedWorkbook:
    parsed = ParsedWorkbook()
    wb_values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    wb_formulas = None
    try:
        by_name = {}
        for ws in wb_values.worksheets:
            by_name.setdefault(normalize_label(ws.title), ws)

        found = {}
        for canonical in SHEET_ORDER:
            spec = SPECS[canonical]
            ws = next((by_name[a] for a in spec.aliases if a in by_name), None)
            if ws is None:
                continue
            sheet = _read_sheet(ws, spec, canonical, ws.title, parsed.problems)
            if sheet is not None:
                parsed.sheets[canonical] = sheet
                found[canonical] = ws

        _choose_catalog_source(parsed)

        # Formulas whose result was never calculated (workbook produced by a script). The second
        # load (with formulas) doubles the memory of the pass, so it only happens when the sheets
        # actually contain formulas.
        if parsed.sheets and _has_formulas(data):
            wb_formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            formula_sheets = {normalize_label(ws.title): ws for ws in wb_formulas.worksheets}
            for canonical, sheet in parsed.sheets.items():
                formula_ws = formula_sheets.get(normalize_label(sheet.title))
                if formula_ws is None:
                    continue
                total, uncached = _count_uncached_formulas(found[canonical], formula_ws, sheet)
                if total and uncached / total >= 0.5:
                    parsed.problems.append(make_problem(
                        'LIBRO_SIN_VALORES_CALCULADOS', sheet=canonical, count=uncached,
                    ))
    finally:
        wb_values.close()
        if wb_formulas is not None:
            wb_formulas.close()

    catalog = parsed.sheets.get(SHEET_CATALOG)
    if catalog and 'shipment_quantity' in catalog.columns and any(
        r.values.get('shipment_quantity') is not None for r in catalog.rows
    ):
        parsed.problems.append(make_problem('SHIPMENT_QUANTITY_IGNORADA', sheet=SHEET_CATALOG))
    return parsed


def _choose_catalog_source(parsed: ParsedWorkbook) -> None:
    """MATRIZ DE RECURSOS is the catalog; a SYSPCC export has none, so its summary serves."""
    if SHEET_CATALOG in parsed.sheets:
        parsed.catalog_sheet = SHEET_CATALOG
        return
    summary = parsed.sheets.get(SHEET_SUMMARY)
    if summary and 'description' in summary.columns and 'unit' in summary.columns:
        parsed.catalog_sheet = SHEET_SUMMARY
        parsed.problems.append(make_problem('CATALOGO_DESDE_RESUMEN', sheet=SHEET_SUMMARY))
        return
    # A MATRIZ with bad headers already produced CABECERA_FALTANTE; otherwise the sheet is absent.
    if not any(p.code == 'CABECERA_FALTANTE' and p.sheet == SHEET_CATALOG for p in parsed.problems):
        parsed.problems.append(make_problem('HOJA_FALTANTE', sheet=SHEET_CATALOG))
