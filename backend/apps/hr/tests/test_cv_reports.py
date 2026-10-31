"""PDF report, ZIP streaming, Excel export, deletion and orphan purge (SYSPCC-023)."""

import io
import zipfile
from pathlib import Path

import pytest
from django.core.signals import request_finished
from django.db import close_old_connections
from openpyxl import load_workbook

from apps.hr.enums import CVItemStatus
from apps.hr.models import CVAnalysis, CVAnalysisItem
from apps.hr.services import cv_analysis_service as svc
from apps.hr.services import cv_export, cv_report_pdf, cv_service
from apps.hr.tests.cv_helpers import FakeCVClient, default_report, make_pdf, upload

pytestmark = pytest.mark.django_db

BASE = '/api/v1/hr/cv-analyses/'
EVIL = '=cmd|\'/c calc\'!A0'


def evil_report(kw, **extra):
    return default_report(kw['output_format'], name=EVIL, headline='+1+1', summary='<img src="file:///etc/passwd"> <b>x</b>',
                          skills=['@SUM(1,1)', '-2+3'], **extra)


@pytest.fixture
def hostile(cv_env, hr_manager, monkeypatch, settings):
    """Analysis with one DONE item whose AI-extracted text is hostile."""
    import anthropic
    client = FakeCVClient()
    client.script = [lambda kw: client.response(evil_report(kw))]
    monkeypatch.setattr(anthropic, 'Anthropic', lambda **kw: client)
    analysis = svc.create_analysis(user=hr_manager, name='Análisis =SUM(A1)')
    res = svc.upload_files(analysis=analysis, user=hr_manager, files=[upload('=evil.pdf', make_pdf('h'))])
    pk = res['accepted'][0]['id']
    cv_service.run_item(pk)
    return analysis, CVAnalysisItem.objects.get(pk=pk)


class TestPdfReport:
    def test_valid_pdf_with_candidate_name(self, cv_env, hr_manager, fake_cv_client):
        analysis = svc.create_analysis(user=hr_manager, name='P', requirements_text='AutoCAD (obligatorio)')
        pk = svc.upload_files(analysis=analysis, user=hr_manager, files=[upload('a.pdf', make_pdf())])['accepted'][0]['id']
        cv_service.run_item(pk)
        item = CVAnalysisItem.objects.get(pk=pk)
        data = cv_report_pdf.build(item)
        assert data.startswith(b'%PDF')
        from pypdf import PdfReader
        text = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        assert 'Ana Pérez Gómez' in text or 'Ana P' in text
        assert 'La IA ayuda a ordenar; la decisión final es de RR. HH.' in text
        assert 'Evaluación frente al puesto' in text and '100 / 100' in text and 'Avanzar' in text

    def test_markup_in_cv_text_is_printed_not_interpreted(self, hostile):
        _, item = hostile
        data = cv_report_pdf.build(item)
        from pypdf import PdfReader
        text = ''.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        assert '<img src="file:///etc/passwd">' in text and '<b>x</b>' in text  # literal text
        assert b'/Subtype /Image' not in data and b'/Subtype/Image' not in data

    def test_esc_neutralizes_markup_and_non_cp1252(self):
        assert cv_report_pdf.esc('<img src="http://x">&') == '&lt;img src="http://x"&gt;&amp;'
        assert cv_report_pdf.esc('Ωmega ✓ 日本') == '?mega ? ??'
        assert cv_report_pdf.esc(None) == '' and cv_report_pdf.esc('a\x00b\x07') == 'ab'

    def test_exotic_characters_do_not_break_build(self, hostile):
        _, item = hostile
        item.result['report']['summary'] = 'Ωmega ✓ 日本 </para> <font size="99999">'
        item.result['report']['candidate']['full_name'] = 'Ñandú ✓'
        assert cv_report_pdf.build(item).startswith(b'%PDF')

    def test_report_endpoint_generates_when_missing_and_requires_done(self, hostile, hr_client):
        analysis, item = hostile
        item.report_pdf.delete(save=True)
        resp = hr_client.get(f'{BASE}{analysis.pk}/items/{item.pk}/report/')
        assert resp.status_code == 200 and resp['Content-Type'] == 'application/pdf'
        assert b''.join(resp.streaming_content).startswith(b'%PDF')
        assert resp['Cache-Control'] == 'no-store' and 'attachment' in resp['Content-Disposition']
        assert CVAnalysisItem.objects.get(pk=item.pk).report_pdf  # stored for next time
        CVAnalysisItem.objects.filter(pk=item.pk).update(status=CVItemStatus.FAILED)
        bad = hr_client.get(f'{BASE}{analysis.pk}/items/{item.pk}/report/')
        assert bad.status_code == 409 and bad.data['code'] == 'invalid_state'

    def test_format_query_param_is_not_a_renderer_override(self, hostile, hr_client):
        analysis, item = hostile
        resp = hr_client.get(f'{BASE}{analysis.pk}/items/{item.pk}/report/?format=json')
        assert resp.status_code == 200 and resp['Content-Type'] == 'application/pdf'

    def test_original_file_download_and_missing(self, hostile, hr_client):
        analysis, item = hostile
        resp = hr_client.get(f'{BASE}{analysis.pk}/items/{item.pk}/file/')
        assert resp.status_code == 200 and resp['X-Content-Type-Options'] == 'nosniff'
        assert resp['Content-Disposition'].startswith('attachment') and 'evil' not in resp['Content-Disposition']
        # resp.close() releases the file but also fires request_finished, whose
        # close_old_connections would drop the PostgreSQL connection mid-test.
        request_finished.disconnect(close_old_connections)
        try:
            resp.close()
        finally:
            request_finished.connect(close_old_connections)
        item.file.storage.delete(item.file.name)
        gone = hr_client.get(f'{BASE}{analysis.pk}/items/{item.pk}/file/')
        assert gone.status_code == 404 and gone.data['code'] == 'file_missing'


class TestExcel:
    def test_formulas_are_neutralized_and_numbers_stay_numbers(self, hostile, hr_client):
        analysis, item = hostile
        resp = hr_client.get(f'{BASE}{analysis.pk}/download-xlsx/')
        assert resp.status_code == 200 and 'spreadsheetml' in resp['Content-Type']
        wb = load_workbook(io.BytesIO(b''.join(resp.streaming_content)))
        assert wb.sheetnames == ['Candidatos', 'Con error']
        ws = wb['Candidatos']
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        assert row['Nombre'] == "'" + EVIL and row['Titular'] == "'+1+1" and row['Habilidades'].startswith("'@SUM")
        assert row['Archivo original'].startswith("'=")
        for cell in ws[2]:
            assert cell.data_type != 'f', cell.value
        assert isinstance(row['Años exp.'], (int, float)) and ws.freeze_panes == 'A2'
        assert ws.auto_filter.ref == 'A1:O2'

    def test_safe_text_unit(self):
        for raw in ('=1+1', '+1', '-1', '@x', '\t=1', ' =1', '\r=1'):
            assert cv_export.safe_text(raw).startswith("'"), raw
        assert cv_export.safe_text('Juan') == 'Juan' and cv_export.safe_text(None) == ''
        assert len(cv_export.safe_text('x' * 40000)) == cv_export.MAX_CELL_CHARS
        assert cv_export.safe_text('a\x00b') == 'ab'

    def test_requirements_matrix_and_error_sheet(self, cv_env, hr_manager, fake_cv_client):
        analysis = svc.create_analysis(user=hr_manager, name='P', requirements_text='AutoCAD (obligatorio)\nInglés')
        files = [upload('a.pdf', make_pdf('a')), upload('b.pdf', make_pdf('b'))]
        ids = [a['id'] for a in svc.upload_files(analysis=analysis, user=hr_manager, files=files)['accepted']]
        cv_service.run_item(ids[0])
        CVAnalysisItem.objects.filter(pk=ids[1]).update(
            status='FAILED', error_message='=cmd Reintenta', original_name='roto.pdf')
        wb = load_workbook(io.BytesIO(cv_export.build_xlsx(analysis)))
        assert wb.sheetnames == ['Candidatos', 'Requisitos', 'Con error']
        req = wb['Requisitos']
        assert [c.value for c in req[1]][1:] == ['1. AutoCAD (obligatorio)', '2. Inglés']
        assert [c.value for c in req[2]][1:] == ['Cumple', 'Cumple']
        err = wb['Con error']
        assert err['A2'].value == 'roto.pdf' and err['B2'].value == "'=cmd Reintenta"
        main = wb['Candidatos']
        assert main['K2'].value == 100 and main['A2'].value == 1 and main['L2'].value == 'Avanzar'

    def test_nothing_to_download_409(self, cv_env, hr_client):
        a = hr_client.post(BASE, {'name': 'Vacío total'}, format='json').data['id']
        for kind in ('download-zip', 'download-xlsx'):
            resp = hr_client.get(f'{BASE}{a}/{kind}/')
            assert resp.status_code == 409 and resp.data['code'] == 'nothing_to_download'

    def test_ranks_share_ties(self):
        class I:
            def __init__(self, pk, score):
                self.pk, self.score = pk, score
        ranks = cv_export.compute_ranks([I(1, 90), I(2, 90), I(3, 70), I(4, None)])
        assert ranks == {1: 1, 2: 1, 3: 3}


class TestZip:
    def make_many(self, hr_manager, n=3):
        analysis = svc.create_analysis(user=hr_manager, name='Lote grande')
        files = [upload(f'{i}.pdf', make_pdf(str(i))) for i in range(n)]
        ids = [a['id'] for a in svc.upload_files(analysis=analysis, user=hr_manager, files=files)['accepted']]
        for pk in ids:
            cv_service.run_item(pk)
        return analysis, ids

    def test_zip_is_streamed_valid_and_names_are_safe(self, cv_env, hr_manager, fake_cv_client, hr_client):
        analysis, ids = self.make_many(hr_manager)
        CVAnalysisItem.objects.filter(pk=ids[2]).update(status='FAILED')  # only DONE items go in
        resp = hr_client.get(f'{BASE}{analysis.pk}/download-zip/')
        assert resp.streaming and resp['Content-Type'] == 'application/zip'
        assert 'attachment' in resp['Content-Disposition'] and 'CVs_Lote-grande_' in resp['Content-Disposition']
        chunks = list(resp.streaming_content)
        assert len(chunks) >= 2  # streamed in pieces, not a single buffer
        zf = zipfile.ZipFile(io.BytesIO(b''.join(chunks)))
        assert zf.testzip() is None
        names = zf.namelist()
        assert names[0] == 'Resumen.xlsx' and len(names) == 3 and len(set(names)) == 3
        for n in names:
            assert '..' not in n and '/' not in n and '\\' not in n
        pdfs = [n for n in names if n.endswith('.pdf')]
        assert all(n.split('_')[0].isdigit() and zf.read(n).startswith(b'%PDF') for n in pdfs)

    def test_stream_memory_is_chunked_not_total(self, cv_env, hr_manager, fake_cv_client):
        analysis, ids = self.make_many(hr_manager, 2)
        gen = cv_export.stream_zip(analysis, ids, include_summary=False)
        assert hasattr(gen, '__next__')  # generator: nothing built until consumed
        first = next(gen)
        assert isinstance(first, bytes) and len(first) <= cv_export.ZIP_CHUNK + 1024
        list(gen)

    def test_item_deleted_while_streaming_is_skipped(self, cv_env, hr_manager, fake_cv_client):
        analysis, ids = self.make_many(hr_manager, 2)
        gen = cv_export.stream_zip(analysis, ids, include_summary=False)
        CVAnalysisItem.objects.filter(pk=ids[1]).delete()
        zf = zipfile.ZipFile(io.BytesIO(b''.join(gen)))
        assert len(zf.namelist()) == 1

    def test_unique_names_for_same_candidate(self, cv_env, hr_manager, fake_cv_client):
        analysis, ids = self.make_many(hr_manager, 2)  # both are "Ana Pérez Gómez"
        zf = zipfile.ZipFile(io.BytesIO(b''.join(cv_export.stream_zip(analysis, ids, include_summary=False))))
        assert len(set(zf.namelist())) == 2
        assert all(n.startswith(('001_Ana-Perez-Gomez_', '002_Ana-Perez-Gomez_')) for n in zf.namelist())


class TestDeletion:
    def stored_files(self, root):
        return {p for p in Path(root).rglob('*') if p.is_file()}

    def test_delete_analysis_removes_rows_and_all_files(self, hostile, hr_client, settings,
                                                         django_capture_on_commit_callbacks):
        analysis, item = hostile
        assert len(self.stored_files(settings.HR_PRIVATE_ROOT)) == 2  # original + report
        with django_capture_on_commit_callbacks(execute=True):
            resp = hr_client.delete(f'{BASE}{analysis.pk}/')
        assert resp.status_code == 204
        assert not CVAnalysis.objects.exists() and not CVAnalysisItem.objects.exists()
        assert self.stored_files(settings.HR_PRIVATE_ROOT) == set()

    def test_delete_single_item_removes_its_files(self, hostile, hr_client, settings, django_capture_on_commit_callbacks):
        analysis, item = hostile
        with django_capture_on_commit_callbacks(execute=True):
            assert hr_client.delete(f'{BASE}{analysis.pk}/items/{item.pk}/').status_code == 204
        assert self.stored_files(settings.HR_PRIVATE_ROOT) == set() and CVAnalysis.objects.exists()

    def test_purge_orphans_command(self, hostile, settings):
        from django.core.management import call_command
        orphan = Path(settings.HR_PRIVATE_ROOT) / 'hr' / 'cv' / 'huerfano.pdf'
        orphan.write_bytes(b'%PDF-orphan')
        out = io.StringIO()
        call_command('hr_cv_purge_orphans', '--dry-run', stdout=out)
        assert orphan.exists() and 'huérfanos: 1' in out.getvalue()
        call_command('hr_cv_purge_orphans', stdout=io.StringIO())
        assert not orphan.exists() and len(self.stored_files(settings.HR_PRIVATE_ROOT)) == 2  # referenced stay

    def test_retention_purges_old_analyses_only_when_enabled(self, hostile, settings, django_capture_on_commit_callbacks):
        from django.core.management import call_command
        from datetime import timedelta
        from django.utils import timezone
        analysis, _ = hostile
        CVAnalysis.objects.filter(pk=analysis.pk).update(created_at=timezone.now() - timedelta(days=100))
        call_command('hr_cv_purge_orphans', stdout=io.StringIO())
        assert CVAnalysis.objects.exists()  # retention off by default
        settings.HR_CV_RETENTION_DAYS = 30
        with django_capture_on_commit_callbacks(execute=True):
            call_command('hr_cv_purge_orphans', stdout=io.StringIO())
        assert not CVAnalysis.objects.exists()
        assert self.stored_files(settings.HR_PRIVATE_ROOT) == set()

    def test_admin_exposes_no_pii(self):
        from django.contrib import admin
        for model in (CVAnalysisItem, CVAnalysis):
            ma = admin.site._registry[model]
            assert not (ma.has_add_permission(None) or ma.has_change_permission(None) or ma.has_delete_permission(None))
            assert 'result' not in ma.fields and 'candidate_name' not in ma.fields
            assert 'original_name' not in ma.fields and 'file' not in ma.fields
