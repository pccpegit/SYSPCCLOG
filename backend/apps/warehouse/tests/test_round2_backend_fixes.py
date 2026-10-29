"""
SYSPCC-025 · django-backend ronda 2 — pruebas propias de los arreglos (más allá de los tests "DEFECTO QA-n").

Todo es inventado: ningún nombre ni valor sale del libro real.

* QA-33: un flujo deflate dañado (y cualquier error de la capa zip) es ARCHIVO_INVALIDO.
* QA-18/19: la fecha que se muestra es la operativa (``movement_date``) y ``Sin fecha`` si es NULL.
* QA-21: una sola regla para la carga inicial: en INGRESOS salvo que INV. INICIAL ya la lleve.
* QA-22/25/26/27/29...: bordes exactos de los topes nuevos.
"""
import io
import struct
import zipfile
import zlib
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

import openpyxl
import pytest
from django.db import DatabaseError
from django.utils import timezone
from pypdf import PdfReader

from apps.warehouse.models import InventoryImport, InventoryImportStatus, InventoryMovement
from apps.warehouse.models.inventory import MAX_STOCK_QUANTITY
from apps.warehouse.services import kardex
from apps.warehouse.services import movements as svc
from apps.warehouse.services.dates import NO_DATE_LABEL, operational_date_label
from apps.warehouse.services.imports import parser, service
from apps.warehouse.services.imports.errors import ImportFileError

from .import_workbooks import (
    D1, WORKER_A, build_workbook, catalog_row, entry_row, standard_workbook,
)
from .qa_support import IMPORTS, INVENTORY, MOVEMENTS, confirm, snapshot_database, tolerant_client, upload

pytestmark = pytest.mark.django_db

SUMMARY = INVENTORY + 'movement-summary/'


@pytest.fixture(autouse=True)
def _quiet_side_effects(monkeypatch):
    from rest_framework.throttling import SimpleRateThrottle

    monkeypatch.setattr(
        SimpleRateThrottle, 'THROTTLE_RATES',
        {**SimpleRateThrottle.THROTTLE_RATES, 'user': '1000000/min', 'warehouse_import': '1000000/hour'},
    )
    with mock.patch('apps.warehouse.views._enqueue_group_pdf_upload'):
        yield


@pytest.fixture
def client(central_warehouse):
    return tolerant_client(central_warehouse)


@pytest.fixture
def item(make_item):
    return make_item('EPP000001', 'GUANTE T: L', 'EA', stock=0)


def _entry(item, qty, user, when, source='PURCHASE', undated=False):
    return svc.register_entry_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        source_type=source, supplier_name='', invoice_number='', notes='', registered_by=user,
        movement_date=when, undated=undated)


def _exit(item, qty, user, when):
    return svc.register_exit_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse='CENTRAL', project_id=None,
        destination_type='EMPLOYEE', destination_detail='Persona Prueba', requested_by_id=None,
        authorized_by_id=None, notes='', registered_by=user, movement_date=when)


def _sheet_sum(book, name, column):
    return sum(row[column] or 0 for row in book[name].iter_rows(min_row=2, values_only=True))


def _load(response):
    assert response.status_code == 200, response.content[:200]
    return openpyxl.load_workbook(io.BytesIO(response.content))


# ---------------------------------------------------------------------------
# QA-33: errores de la capa zip
# ---------------------------------------------------------------------------

class TestZipLayerErrorsAreARejectedFile:
    @pytest.mark.parametrize('error', [
        zlib.error('Error -3 while decompressing data'), struct.error('unpack requires a buffer'),
        ValueError('invalid literal'), EOFError(), OSError('disk'), RuntimeError('bad password'),
    ])
    def test_any_error_reading_a_part_is_archivo_invalido(self, error):
        with mock.patch.object(parser, '_contains_dtd', side_effect=error):
            with pytest.raises(ImportFileError) as exc:
                parser.inspect_zip(standard_workbook())
        assert exc.value.code == 'ARCHIVO_INVALIDO'

    def test_a_deliberate_rejection_keeps_its_own_message(self):
        """``ImportFileError`` is not one of the translated errors: the DTD message must not be flattened."""
        with mock.patch.object(parser, '_contains_dtd', return_value=True):
            with pytest.raises(ImportFileError) as exc:
                parser.inspect_zip(standard_workbook())
        assert 'XML' in exc.value.message

    def test_a_manifest_larger_than_the_limit_is_refused(self):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(standard_workbook())) as src, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                content = src.read(info.filename)
                if info.filename == '[Content_Types].xml':
                    padding = b'<!--' + b'x' * (parser.MAX_CONTENT_TYPES_BYTES + 10) + b'-->'
                    content = content.replace(b'</Types>', padding + b'</Types>')
                dst.writestr(info.filename, content)
        with pytest.raises(ImportFileError) as exc:
            parser.inspect_zip(out.getvalue())
        assert exc.value.code == 'ARCHIVO_INVALIDO'

    def test_a_manifest_that_is_not_xml_is_refused(self):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(standard_workbook())) as src, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                dst.writestr(info.filename, b'<Types><Override' if info.filename == '[Content_Types].xml' else src.read(info.filename))
        with pytest.raises(ImportFileError) as exc:
            parser.inspect_zip(out.getvalue())
        assert exc.value.code == 'ARCHIVO_INVALIDO'

    def test_a_standard_workbook_still_passes(self):
        parser.inspect_zip(standard_workbook())


class TestFileNameKeepsItsExtension:
    def test_a_long_name_is_cut_in_the_stem(self):
        name = parser.sanitize_file_name('inventario_' + 'x' * 400 + '.xlsx')
        assert len(name) == parser.MAX_FILE_NAME_LENGTH and name.endswith('.xlsx')

    def test_a_long_name_without_extension_is_just_cut(self):
        assert len(parser.sanitize_file_name('y' * 400)) == parser.MAX_FILE_NAME_LENGTH

    def test_a_short_name_is_untouched(self):
        assert parser.sanitize_file_name('libro.xlsx') == 'libro.xlsx'

    def test_an_absurdly_long_extension_cannot_eat_the_whole_name(self):
        name = parser.sanitize_file_name('a.' + 'z' * 400)
        assert len(name) <= parser.MAX_FILE_NAME_LENGTH and name.startswith('a.')


# ---------------------------------------------------------------------------
# QA-18 / QA-19: fecha operativa
# ---------------------------------------------------------------------------

class TestOperationalDateLabel:
    def test_a_date_is_day_month_year(self):
        assert operational_date_label(date(2026, 10, 5)) == '05/10/2026'

    def test_a_missing_date_says_so(self):
        assert operational_date_label(None) == NO_DATE_LABEL == 'Sin fecha'


class TestUndatedMovementsInDocuments:
    @pytest.fixture
    def undated(self, item, central_warehouse):
        _entry(item, 5, central_warehouse, None, undated=True)
        movement = InventoryMovement.objects.get()
        assert movement.movement_date is None
        return movement

    def test_the_plain_export_says_sin_fecha_instead_of_the_registration_time(self, client, undated):
        sheet = _load(client.get(MOVEMENTS + 'export/')).active
        header = [c.value for c in sheet[1]]
        assert sheet[2][header.index('Fecha')].value == 'Sin fecha'

    def test_the_movement_pdf_says_sin_fecha(self, client, undated):
        response = client.get(f'{MOVEMENTS}{undated.pk}/pdf/')
        assert response.status_code == 200
        text = ' '.join(' '.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(response.content)).pages).split())
        assert 'Sin fecha' in text
        assert undated.created_at.strftime('%d/%m/%Y %H:%M') not in text

    def test_the_plain_export_shows_the_operational_date_without_time(self, client, item, central_warehouse):
        when = timezone.localdate() - timedelta(days=9)
        _entry(item, 5, central_warehouse, when)
        sheet = _load(client.get(MOVEMENTS + 'export/')).active
        header = [c.value for c in sheet[1]]
        assert sheet[2][header.index('Fecha')].value == when.strftime('%d/%m/%Y')


# ---------------------------------------------------------------------------
# QA-21: una sola regla para la carga inicial
# ---------------------------------------------------------------------------

LOAD_DAY, ENTRY_DAY, EXIT_DAY = date(2026, 10, 5), date(2026, 10, 8), date(2026, 10, 12)


@pytest.fixture
def history(item, central_warehouse):
    _entry(item, 10, central_warehouse, LOAD_DAY, source='INITIAL_LOAD')
    _entry(item, 4, central_warehouse, ENTRY_DAY)
    _exit(item, 3, central_warehouse, EXIT_DAY)
    return item


def _summary_row(client, **params):
    return client.get(SUMMARY, {'no_page': 'true', **params}).json()['results'][0]


class TestInitialLoadIsWrittenWhereTheSummaryCountsIt:
    def test_the_rule_depends_only_on_date_from(self):
        assert kardex.initial_loads_count_as_entries(None) is False
        assert kardex.initial_loads_count_as_entries(date(2026, 10, 5)) is True

    @pytest.mark.parametrize('params, initial, entries', [
        ({}, 10, 4),                                                     # travels in INV. INICIAL
        ({'date_to': '2026-10-10'}, 10, 4),                              # date_to alone does not move it
        ({'date_from': '2026-10-01'}, 0, 14),                            # inside the window: an entry
        ({'date_from': '2026-10-05'}, 0, 14),                            # the day of the load
        ({'date_from': '2026-10-06'}, 10, 4),                            # before the window: balance
        ({'date_from': '2026-10-01', 'date_to': '2026-10-31'}, 0, 14),
    ])
    def test_sheets_and_summary_columns_agree_in_every_window(self, client, history, params, initial, entries):
        row = _summary_row(client, **params)
        assert (row['initial_stock'], row['entries']) == (initial, entries)

        book = _load(client.get(SUMMARY + 'export/', {'include_movements': 'true', **params}))
        summary_sheet = book['RESUMEN DE MOVIMIENTO']
        header = [c.value for c in summary_sheet[1]]
        exported = next(summary_sheet.iter_rows(min_row=2, values_only=True))
        assert exported[header.index('INV. INICIAL')] == initial
        # the same rule on both sides: the entries column IS the INGRESOS sheet
        assert _sheet_sum(book, 'INGRESOS', 5) == pytest.approx(entries)
        assert exported[header.index('INGRESOS')] == pytest.approx(entries)
        assert initial + _sheet_sum(book, 'INGRESOS', 5) - _sheet_sum(book, 'EGRESOS', 5) == pytest.approx(
            exported[header.index('STOCK ACTUAL')])

    def test_the_movements_only_export_has_no_initial_column_so_it_writes_the_load(self, client, history):
        """``layout=excel_pcc`` has no RESUMEN sheet: dropping the load there would lose 10 units."""
        book = _load(client.get(MOVEMENTS + 'export/', {'layout': 'excel_pcc'}))
        assert book.sheetnames == ['INGRESOS', 'EGRESOS']
        assert sorted(r[5] for r in book['INGRESOS'].iter_rows(min_row=2, values_only=True)) == [4, 10]

    def test_a_workbook_exported_with_a_window_that_holds_the_load_reimports_to_the_same_stock(
        self, client, history,
    ):
        window = {'date_from': '2026-10-01'}
        exported = client.get(SUMMARY + 'export/', {'include_movements': 'true', **window}).content
        expected = Decimal('11')                                          # 10 + 4 - 3
        InventoryMovement.objects.all().delete()
        history.stocks.all().delete()
        history.delete()

        preview = upload(client, exported)
        assert preview.status_code == 201, preview.content
        done = confirm(client, preview.json(), project_map={})
        assert done.status_code == 200, done.content
        from apps.warehouse.models import InventoryStock
        assert InventoryStock.objects.get(inventory__product_code='EPP000001').quantity == expected


# ---------------------------------------------------------------------------
# QA-22: el stock no puede salirse de DECIMAL(12, 3)
# ---------------------------------------------------------------------------

class TestStockColumnLimit:
    def test_an_entry_that_lands_exactly_on_the_maximum_is_accepted(self, client, make_item):
        make_item('LYP000009', 'LÍQUIDO', 'LT', stock=Decimal('999999998.000'))
        item = make_item('LYP000010', 'LÍQUIDO B', 'LT', stock=Decimal('999999998.000'))
        response = client.post(MOVEMENTS + 'entry/', {
            'inventory': item.pk, 'quantity': '1.999', 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE',
        }, format='json')
        assert response.status_code == 201, response.content
        stock = item.stocks.get()
        assert stock.quantity == MAX_STOCK_QUANTITY

    def test_one_unit_more_is_refused_with_a_message_and_writes_nothing(self, client, make_item):
        item = make_item('LYP000011', 'LÍQUIDO', 'LT', stock=Decimal('999999999.000'))
        before = snapshot_database()
        response = client.post(MOVEMENTS + 'entry/', {
            'inventory': item.pk, 'quantity': '1.000', 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE',
        }, format='json')
        assert response.status_code == 400
        assert 'stock máximo' in response.json()['detail']
        assert snapshot_database() == before

    def test_repeated_lines_of_one_item_are_added_before_checking(self, client, make_item):
        item = make_item('LYP000012', 'LÍQUIDO', 'LT', stock=Decimal('999999990'))
        lines = [{'inventory': item.pk, 'quantity': '6'}, {'inventory': item.pk, 'quantity': '6'}]
        response = client.post(MOVEMENTS + 'batch-entry/', {
            'items': lines, 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}, format='json')
        assert response.status_code == 400
        assert item.stocks.get().quantity == Decimal('999999990')

    def test_an_adjustment_cannot_overflow_either(self, make_item, central_warehouse):
        item = make_item('LYP000013', 'LÍQUIDO', 'LT', stock=Decimal('999999999.000'))
        with pytest.raises(ValueError, match='stock máximo'):
            svc.register_adjustment(item, Decimal('1'), 'CENTRAL', None, 'ajuste', central_warehouse)
        assert item.stocks.get().quantity == Decimal('999999999.000')
        assert not InventoryMovement.objects.exists()

    def test_the_planner_blocks_a_workbook_whose_rows_add_up_past_the_column(self, client):
        book = build_workbook(
            catalog=[catalog_row('LYP000014', 'LÍQUIDO', unit='LT')],
            entries=[entry_row('LYP000014', '999999999', D1), entry_row('LYP000014', '1', D1, po='OTRA')],
        )
        preview = upload(client, book).json()
        assert preview['can_confirm'] is False
        assert preview['summary']['by_code'].get('STOCK_EXCEDIDO') == 1
        assert preview['summary']['counts']['blocker'] >= 1

    def test_a_workbook_that_lands_on_the_maximum_is_not_blocked(self, client):
        book = build_workbook(
            catalog=[catalog_row('LYP000015', 'LÍQUIDO', unit='LT')],
            entries=[entry_row('LYP000015', '999999998', D1), entry_row('LYP000015', '1.999', D1, po='OTRA')],
        )
        preview = upload(client, book).json()
        assert 'STOCK_EXCEDIDO' not in preview['summary']['by_code']


# ---------------------------------------------------------------------------
# QA-26 / QA-29 / QA-16: bordes exactos
# ---------------------------------------------------------------------------

class TestBoundaries:
    def test_a_vale_of_exactly_two_hundred_lines_is_accepted(self, client, make_item):
        item = make_item('EPP000020', 'GUANTE', 'EA', stock=10)
        lines = [{'inventory': item.pk, 'quantity': 1} for _ in range(200)]
        response = client.post(MOVEMENTS + 'batch-entry/', {
            'items': lines, 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}, format='json')
        assert response.status_code == 201, response.content[:200]

    def test_notes_of_exactly_the_limit_are_accepted_and_one_more_is_not(self, client, item):
        body = {'inventory': item.pk, 'quantity': 1, 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}
        assert client.post(MOVEMENTS + 'entry/', {**body, 'notes': 'x' * 2000}, format='json').status_code == 201
        assert client.post(MOVEMENTS + 'entry/', {**body, 'notes': 'x' * 2001}, format='json').status_code == 400

    def test_the_first_accepted_operational_date_is_the_1st_of_january_2000(self, client, item):
        body = {'inventory': item.pk, 'quantity': 1, 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}
        assert client.post(MOVEMENTS + 'entry/', {**body, 'movement_date': '2000-01-01'}, format='json').status_code == 201
        assert client.post(MOVEMENTS + 'entry/', {**body, 'movement_date': '1999-12-31'}, format='json').status_code == 400

    def test_the_smallest_quantity_the_column_holds_is_accepted_by_entry_and_exit(self, client, make_item):
        """``min_value=0.001`` (a float) used to reject exactly 0.001: Decimal('0.001') < 0.001."""
        item = make_item('LYP000030', 'LÍQUIDO', 'LT', stock=5)
        entry = client.post(MOVEMENTS + 'entry/', {
            'inventory': item.pk, 'quantity': '0.001', 'warehouse': 'CENTRAL', 'source_type': 'PURCHASE'}, format='json')
        out = client.post(MOVEMENTS + 'exit/', {
            'inventory': item.pk, 'quantity': '0.001', 'warehouse': 'CENTRAL', 'destination_type': 'PROJECT',
            'destination_detail': 'Obra'}, format='json')
        assert (entry.status_code, out.status_code) == (201, 201), (entry.content[:160], out.content[:160])

    def test_check_stock_accepts_a_real_project_id(self, client, epp_item):
        response = client.get(INVENTORY + 'check-stock/', {'descriptions': 'guante', 'project_id': '1'})
        assert response.status_code == 200
        assert response.json()[0]['matched_product'] == epp_item.product_code


# ---------------------------------------------------------------------------
# QA-27: previsualizaciones vencidas
# ---------------------------------------------------------------------------

def _expire(import_id):
    InventoryImport.objects.filter(pk=import_id).update(expires_at=timezone.now() - timedelta(hours=1))


class TestExpiredPreviewsAreScrubbedOnTheRequestPath:
    def test_listing_the_history_purges_expired_previews(self, client):
        preview = upload(client, standard_workbook()).json()
        _expire(preview['id'])
        assert client.get(IMPORTS).status_code == 200
        stale = InventoryImport.objects.get(pk=preview['id'])
        assert stale.status == InventoryImportStatus.EXPIRED and not stale.file_bytes
        assert 'unmatched_destinations' not in stale.summary

    def test_reading_an_expired_preview_purges_it_and_shows_no_names(self, client):
        preview = upload(client, standard_workbook()).json()
        _expire(preview['id'])
        body = client.get(f"{IMPORTS}{preview['id']}/")
        assert body.status_code == 200
        assert body.json()['status'] == 'EXPIRED' and body.json()['unmatched_destinations'] == []
        assert WORKER_A.split(',')[0].upper() not in body.content.decode().upper()

    def test_a_failing_purge_never_fails_the_upload(self, client):
        with mock.patch.object(service, 'purge_stale_imports', side_effect=DatabaseError('lock timeout')):
            response = upload(client, standard_workbook())
        assert response.status_code == 201, response.content[:200]

    def test_a_fresh_preview_is_not_touched_by_the_purge(self, client):
        preview = upload(client, standard_workbook()).json()
        client.get(IMPORTS)
        fresh = InventoryImport.objects.get(pk=preview['id'])
        assert fresh.status == InventoryImportStatus.PREVIEWED and fresh.file_bytes
        assert client.get(f"{IMPORTS}{preview['id']}/").json()['unmatched_destinations']


# ---------------------------------------------------------------------------
# QA-17: la huella ajena nunca rompe
# ---------------------------------------------------------------------------

class TestForeignDigestIsAConflict:
    @pytest.mark.parametrize('digest', ['x', 'ñandú', '😀', 'é' * 64])
    def test_it_is_a_409_plan_changed_and_the_import_stays_previewed(self, client, digest):
        preview = upload(client, standard_workbook()).json()
        before = snapshot_database()
        response = client.post(f"{IMPORTS}{preview['id']}/confirm/", {
            'expected_plan_digest': digest, 'project_map': {'ON PLANT': None},
            'accept_warnings': True, 'accept_skipped_rows': True}, format='json')
        assert response.status_code == 409, (response.status_code, response.content[:160])
        assert response.json()['code'] == 'plan_changed'
        assert snapshot_database() == before
        assert InventoryImport.objects.get().status == InventoryImportStatus.PREVIEWED


# ---------------------------------------------------------------------------
# QA-14 / QA-15: ids acotados
# ---------------------------------------------------------------------------

class TestIdBounds:
    def test_the_largest_valid_id_is_looked_up_normally_and_is_a_clean_error(self, client):
        preview = upload(client, standard_workbook()).json()
        response = client.post(f"{IMPORTS}{preview['id']}/confirm/", {
            'expected_plan_digest': preview['plan_digest'], 'project_map': {'ON PLANT': 2 ** 63 - 1},
            'accept_warnings': True, 'accept_skipped_rows': True}, format='json')
        assert response.status_code == 400 and response.json()['code'] == 'project_not_found'

    @pytest.mark.parametrize('value', [0, -1])
    def test_zero_and_negative_ids_are_refused_by_the_serializer(self, client, value):
        preview = upload(client, standard_workbook()).json()
        response = client.post(f"{IMPORTS}{preview['id']}/confirm/", {
            'expected_plan_digest': preview['plan_digest'], 'project_map': {'ON PLANT': value},
            'accept_warnings': True, 'accept_skipped_rows': True}, format='json')
        assert response.status_code == 400

    def test_a_null_project_still_means_do_not_assign_one(self, client):
        preview = upload(client, standard_workbook()).json()
        response = confirm(client, preview)
        assert response.status_code == 200, response.content[:200]
