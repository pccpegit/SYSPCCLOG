"""SYSPCC-025 — resumen de movimiento (kardex agregado) y export con el formato de su hoja."""
import io
from datetime import date
from decimal import Decimal
from unittest import mock

import openpyxl
import pytest

from apps.warehouse.models import InventoryStock
from apps.warehouse.services import movements as svc
from .conftest import client_for

URL = '/api/v1/warehouse/inventory/movement-summary/'
EXPORT = URL + 'export/'


def _entry(item, qty, user, d=None, source='PURCHASE', warehouse='CENTRAL', **kw):
    kw = ({'undated': True} if d is None else {'movement_date': d}) | kw
    return svc.register_entry_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
        source_type=source, supplier_name='', invoice_number='', notes='', registered_by=user, **kw)


def _exit(item, qty, user, d=None, warehouse='CENTRAL', **kw):
    kw = ({'undated': True} if d is None else {'movement_date': d}) | kw
    return svc.register_exit_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
        destination_type='EMPLOYEE', destination_detail=kw.pop('destination_detail', 'Persona Test'),
        requested_by_id=None, authorized_by_id=None, notes='', registered_by=user, **kw)


def _rows(resp):
    return {r['product_code']: r for r in resp.data['results']}


@pytest.fixture
def ledger(make_item, central_warehouse, project):
    """Item A: inicial 10, +5 (05/10), -3 (08/10), -2 (15/10), +4 sin fecha. Item B sin movimientos."""
    a = make_item('EPP0001', 'GUANTE T: L', 'EA', project=project, floor=2, observation='obs A', size='L')
    b = make_item('EPP0002', 'CASCO', 'EA', location='F3-2', floor=2)
    u = central_warehouse
    _entry(a, 10, u, date(2026, 10, 1), source='INITIAL_LOAD')
    _entry(a, 5, u, date(2026, 10, 5))
    _exit(a, 3, u, date(2026, 10, 8))
    _exit(a, 2, u, date(2026, 10, 15))
    _entry(a, 4, u, None)
    return a, b


@pytest.mark.django_db
class TestMovementSummary:
    def test_closing_stock_is_initial_plus_entries_minus_exits(self, ledger, central_warehouse):
        resp = client_for(central_warehouse).get(URL)
        assert resp.status_code == 200
        a = _rows(resp)['EPP0001']
        assert (a['initial_stock'], a['entries'], a['exits']) == (10, 9, 5)       # el sin fecha cuenta sin filtro
        assert a['closing_stock'] == a['initial_stock'] + a['entries'] - a['exits'] == 14
        assert a['current_stock'] == 14                                            # coincide con InventoryStock
        assert resp.headers['Cache-Control'] == 'no-store'

    def test_row_shape_and_derived_location(self, ledger, central_warehouse, project):
        rows = _rows(client_for(central_warehouse).get(URL))
        a, b = rows['EPP0001'], rows['EPP0002']
        for key in ('inventory_id', 'product_code', 'description', 'unit', 'category', 'size',
                    'initial_stock', 'entries', 'exits', 'closing_stock', 'current_stock',
                    'location', 'location_status', 'floor', 'observation', 'project_code', 'project_name'):
            assert key in a, key
        assert (a['location'], a['location_status'], a['size']) == ('PISO 2', 'FLOOR_ONLY', 'L')
        assert (a['project_code'], a['observation']) == (project.code, 'obs A')
        assert (b['location'], b['location_status'], b['project_code']) == ('F3-2', 'EXACT', '')
        assert (b['initial_stock'], b['closing_stock'], b['current_stock']) == (0, 0, 0)

    def test_period_filter_moves_prior_balance_into_initial_stock(self, ledger, central_warehouse):
        resp = client_for(central_warehouse).get(URL, {'date_from': '2026-10-07', 'date_to': '2026-10-31'})
        a = _rows(resp)['EPP0001']
        assert a['initial_stock'] == 15            # 10 + 5 antes del 07/10
        assert (a['entries'], a['exits']) == (0, 5)
        assert a['closing_stock'] == 10
        assert a['current_stock'] == 14            # stock real, independiente del periodo
        assert resp.data['meta']['undated_movements'] == 1
        assert resp.data['meta']['totals'] == {
            'initial': 15.0, 'entries': 0.0, 'exits': 5.0, 'closing': 10.0, 'current': 14.0,
        }

    def test_date_to_only_cuts_the_ledger(self, ledger, central_warehouse):
        a = _rows(client_for(central_warehouse).get(URL, {'date_to': '2026-10-08'}))['EPP0001']
        assert (a['initial_stock'], a['entries'], a['exits'], a['closing_stock']) == (10, 5, 3, 12)

    def test_undated_are_counted_only_without_date_filter(self, ledger, central_warehouse):
        no_filter = client_for(central_warehouse).get(URL)
        assert no_filter.data['meta']['undated_movements'] == 0
        assert _rows(no_filter)['EPP0001']['entries'] == 9
        filtered = client_for(central_warehouse).get(URL, {'date_from': '2026-01-01'})
        # (a dated initial load inside the period counts as an entry of that period)
        assert _rows(filtered)['EPP0001']['entries'] == 15
        assert filtered.data['meta']['undated_movements'] == 1

    def test_adjustments_count_in_the_right_column(self, make_item, central_warehouse):
        item = make_item('EPP5', 'x', stock=None)
        _entry(item, 10, central_warehouse, date(2026, 10, 1), source='INITIAL_LOAD')
        svc.register_adjustment(item, Decimal('3'), 'CENTRAL', None, 'sobrante', central_warehouse)
        svc.register_adjustment(item, Decimal('-4'), 'CENTRAL', None, 'merma', central_warehouse)
        row = _rows(client_for(central_warehouse).get(URL))['EPP5']
        assert (row['initial_stock'], row['entries'], row['exits'], row['closing_stock']) == (10, 3, 4, 9)
        assert row['current_stock'] == 9

    def test_warehouse_filter_scopes_movements_and_stock(self, make_item, central_warehouse):
        item = make_item('EPP5', 'x')
        _entry(item, 6, central_warehouse, date(2026, 10, 1), warehouse='CENTRAL')
        _entry(item, 2, central_warehouse, date(2026, 10, 1), warehouse='SITE')
        client = client_for(central_warehouse)
        assert _rows(client.get(URL))['EPP5']['closing_stock'] == 8
        site = _rows(client.get(URL, {'warehouse': 'SITE'}))['EPP5']
        assert (site['entries'], site['current_stock']) == (2, 2)
        assert client.get(URL, {'warehouse': 'XX'}).status_code == 400

    def test_catalog_filters_and_flags(self, ledger, central_warehouse, project):
        client = client_for(central_warehouse)
        assert list(_rows(client.get(URL, {'search': 'casco'}))) == ['EPP0002']
        assert list(_rows(client.get(URL, {'project': project.pk}))) == ['EPP0001']
        assert list(_rows(client.get(URL, {'location_status': 'exact'}))) == ['EPP0002']
        assert list(_rows(client.get(URL, {'only_with_movement': 'true'}))) == ['EPP0001']
        assert list(_rows(client.get(URL, {'only_zero_stock': 'true'}))) == ['EPP0002']

    def test_ordering_whitelist(self, ledger, central_warehouse):
        client = client_for(central_warehouse)
        codes = lambda **p: [r['product_code'] for r in client.get(URL, p).data['results']]
        assert codes(ordering='-closing_stock') == ['EPP0001', 'EPP0002']
        assert codes(ordering='closing_stock') == ['EPP0002', 'EPP0001']
        assert codes(ordering='-product_code') == ['EPP0002', 'EPP0001']
        assert codes(ordering='password; DROP TABLE') == ['EPP0001', 'EPP0002']   # inválido -> por defecto

    def test_pagination_and_no_page(self, make_item, central_warehouse):
        for i in range(30):
            make_item(f'EPP{i:04d}', f'item {i}')
        client = client_for(central_warehouse)
        page = client.get(URL, {'page_size': 10})
        assert page.data['count'] == 30 and len(page.data['results']) == 10 and 'meta' in page.data
        full = client.get(URL, {'no_page': 'true'})
        assert len(full.data['results']) == 30 and full.data['meta']['truncated'] is False

    @pytest.mark.parametrize('params', [{'date_from': 'x'}, {'date_from': '2026-10-09', 'date_to': '2026-10-01'}])
    def test_invalid_dates_are_400(self, central_warehouse, params):
        assert client_for(central_warehouse).get(URL, params).status_code == 400


def _load(response):
    return openpyxl.load_workbook(io.BytesIO(response.content))


@pytest.mark.django_db
class TestMovementSummaryExport:
    def test_summary_sheet_has_the_exact_headers_of_the_real_workbook(self, ledger, central_warehouse):
        resp = client_for(central_warehouse).get(EXPORT)
        assert resp.status_code == 200
        assert resp.headers['Cache-Control'] == 'no-store'
        assert 'RESUMEN_MOVIMIENTO_' in resp.headers['Content-Disposition']
        wb = _load(resp)
        assert wb.sheetnames == ['RESUMEN DE MOVIMIENTO']
        ws = wb['RESUMEN DE MOVIMIENTO']
        assert [c.value for c in ws[1]] == [
            'N°', 'ITEM', 'DESCRIPCION', 'UOM', 'INV. INICIAL', 'INGRESOS', 'EGRESOS',
            'STOCK ACTUAL', 'LOCACIÓN', 'OBSERVACIÓN', 'PROYECTO']
        row = [c.value for c in ws[2]]
        assert row[1:8] == ['EPP0001', 'GUANTE T: L', 'EA', 10, 9, 5, 14]
        assert row[8] == 'PISO 2'
        # values, not formulas
        assert all(not (isinstance(c.value, str) and c.value.startswith('=')) for r in ws.iter_rows() for c in r)

    def test_stock_actual_equals_inicial_plus_ingresos_minus_egresos_in_every_row(self, ledger, central_warehouse):
        ws = _load(client_for(central_warehouse).get(EXPORT))['RESUMEN DE MOVIMIENTO']
        for r in list(ws.iter_rows(min_row=2, values_only=True)):
            assert r[7] == r[4] + r[5] - r[6]

    def test_export_respects_filters(self, ledger, central_warehouse):
        ws = _load(client_for(central_warehouse).get(EXPORT, {'search': 'casco'}))['RESUMEN DE MOVIMIENTO']
        assert ws.max_row == 2 and ws.cell(2, 2).value == 'EPP0002'

    def test_include_movements_adds_ingresos_and_egresos_sheets(self, ledger, central_warehouse):
        wb = _load(client_for(central_warehouse).get(EXPORT, {'include_movements': 'true'}))
        assert wb.sheetnames == ['RESUMEN DE MOVIMIENTO', 'INGRESOS', 'EGRESOS']
        assert [c.value for c in wb['INGRESOS'][1]] == [
            'N°', 'FECHA', 'TAG', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD INGRESADA', 'UOM', 'LOCACIÓN',
            'PO', 'GUÍA DE REMISIÓN', 'TRANSFERENCIA INTERNA', 'OBSERVACIÓN']
        assert [c.value for c in wb['EGRESOS'][1]] == [
            'N°', 'Nª VALE', 'FECHA', 'ITEM', 'DESCRIPCIÓN', 'CANTIDAD A EGRESAR', 'LOCACIÓN',
            'TRANSFERENCIA INTERNA', 'USUARIO/ DESTINO', 'OBSERVACIÓN']
        # la carga inicial viaja en INV. INICIAL, no se duplica en INGRESOS
        entries = list(wb['INGRESOS'].iter_rows(min_row=2, values_only=True))
        assert sorted(e[5] for e in entries) == [4, 5]
        exits = list(wb['EGRESOS'].iter_rows(min_row=2, values_only=True))
        assert sorted(e[5] for e in exits) == [2, 3]

    def test_include_movements_honours_the_date_range(self, ledger, central_warehouse):
        wb = _load(client_for(central_warehouse).get(
            EXPORT, {'include_movements': 'true', 'date_from': '2026-10-07', 'date_to': '2026-10-10'}))
        assert [r[5] for r in wb['EGRESOS'].iter_rows(min_row=2, values_only=True)] == [3]
        assert wb['INGRESOS'].max_row == 1       # solo cabecera

    def test_formula_injection_is_neutralized(self, make_item, central_warehouse):
        item = make_item('EPP9', '=HYPERLINK("http://evil")', observation='+cmd|calc')
        _entry(item, 1, central_warehouse, date(2026, 10, 1), source='INITIAL_LOAD')
        _exit(item, 1, central_warehouse, date(2026, 10, 2), destination_detail='@SUM(A1)')
        wb = _load(client_for(central_warehouse).get(EXPORT, {'include_movements': 'true'}))
        assert wb['RESUMEN DE MOVIMIENTO'].cell(2, 3).value == "'=HYPERLINK(\"http://evil\")"
        assert wb['RESUMEN DE MOVIMIENTO'].cell(2, 10).value == "'+cmd|calc"
        assert wb['EGRESOS'].cell(2, 9).value == "'@SUM(A1)"

    def test_truncation_writes_a_visible_warning_row(self, ledger, central_warehouse):
        with mock.patch('apps.core.utils.excel.MAX_EXPORT_ROWS', 1):
            ws = _load(client_for(central_warehouse).get(EXPORT))['RESUMEN DE MOVIMIENTO']
        assert 'ADVERTENCIA' in ws.cell(1, 1).value
        assert ws.cell(2, 1).value == 'N°'
        assert ws.max_row == 3

    def test_export_log_has_no_names(self, ledger, central_warehouse, app_logs):
        client_for(central_warehouse).get(EXPORT, {'include_movements': 'true'})
        assert 'warehouse.export' in app_logs.events()
        assert 'Persona Test' not in app_logs.dump()

    def test_movements_export_excel_pcc_layout(self, ledger, central_warehouse):
        resp = client_for(central_warehouse).get('/api/v1/warehouse/movements/export/', {'layout': 'excel_pcc'})
        assert resp.status_code == 200
        wb = _load(resp)
        assert wb.sheetnames == ['INGRESOS', 'EGRESOS']
        assert wb['EGRESOS'].max_row == 3

    def test_movements_export_default_layout_is_unchanged(self, ledger, central_warehouse):
        wb = _load(client_for(central_warehouse).get('/api/v1/warehouse/movements/export/'))
        assert wb.sheetnames == ['Movimientos']


@pytest.mark.django_db
def test_summary_totals_match_stock_table(ledger, central_warehouse):
    resp = client_for(central_warehouse).get(URL)
    total_stock = sum(s.quantity for s in InventoryStock.objects.all())
    assert Decimal(str(resp.data['meta']['totals']['closing'])) == total_stock
