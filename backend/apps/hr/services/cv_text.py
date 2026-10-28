"""Text/image extraction helpers shared by CV analysis and job-posting reading (SYSPCC-023).

Pure functions: bytes in, text/base64 out. They never log content (PII) and
never raise on hostile input from `pdf_text` (returns '' instead).
"""

import base64
import io

MAX_DOCX_CHARS = 60_000
MAX_IMAGE_SIDE = 2000


def docx_text(data: bytes, max_chars: int = MAX_DOCX_CHARS) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []

    def table_text(table):
        seen = set()
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                if cell.text.strip():
                    cells.append(cell.text.strip())
            if cells:
                parts.append(' | '.join(cells))

    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
    try:
        blocks = list(doc.iter_inner_content())
    except AttributeError:  # pragma: no cover - old python-docx
        blocks = list(doc.paragraphs) + list(doc.tables)
    for block in blocks:
        if hasattr(block, 'rows'):
            table_text(block)
        elif block.text.strip():
            parts.append(block.text.strip())
    return '\n'.join(parts)[:max_chars]


def pdf_text(data: bytes, max_pages: int, max_chars: int) -> str:
    """Text layer of a PDF (first `max_pages` pages, at most `max_chars`).
    A scanned PDF or a hostile one yields '' (hostile PDFs raise anything,
    including RecursionError, hence the broad catch)."""
    from pypdf import PdfReader

    parts: list[str] = []
    total = 0
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            return ''
        for page in list(reader.pages)[:max_pages]:
            try:
                chunk = page.extract_text() or ''
            except BaseException as exc:  # noqa: BLE001
                if isinstance(exc, (KeyboardInterrupt, SystemExit)):
                    raise
                continue
            if chunk.strip():
                parts.append(chunk.strip())
                total += len(chunk)
            if total >= max_chars:
                break
    except BaseException as exc:  # noqa: BLE001 - hostile PDF
        if isinstance(exc, (KeyboardInterrupt, SystemExit)):
            raise
        return ''
    return '\n'.join(parts)[:max_chars]


def image_b64(data: bytes) -> str:
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(data)) as img:
        img = ImageOps.exif_transpose(img)
        if img.mode in ('RGBA', 'LA', 'P'):
            img = img.convert('RGBA')
            background = Image.new('RGB', img.size, (255, 255, 255))
            background.paste(img, mask=img.split()[-1])
            img = background
        else:
            img = img.convert('RGB')
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        out = io.BytesIO()
        img.save(out, format='JPEG', quality=85)  # no EXIF is written
    return base64.standard_b64encode(out.getvalue()).decode('ascii')
