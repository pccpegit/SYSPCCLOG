"""
Excel del Master de RQ (SYSPCC-030), generado en el servidor con openpyxl en modo `write_only`.

Hojas
-----
* ``MASTER``: una fila por ítem, con las columnas del Excel «MASTER DE RQ» de gerencia más
  Línea, Seguimiento, Etapa del RQ y Alertas.
* ``RESUMEN POR RQ``: una fila por RQ con su % de atención y los conteos por estado.
* ``ENTREGAS``: una fila por entrega (o reversa) registrada por el almacén: RQ, ítem, fecha,
  guía, cantidad y responsable (usuario del sistema). Sin nombres de trabajadores ni DNI.

Mismos filtros y mismo alcance de visibilidad que el listado. Todo texto pasa por
`sanitize_excel_value` (defensa contra inyección de fórmulas) y se limpian los caracteres de
control que openpyxl rechazaría. Las fechas van como fechas reales (America/Lima, dd/mm/aaaa).

Tope de filas: `MAX_MASTER_ROWS` ítems (y `MAX_DELIVERY_ROWS` entregas). Si se supera, el libro
sale truncado con una fila de ADVERTENCIA antes del encabezado de cada hoja afectada (misma
convención que los exportes de almacén) y `ExportResult.truncated` es True.
"""
import io
import logging
import re
from dataclasses import dataclass

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.core.utils.excel import sanitize_excel_value
from apps.rq.services import item_tracking as rules
from apps.rq.services import request_master as master

logger = logging.getLogger(__name__)

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MASTER_SHEET = 'MASTER'
SUMMARY_SHEET = 'RESUMEN POR RQ'
DELIVERIES_SHEET = 'ENTREGAS'

MAX_MASTER_ROWS = 20000
MAX_DELIVERY_ROWS = 50000
DATE_FORMAT = 'dd/mm/yyyy'
PERCENT_FORMAT = '0.0'
TRUNCATED_WARNING = (
    'ADVERTENCIA: la exportación se truncó. Aplique filtros adicionales para ver el resto.'
)
NO_VALUE = '—'

MASTER_HEADERS = [
    'ID RQ', 'PROYECTO / DEPARTAMENTO', 'DESCRIPCIÓN', 'ÁREA', 'UNIDAD', 'CANT. SOLICITADA',
    'CANT. ATENDIDA', 'STOCK ALMACÉN CENTRAL', 'CANT. FALTANTE', 'FUENTE', 'PRIORIDAD',
    'RESPONSABLE (SOLICITANTE)', 'FECHA DE APROBACIÓN', 'FECHA REQUERIDA', 'FECHA ATENDIDA',
    'ESTADO DEL ÍTEM', 'N° DE GUÍA', 'COMENTARIO', 'PRESUPUESTO / ADICIONAL',
    'RESPONSABLE DE ATENCIÓN', 'APROBACIÓN DE GERENCIA', '% ATENCIÓN DEL RQ', 'LÍNEA',
    'SEGUIMIENTO', 'ETAPA DEL RQ', 'ALERTAS',
]
SUMMARY_HEADERS = [
    'ID RQ', 'PROYECTO / DEPARTAMENTO', 'FLUJO', 'ÁREA', 'SOLICITANTE', 'ETAPA', 'ESTADO DEL RQ',
    'FECHA DE APROBACIÓN', 'FECHA REQUERIDA', 'APROBACIÓN DE GERENCIA', '% ATENCIÓN',
    'ÍTEMS', 'ATENDIDOS', 'PARCIALES', 'PENDIENTES', 'STAND BY', 'SIN VINCULAR',
    'DESESTIMADOS', 'ANULADOS',
]
DELIVERIES_HEADERS = [
    'ID RQ', 'LÍNEA', 'DESCRIPCIÓN', 'TIPO', 'FECHA', 'N° DE GUÍA', 'CANTIDAD', 'RESPONSABLE',
]

_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
_MAX_CELL_CHARS = 32000


@dataclass
class ExportResult:
    content: bytes
    filename: str
    truncated: bool
    item_rows: int
    delivery_rows: int
    rq_rows: int


def _text(value) -> str:
    """Texto seguro para una celda: sin caracteres de control, acotado y sin fórmulas."""
    if value is None:
        return ''
    cleaned = _ILLEGAL.sub('', str(value))[:_MAX_CELL_CHARS]
    # sanitize_excel_value cubre = + - @ ; el tabulador y el retorno de carro también disparan fórmulas.
    if cleaned.startswith(('\t', '\r')):
        cleaned = "'" + cleaned
    return sanitize_excel_value(cleaned)


def _cell(ws, value, *, date=False, number_format=None):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Sí' if value else 'No'
    if date or number_format:
        cell = WriteOnlyCell(ws, value=value)
        cell.number_format = DATE_FORMAT if date else number_format
        return cell
    if isinstance(value, str):
        return _text(value)
    return value


def _start(ws, headers, truncated: bool):
    # En modo write_only las vistas de la hoja se escriben con la primera fila: va antes.
    ws.freeze_panes = f'A{3 if truncated else 2}'
    if truncated:
        ws.append([TRUNCATED_WARNING])
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(ws, value=header)
        cell.font = Font(bold=True)
        header_cells.append(cell)
    ws.append(header_cells)


def _master_row(ws, row: dict) -> list:
    guides = ' / '.join(row['guides'])
    comment = ' | '.join(part for part in (
        row['comments'],
        f"{row['attention_override_label']}: {row['attention_note']}" if row['attention_override'] else '',
    ) if part)
    progress = row['rq_progress']
    percent = progress['percent'] if progress else None
    missing = row['quantity_missing']
    gm = row['gm_approval_label'] + (f" ({row['gm_approval_kind_label']})" if row['gm_approval_kind'] else '')
    budget = row['budget_type'] or row['budget_classification_label']
    return [
        _cell(ws, row['rq_number']),
        _cell(ws, row['origin']),
        _cell(ws, row['description']),
        _cell(ws, row['requesting_area_label']),
        _cell(ws, row['unit']),
        _cell(ws, row['quantity']),
        _cell(ws, row['quantity_attended']),
        _cell(ws, row['central_stock']),
        NO_VALUE if missing is None else _cell(ws, missing),
        _cell(ws, row['item_source_label']),
        _cell(ws, row['priority_label']),
        _cell(ws, row['requested_by_name']),
        _cell(ws, row['approval_date'], date=True),
        _cell(ws, row['need_date'], date=True),
        _cell(ws, row['attended_date'], date=True),
        _cell(ws, row['status_label']),
        _cell(ws, guides),
        _cell(ws, comment),
        _cell(ws, budget),
        _cell(ws, row['attention_responsible']),
        _cell(ws, gm),
        _cell(ws, percent, number_format=PERCENT_FORMAT) if percent is not None else NO_VALUE,
        _cell(ws, row['line_number']),
        _cell(ws, row['tracking_label']),
        _cell(ws, row['stage_label']),
        _cell(ws, '; '.join(alert['label'] for alert in row['alerts'])),
    ]


def _summary_row(ws, row: dict) -> list:
    gm = row['gm_approval_label'] + (f" ({row['gm_approval_kind_label']})" if row['gm_approval_kind'] else '')
    return [
        _cell(ws, row['rq_number']),
        _cell(ws, row['origin']),
        _cell(ws, row['flow_label']),
        _cell(ws, row['requesting_area_label']),
        _cell(ws, row['requested_by_name']),
        _cell(ws, row['stage_label']),
        _cell(ws, row['rq_status_label']),
        _cell(ws, row['approval_date'], date=True),
        _cell(ws, row['need_date'], date=True),
        _cell(ws, gm),
        _cell(ws, row['percent'], number_format=PERCENT_FORMAT) if row['percent'] is not None else NO_VALUE,
        row['items_total'], row['items_attended'], row['items_partial'], row['items_pending'],
        row['items_stand_by'], row['items_unlinked'], row['items_discarded'], row['items_voided'],
    ]


def _delivery_row(ws, row: dict, delivery: dict) -> list:
    return [
        _cell(ws, row['rq_number']),
        _cell(ws, row['line_number']),
        _cell(ws, row['description']),
        'REVERSA' if delivery['is_reversal'] else 'ENTREGA',
        _cell(ws, delivery['date'], date=True),
        _cell(ws, delivery['guide']),
        _cell(ws, delivery['quantity']),
        _cell(ws, delivery['responsible']),
    ]


def _filename(user, params: master.MasterParams) -> str:
    from apps.rq.permissions import get_accessible_requests_queryset

    label = 'TODOS'
    if params.project:
        # El código sale solo si el usuario ve algún RQ de ese proyecto (no revela otros).
        code = (
            get_accessible_requests_queryset(user).filter(project_id=params.project)
            .values_list('project__code', flat=True).first()
        )
        label = re.sub(r'[^A-Za-z0-9_-]+', '', code or '') or f'PROYECTO{params.project}'
    return f'Master_RQ_{label}_{timezone.localdate():%Y-%m-%d}.xlsx'


def build_master_workbook(user, params: master.MasterParams) -> ExportResult:
    """Libro completo del Master para `user` con los filtros `params`."""
    matching = master.select_matching(user, params)
    truncated_items = len(matching) > MAX_MASTER_ROWS
    exported = matching[:MAX_MASTER_ROWS]

    wb = Workbook(write_only=True)
    ws_master = wb.create_sheet(MASTER_SHEET)
    ws_summary = wb.create_sheet(SUMMARY_SHEET)
    ws_deliveries = wb.create_sheet(DELIVERIES_SHEET)

    _start(ws_master, MASTER_HEADERS, truncated_items)

    # Se recorre en bloques: cada uno trae el detalle de sus ítems (filas, guías, responsables).
    delivery_rows = 0
    truncated_deliveries = False
    delivery_buffer = []
    can_override = master.user_is_logistics(user)
    # Cada ítem se clasifica UNA vez: los de los RQs completos alimentan a la vez el % de
    # atención de cada fila, la hoja de resumen y los estados de la hoja MASTER.
    request_ids = list(dict.fromkeys(light.request_id for light in matching))
    lights = master.lights_by_request(request_ids)
    results = master.results_of(lights)
    progress = master.progress_from_results(results)
    classified = {light.pk: light.result for items in lights.values() for light in items}
    for chunk in master._chunks([light.pk for light in exported]):
        by_pk = {item.pk: item for item in master.detail_queryset().filter(pk__in=chunk)}
        rows = master.build_rows(
            [by_pk[pk] for pk in chunk if pk in by_pk], can_override_role=can_override,
            progress=progress, classified=classified,
        )
        for row in rows:
            ws_master.append(_master_row(ws_master, row))
            for delivery in row['deliveries']:
                if delivery_rows >= MAX_DELIVERY_ROWS:
                    truncated_deliveries = True
                    break
                delivery_buffer.append(_delivery_row(ws_deliveries, row, delivery))
                delivery_rows += 1

    _start(ws_deliveries, DELIVERIES_HEADERS, truncated_deliveries)
    for line in delivery_buffer:
        ws_deliveries.append(line)

    summary, _kpis = master.summary_rows(user, params, matching=matching, results=results)
    _start(ws_summary, SUMMARY_HEADERS, False)
    for row in summary:
        ws_summary.append(_summary_row(ws_summary, row))

    buffer = io.BytesIO()
    wb.save(buffer)
    truncated = truncated_items or truncated_deliveries
    if truncated:
        logger.warning(
            'rq.master.export.truncated',
            extra={'user_id': user.pk, 'items_total': len(matching), 'max_rows': MAX_MASTER_ROWS},
        )
    return ExportResult(
        content=buffer.getvalue(),
        filename=_filename(user, params),
        truncated=truncated,
        item_rows=len(exported),
        delivery_rows=delivery_rows,
        rq_rows=len(summary),
    )
