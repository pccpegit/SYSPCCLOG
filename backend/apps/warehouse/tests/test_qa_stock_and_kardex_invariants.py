"""
SYSPCC-025 · QA ronda 1 — invariantes de stock y kardex.

* El stock nunca queda negativo por ninguna vía (exit, batch-exit, epp-delivery, ajuste), y un
  rechazo no deja rastro.
* El kardex cuadra: INICIAL + INGRESOS - EGRESOS = STOCK ACTUAL = suma de InventoryStock = saldo
  final del kardex, sin filtros y con cualquier ventana de fechas; las ventanas consecutivas se
  encadenan y el export suma lo mismo que el resumen.

Los tests marcados "DEFECTO QA-n" describen un comportamiento incorrecto de producción que esta
ronda encontró: fallan a propósito hasta que se corrija.
"""
import io
from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from itertools import count
from unittest import mock

import openpyxl
import pytest
from django.utils import timezone

from apps.warehouse.models import InventoryMovement, InventoryStock, MovementGroup
from apps.warehouse.services import movements as svc

from .conftest import client_for
from .qa_support import INVENTORY, MOVEMENTS, snapshot_database, tolerant_client

pytestmark = pytest.mark.django_db

SUMMARY = INVENTORY + 'movement-summary/'
D = date


@pytest.fixture(autouse=True)
def _no_pdf_task():
    with mock.patch('apps.warehouse.views._enqueue_group_pdf_upload') as enqueue:
        yield enqueue


# ---------------------------------------------------------------------------
# Stock no negativo
# ---------------------------------------------------------------------------

_vouchers = count(7000)


def call_exit(client, item, quantity, warehouse='CENTRAL'):
    return client.post(MOVEMENTS + 'exit/', {
        'inventory': item.pk, 'quantity': quantity, 'warehouse': warehouse,
        'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona Prueba',
    }, format='json')


def call_batch_exit(client, item, quantity, warehouse='CENTRAL'):
    return client.post(MOVEMENTS + 'batch-exit/', {
        'items': [{'inventory': item.pk, 'quantity': quantity}], 'warehouse': warehouse,
        'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona Prueba',
    }, format='json')


def call_epp(client, item, quantity, warehouse='CENTRAL'):
    return client.post(MOVEMENTS + 'epp-delivery/', {
        'warehouse': warehouse, 'voucher_number': str(next(_vouchers)),
        'movement_date': timezone.localdate().isoformat(),
        'recipient': {'name': 'Persona Prueba'}, 'items': [{'inventory': item.pk, 'quantity': quantity}],
    }, format='json')


EXIT_CALLS = pytest.mark.parametrize('call', [call_exit, call_batch_exit, call_epp],
                                     ids=['exit', 'batch-exit', 'epp-delivery'])


def stock_now(item, warehouse='CENTRAL'):
    row = InventoryStock.objects.filter(inventory=item, warehouse_type=warehouse).first()
    return row.quantity if row else Decimal('0')


@pytest.fixture
def stocked(make_item):
    return make_item('EPP000001', 'GUANTE T: L', 'EA', stock=5)


@pytest.fixture
def client(central_warehouse):
    return tolerant_client(central_warehouse)


@EXIT_CALLS
class TestStockNeverGoesNegative:
    def test_taking_exactly_what_exists_leaves_zero(self, call, client, stocked):
        assert call(client, stocked, '5').status_code == 201
        assert stock_now(stocked) == 0

    def test_one_piece_more_than_exists_is_a_400_and_changes_nothing(self, call, client, stocked):
        before = snapshot_database()
        response = call(client, stocked, '6')
        assert response.status_code == 400, response.status_code
        assert snapshot_database() == before

    def test_taking_everything_twice_only_works_once(self, call, client, stocked):
        assert call(client, stocked, '5').status_code == 201
        before = snapshot_database()
        assert call(client, stocked, '5').status_code == 400
        assert snapshot_database() == before
        assert stock_now(stocked) == 0

    @pytest.mark.parametrize('quantity', ['0', '-1', '-0.001', 'abc', '', None, '1e12', '999999999999999'])
    def test_nonsense_quantities_are_a_400_never_a_500(self, call, client, stocked, quantity):
        before = snapshot_database()
        response = call(client, stocked, quantity)
        assert response.status_code == 400, (quantity, response.status_code)
        assert snapshot_database() == before

    def test_a_fraction_of_a_piece_is_rejected(self, call, client, stocked):
        before = snapshot_database()
        assert call(client, stocked, '2.5').status_code == 400
        assert snapshot_database() == before

    def test_the_largest_representable_quantity_is_insufficient_stock_not_a_crash(self, call, client, stocked):
        assert call(client, stocked, '999999999.000').status_code == 400

    def test_an_item_that_never_had_stock_cannot_be_taken(self, call, client, make_item):
        empty = make_item('EPP000009', 'SIN STOCK', 'EA')
        before = snapshot_database()
        assert call(client, empty, '1').status_code == 400
        assert not InventoryStock.objects.filter(quantity__lt=0).exists()
        assert snapshot_database()['movements'] == before['movements']

    def test_stock_of_another_warehouse_does_not_cover_the_exit(self, call, client, make_item):
        site_only = make_item('EPP000010', 'SOLO OBRA', 'EA', stock=5, warehouse='SITE')
        assert call(client, site_only, '1', warehouse='CENTRAL').status_code == 400
        assert stock_now(site_only, 'SITE') == 5

    def test_decimal_units_accept_fractions_and_still_cannot_overdraw(self, call, client, make_item):
        liquid = make_item('LYPCC0001', 'LIQUIDO', 'LT', stock='2.5')
        assert call(client, liquid, '1.25').status_code == 201
        assert call(client, liquid, '1.251').status_code == 400
        assert call(client, liquid, '1.25').status_code == 201
        assert stock_now(liquid) == 0


class TestRepeatedLinesAreSummed:
    def test_batch_with_two_lines_of_one_item_cannot_exceed_the_stock(self, client, stocked):
        before = snapshot_database()
        response = client.post(MOVEMENTS + 'batch-exit/', {
            'items': [{'inventory': stocked.pk, 'quantity': 3}, {'inventory': stocked.pk, 'quantity': 3}],
            'warehouse': 'CENTRAL', 'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona Prueba',
        }, format='json')
        assert response.status_code == 400
        assert snapshot_database() == before

    def test_epp_delivery_with_two_lines_of_one_item_cannot_exceed_the_stock(self, client, stocked):
        before = snapshot_database()
        response = client.post(MOVEMENTS + 'epp-delivery/', {
            'warehouse': 'CENTRAL', 'voucher_number': str(next(_vouchers)),
            'movement_date': timezone.localdate().isoformat(), 'recipient': {'name': 'Persona Prueba'},
            'items': [{'inventory': stocked.pk, 'quantity': 3}, {'inventory': stocked.pk, 'quantity': 3}],
        }, format='json')
        assert response.status_code == 400
        assert snapshot_database() == before

    def test_two_lines_that_add_up_to_the_stock_are_fine(self, client, stocked):
        response = client.post(MOVEMENTS + 'batch-exit/', {
            'items': [{'inventory': stocked.pk, 'quantity': 3}, {'inventory': stocked.pk, 'quantity': 2}],
            'warehouse': 'CENTRAL', 'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona Prueba',
        }, format='json')
        assert response.status_code == 201
        assert stock_now(stocked) == 0


class TestAdjustmentsCannotOverdraw:
    def test_a_negative_adjustment_larger_than_the_stock_is_refused(self, stocked, central_warehouse):
        with pytest.raises(ValueError):
            svc.register_adjustment(stocked, Decimal('-6'), 'CENTRAL', None, 'conteo', central_warehouse)
        assert stock_now(stocked) == 5
        assert not InventoryMovement.objects.filter(movement_type='ADJUSTMENT').exists()

    def test_an_adjustment_to_exactly_zero_is_allowed(self, stocked, central_warehouse):
        svc.register_adjustment(stocked, Decimal('-5'), 'CENTRAL', None, 'conteo', central_warehouse)
        assert stock_now(stocked) == 0


# ---------------------------------------------------------------------------
# El kardex cuadra
# ---------------------------------------------------------------------------

def _entry(item, qty, user, when=None, source='PURCHASE', warehouse='CENTRAL'):
    extra = {'undated': True} if when is None else {'movement_date': when}
    return svc.register_entry_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
        source_type=source, supplier_name='', invoice_number='', notes='', registered_by=user, **extra)


def _exit(item, qty, user, when, warehouse='CENTRAL'):
    return svc.register_exit_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': qty}], warehouse=warehouse, project_id=None,
        destination_type='EMPLOYEE', destination_detail='Persona Prueba', requested_by_id=None,
        authorized_by_id=None, notes='', registered_by=user, movement_date=when)


@pytest.fixture
def history(make_item, central_warehouse):
    """A scripted life of three items: carga inicial, ingresos, egresos, ajustes, un sin fecha y obra."""
    u = central_warehouse
    a = make_item('EPP000001', 'GUANTE T: L', 'EA')
    b = make_item('EPP000002', 'CASCO', 'EA')
    c = make_item('LYPCC0001', 'LIQUIDO', 'LT')
    _entry(a, 20, u, D(2026, 9, 30), source='INITIAL_LOAD')
    _entry(a, 10, u, D(2026, 10, 5))
    _exit(a, 4, u, D(2026, 10, 8))
    _entry(a, 6, u, D(2026, 10, 12))
    _exit(a, 7, u, D(2026, 10, 15))
    svc.register_adjustment(a, Decimal('-2'), 'CENTRAL', None, 'conteo', u, movement_date=D(2026, 10, 16))
    svc.register_adjustment(a, Decimal('1'), 'CENTRAL', None, 'conteo', u, movement_date=D(2026, 10, 17))
    _entry(a, 3, u, None)                                              # sin fecha
    _entry(a, 4, u, D(2026, 10, 10), warehouse='SITE')                 # otro almacén
    _entry(b, 5, u, D(2026, 10, 5))
    _exit(b, 5, u, D(2026, 10, 6))
    _entry(c, Decimal('2.5'), u, D(2026, 10, 7))
    _exit(c, Decimal('0.75'), u, D(2026, 10, 9))
    return a, b, c


def summary_rows(client, **params):
    response = client.get(SUMMARY, {'no_page': 'true', **params})
    assert response.status_code == 200, response.content
    return {r['product_code']: r for r in response.json()['results']}, response.json()


def kardex(client, item, **params):
    response = client.get(f'{INVENTORY}{item.pk}/kardex/', params)
    assert response.status_code == 200, response.content
    return response.json()


def db_stock(item, warehouse=None) -> Decimal:
    qs = InventoryStock.objects.filter(inventory=item)
    if warehouse:
        qs = qs.filter(warehouse_type=warehouse)
    return sum((s.quantity for s in qs), Decimal('0'))


class TestKardexMatchesStockWithoutFilters:
    def test_closing_stock_equals_current_stock_equals_the_database_equals_the_kardex(self, history, central_warehouse):
        client = client_for(central_warehouse)
        rows, _ = summary_rows(client)
        for item in history:
            row = rows[item.product_code]
            card = kardex(client, item)
            expected = float(db_stock(item))
            assert row['closing_stock'] == row['current_stock'] == expected
            assert card['closing_stock'] == card['movements'][-1]['balance_after'] == expected
            assert row['initial_stock'] + row['entries'] - row['exits'] == pytest.approx(row['closing_stock'])

    @pytest.mark.parametrize('warehouse', ['CENTRAL', 'SITE'])
    def test_a_warehouse_filter_reconciles_with_that_warehouse_only(self, history, central_warehouse, warehouse):
        client = client_for(central_warehouse)
        rows, _ = summary_rows(client, warehouse=warehouse)
        a = history[0]
        assert rows[a.product_code]['closing_stock'] == float(db_stock(a, warehouse))
        assert kardex(client, a, warehouse=warehouse)['closing_stock'] == float(db_stock(a, warehouse))

    def test_each_kardex_line_moves_the_balance_by_its_signed_quantity(self, history, central_warehouse):
        card = kardex(client_for(central_warehouse), history[0])
        balance = card['opening_stock']
        for line in card['movements']:
            balance += line['delta']
            assert line['balance_after'] == pytest.approx(balance)
        assert balance == pytest.approx(card['closing_stock'])

    def test_adjustments_move_the_balance_in_the_right_direction(self, history, central_warehouse):
        card = kardex(client_for(central_warehouse), history[0])
        by_type = [(m['movement_type'], m['delta']) for m in card['movements'] if m['movement_type'] == 'ADJUSTMENT']
        assert sorted(d for _t, d in by_type) == [-2, 1]


WINDOWS = [
    (None, D(2026, 10, 7)),
    (D(2026, 10, 8), D(2026, 10, 12)),
    (D(2026, 10, 13), D(2026, 10, 17)),
    (D(2026, 10, 18), None),
    (D(2026, 10, 5), D(2026, 10, 5)),
]


def _params(window):
    start, end = window
    return {k: v.isoformat() for k, v in (('date_from', start), ('date_to', end)) if v}


class TestKardexMatchesStockInAnyWindow:
    @pytest.mark.parametrize('window', WINDOWS, ids=lambda w: f'{w[0]}..{w[1]}')
    def test_initial_plus_entries_minus_exits_is_closing_and_the_kardex_agrees(
        self, history, central_warehouse, window,
    ):
        client = client_for(central_warehouse)
        rows, _ = summary_rows(client, **_params(window))
        for item in history:
            row = rows[item.product_code]
            card = kardex(client, item, **_params(window))
            assert row['initial_stock'] + row['entries'] - row['exits'] == pytest.approx(row['closing_stock'])
            if window[0]:
                assert card['opening_stock'] == pytest.approx(row['initial_stock'])
            assert card['closing_stock'] == pytest.approx(row['closing_stock']) or window[0] is None

    def test_consecutive_windows_chain_the_closing_of_one_into_the_initial_of_the_next(
        self, history, central_warehouse,
    ):
        client = client_for(central_warehouse)
        first, _ = summary_rows(client, date_from='2026-10-01', date_to='2026-10-07')
        second, _ = summary_rows(client, date_from='2026-10-08', date_to='2026-10-12')
        third, _ = summary_rows(client, date_from='2026-10-13', date_to='2026-10-17')
        for code in first:
            assert second[code]['initial_stock'] == pytest.approx(first[code]['closing_stock'])
            assert third[code]['initial_stock'] == pytest.approx(second[code]['closing_stock'])

    def test_dated_history_closes_at_stock_minus_what_has_no_date(self, history, central_warehouse):
        client = client_for(central_warehouse)
        rows, payload = summary_rows(client, date_from='2026-09-01', date_to='2026-12-31')
        a = history[0]
        assert rows[a.product_code]['closing_stock'] == float(db_stock(a)) - 3     # the undated +3
        assert payload['meta']['undated_movements'] == 1

    def test_a_window_with_no_movements_keeps_the_prior_balance(self, history, central_warehouse):
        rows, _ = summary_rows(client_for(central_warehouse), date_from='2026-11-01', date_to='2026-11-30')
        a = rows[history[0].product_code]
        assert (a['entries'], a['exits']) == (0, 0)
        assert a['initial_stock'] == a['closing_stock'] > 0

    def test_totals_in_meta_cover_every_page_not_only_the_current_one(self, history, central_warehouse):
        client = client_for(central_warehouse)
        paged = client.get(SUMMARY, {'page_size': 1}).json()
        rows, _ = summary_rows(client)
        assert len(paged['results']) == 1 and paged['count'] == 3
        assert paged['meta']['totals']['closing'] == pytest.approx(sum(r['closing_stock'] for r in rows.values()))
        assert paged['meta']['totals']['exits'] == pytest.approx(sum(r['exits'] for r in rows.values()))


class TestExportAddsUpLikeTheSummary:
    def test_the_movement_sheets_add_up_to_the_summary_totals(self, history, central_warehouse):
        client = client_for(central_warehouse)
        _, payload = summary_rows(client)
        response = client.get(SUMMARY + 'export/', {'include_movements': 'true'})
        assert response.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        entries = sum(r[5] for r in wb['INGRESOS'].iter_rows(min_row=2, values_only=True))
        exits = sum(r[5] for r in wb['EGRESOS'].iter_rows(min_row=2, values_only=True))
        totals = payload['meta']['totals']
        assert entries == pytest.approx(totals['entries'])
        assert exits == pytest.approx(totals['exits'])

    def test_every_summary_row_of_the_workbook_reconciles(self, history, central_warehouse):
        response = client_for(central_warehouse).get(SUMMARY + 'export/')
        ws = openpyxl.load_workbook(io.BytesIO(response.content))['RESUMEN DE MOVIMIENTO']
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 3
        for _n, _code, _desc, _uom, initial, entries, exits, closing, *_rest in rows:
            assert initial + entries - exits == pytest.approx(closing)

    def test_an_empty_database_exports_a_valid_workbook_with_only_headers(self, central_warehouse):
        response = client_for(central_warehouse).get(SUMMARY + 'export/', {'include_movements': 'true'})
        assert response.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames == ['RESUMEN DE MOVIMIENTO', 'INGRESOS', 'EGRESOS']
        assert all(ws.max_row == 1 for ws in wb.worksheets)


# ---------------------------------------------------------------------------
# Dashboard: "hoy" es la fecha operativa de Lima
# ---------------------------------------------------------------------------

class TestDashboardTodayCounters:
    """
    DEFECTO QA-4: `inventory/summary/` cuenta "hoy" con `created_at__date`, la fecha de AUDITORÍA, y no
    con `movement_date`, la fecha operativa que SYSPCC-025 introdujo. Tras importar el libro (movimientos
    de octubre registrados hoy) el tablero marca "Ingresos hoy / Egresos hoy" con TODO lo importado.

    DEFECTO QA-5: además compara contra `timezone.now().date()` (fecha UTC) mientras `__date` convierte
    a Lima: desde las 19:00 de Lima el contador de "hoy" se queda en 0.
    """

    def test_movements_dated_in_the_past_are_not_counted_as_today(self, history, central_warehouse):
        client = client_for(central_warehouse)
        summary = client.get(INVENTORY + 'summary/').json()
        assert summary['today_entries'] == 0
        assert summary['today_exits'] == 0

    def test_a_movement_dated_today_is_counted_as_today(self, history, central_warehouse):
        item = history[1]
        _entry(item, 1, central_warehouse, timezone.localdate())
        _exit(item, 1, central_warehouse, timezone.localdate())
        summary = client_for(central_warehouse).get(INVENTORY + 'summary/').json()
        assert (summary['today_entries'], summary['today_exits']) == (1, 1)

    def test_after_seven_pm_in_lima_todays_movements_are_still_counted(self, make_item, central_warehouse):
        item = make_item('EPP000077', 'CASCO', 'EA')
        evening = datetime(2026, 10, 20, 1, 30, tzinfo=dt_timezone.utc)       # 20:30 of 19/10 in Lima
        with mock.patch('django.utils.timezone.now', return_value=evening):
            # `_entry` without a date stores NULL (`undated`): the vale must be dated today (Lima: 19/10)
            _entry(item, 2, central_warehouse, timezone.localdate())
            summary = client_for(central_warehouse).get(INVENTORY + 'summary/').json()
        assert summary['today_entries'] == 1


# ---------------------------------------------------------------------------
# Parámetros de consulta hostiles
# ---------------------------------------------------------------------------

JUNK = [
    ('inventory/', {'floor': 'abc'}), ('inventory/', {'project': 'abc'}), ('inventory/', {'low_stock': 'maybe'}),
    ('inventory/', {'location_status': 'zzz'}), ('inventory/', {'page': '0'}), ('inventory/', {'page_size': 'x'}),
    ('inventory/', {'search': "%_'\";--" + 'ñ' * 400}), ('inventory/', {'ordering': 'DROP TABLE'}),
    ('inventory/movement-summary/', {'date_from': 'ayer'}), ('inventory/movement-summary/', {'date_to': '2026-02-30'}),
    ('inventory/movement-summary/', {'ordering': '-closing_stock;DROP'}),
    ('inventory/movement-summary/', {'project': 'abc'}), ('inventory/movement-summary/', {'floor': '99'}),
    ('inventory/movement-summary/', {'warehouse': 'MARTE'}), ('inventory/movement-summary/', {'page': 'ultima'}),
    ('inventory/movement-summary/export/', {'date_from': 'ayer'}),
    ('inventory/movement-summary/export/', {'include_movements': 'true', 'warehouse': 'MARTE'}),
    ('movements/', {'date_from': 'x'}), ('movements/', {'movement_type': 'ZZZ'}), ('movements/', {'warehouse': 'ZZZ'}),
    ('movements/', {'destination_personal': 'abc'}), ('movements/', {'import_batch': 'no-es-uuid'}),
    ('movements/', {'has_date': 'quizas'}), ('movements/', {'search': "'" * 300}),
    ('movements/', {'project': 'abc'}), ('movements/', {'qty_min': 'abc'}), ('movements/', {'qty_max': '1e999'}),
    ('movements/export/', {'layout': 'excel_pcc', 'date_from': 'x'}),
    ('movements/export/', {'qty_min': 'abc'}),
]


@pytest.mark.parametrize('path,params', JUNK, ids=lambda v: str(v)[:40])
def test_junk_query_parameters_never_cause_a_server_error(path, params, central_warehouse, history):
    response = tolerant_client(central_warehouse).get('/api/v1/warehouse/' + path, params)
    assert response.status_code < 500, (path, params, response.status_code)


@pytest.mark.parametrize('params', [{'date_from': 'ayer'}, {'date_from': '2026-10-09', 'date_to': '2026-10-01'}])
def test_kardex_rejects_bad_dates_with_a_400(params, central_warehouse, history):
    response = tolerant_client(central_warehouse).get(f'{INVENTORY}{history[0].pk}/kardex/', params)
    assert response.status_code == 400


def test_kardex_of_an_item_that_does_not_exist_is_a_404(central_warehouse):
    assert tolerant_client(central_warehouse).get(f'{INVENTORY}999999/kardex/').status_code == 404


def test_the_unpaginated_summary_is_capped_and_says_so(central_warehouse, make_item):
    from apps.warehouse.models import Inventory
    Inventory.objects.bulk_create(
        Inventory(product_code=f'X{n:05d}', description='x', unit='EA') for n in range(1001)
    )
    payload = client_for(central_warehouse).get(SUMMARY, {'no_page': 'true'}).json()
    assert payload['count'] == 1001
    assert len(payload['results']) == 1000
    assert payload['meta']['truncated'] is True


def test_page_size_is_capped_at_100(central_warehouse):
    from apps.warehouse.models import Inventory
    Inventory.objects.bulk_create(Inventory(product_code=f'Y{n:04d}', description='y', unit='EA') for n in range(150))
    payload = client_for(central_warehouse).get(SUMMARY, {'page_size': 100000}).json()
    assert len(payload['results']) == 100


def test_ordering_by_closing_stock_descending_really_sorts(history, central_warehouse):
    payload = client_for(central_warehouse).get(SUMMARY, {'ordering': '-closing_stock'}).json()
    closings = [r['closing_stock'] for r in payload['results']]
    assert closings == sorted(closings, reverse=True)


def test_an_unknown_ordering_falls_back_to_the_code(history, central_warehouse):
    payload = client_for(central_warehouse).get(SUMMARY, {'ordering': 'password'}).json()
    codes = [r['product_code'] for r in payload['results']]
    assert codes == sorted(codes)


def test_the_group_numbers_keep_growing_past_9999(central_warehouse, make_item, django_user_model):
    """
    A string sort puts G10000 before G9999, so after the 9 999th vale of a year the next number repeats
    G10000 and the unique constraint turns every registration into a 409.
    """
    item = make_item('EPP000042', 'CASCO', 'EA', stock=5)
    year = timezone.now().year
    for number in (9999, 10000):
        MovementGroup.objects.create(
            group_number=f'SAL-{year}-G{number}', movement_type='EXIT', warehouse='CENTRAL',
            registered_by=central_warehouse,
        )
    group = svc.register_exit_batch(
        items_data=[{'inventory_id': item.pk, 'quantity': 1}], warehouse='CENTRAL', project_id=None,
        destination_type='EMPLOYEE', destination_detail='Persona Prueba', requested_by_id=None,
        authorized_by_id=None, notes='', registered_by=central_warehouse,
    )
    assert group.group_number == f'SAL-{year}-G10001'
