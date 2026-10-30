"""SYSPCC-025 — importer REST endpoints (/api/v1/warehouse/imports/)."""
import io
import uuid
from datetime import date

import pytest
from django.core.cache.backends.locmem import LocMemCache
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.core.enums import RoleChoices
from apps.warehouse.models import Inventory, InventoryImport, InventoryMovement, InventoryStock, MovementGroup
from apps.warehouse.services import movements as movement_svc
from apps.warehouse.services.imports import parser
from apps.warehouse.throttling import WarehouseImportThrottle

from .conftest import client_for, make_user
from .import_workbooks import (
    WORKER_A, WORKER_B, build_workbook, catalog_row, entry_row, exit_row, rewrite_zip, standard_workbook,
)

pytestmark = pytest.mark.django_db

API = '/api/v1/warehouse/imports/'
XLSX = parser.XLSX_CONTENT_TYPE
NAMES = ('Quispe', 'Mamani', 'Ccori', 'Huaman', 'Rosa', 'Pedro')


@pytest.fixture(autouse=True)
def isolated_throttle(monkeypatch):
    """Upload throttle counters live in redis in this stack: give every test its own."""
    monkeypatch.setattr(WarehouseImportThrottle, 'cache', LocMemCache(f'imports-{uuid.uuid4()}', {}))


@pytest.fixture
def client(central_warehouse):
    return client_for(central_warehouse)


def post_upload(client, data, name='libro.xlsx', content_type=XLSX, **extra):
    file = SimpleUploadedFile(name, data, content_type=content_type)
    return client.post(API, {'file': file, **extra}, format='multipart')


def confirm_body(preview, **overrides):
    body = {
        'expected_plan_digest': preview['plan_digest'], 'project_map': {'ON PLANT': None},
        'accept_warnings': True, 'accept_skipped_rows': True,
    }
    body.update(overrides)
    return body


def confirm(client, preview, **overrides):
    return client.post(f"{API}{preview['id']}/confirm/", confirm_body(preview, **overrides), format='json')


class TestUpload:
    def test_upload_returns_the_preview(self, client):
        response = post_upload(client, standard_workbook())
        assert response.status_code == 201
        body = response.json()
        assert body['status'] == 'PREVIEWED' and body['can_confirm'] is True and len(body['plan_digest']) == 64
        assert body['summary']['catalog']['new'] == 3
        assert body['summary']['counts']['blocker'] == 0
        assert body['projects'] == [{'label': 'ON PLANT', 'items': 3, 'matches': [], 'suggested_id': None}]
        assert {d['display'] for d in body['unmatched_destinations']} == {WORKER_A, WORKER_B}
        assert body['requires']['accept_warnings'] is True
        assert 'projects' not in body['summary'] and 'unmatched_destinations' not in body['summary']

    def test_responses_are_never_cached(self, client):
        response = post_upload(client, standard_workbook())
        assert response['Cache-Control'] == 'no-store'
        assert client.get(f"{API}{response.json()['id']}/")['Cache-Control'] == 'no-store'
        assert client.get(API)['Cache-Control'] == 'no-store'

    def test_upload_does_not_change_the_inventory(self, client):
        post_upload(client, standard_workbook())
        assert Inventory.objects.count() == 0 and InventoryMovement.objects.count() == 0

    def test_blocked_book_is_previewed_but_not_confirmable(self, client):
        response = post_upload(client, build_workbook(entries=[entry_row('EPP1', 1)]))
        assert response.status_code == 201 and response.json()['can_confirm'] is False

    @pytest.mark.parametrize('name,content_type,code', [
        ('libro.xlsm', XLSX, 'extension_not_allowed'),
        ('libro.xls', 'application/vnd.ms-excel', 'extension_not_allowed'),
        ('libro.xlsx', 'text/csv', 'ARCHIVO_INVALIDO'),
    ])
    def test_wrong_files_are_rejected_with_a_clear_code(self, client, name, content_type, code):
        response = post_upload(client, standard_workbook(), name=name, content_type=content_type)
        assert response.status_code == 400
        body = response.json()
        assert body['error'] is True and body['status_code'] == 400 and body['code'] == code and body['detail']

    def test_not_a_zip(self, client):
        response = post_upload(client, b'hola mundo')
        assert response.status_code == 400 and response.json()['code'] == 'ARCHIVO_INVALIDO'

    def test_missing_file(self, client):
        response = client.post(API, {}, format='multipart')
        assert response.status_code == 400 and response.json()['error'] is True

    def test_oversize_file(self, client, monkeypatch):
        monkeypatch.setattr(parser, 'MAX_UPLOAD_BYTES', 1000)
        response = post_upload(client, standard_workbook())
        assert response.status_code == 400 and response.json()['code'] == 'file_too_large'

    def test_zip_bomb(self, client):
        bomb = rewrite_zip(standard_workbook(), add={'xl/media/z.bin': b'\0' * (6 * 1024 * 1024)})
        assert post_upload(client, bomb).status_code == 400

    def test_a_book_with_millions_of_shared_strings_is_refused_before_it_is_loaded(self, client):
        """Memory guard: that table is loaded in RAM by openpyxl and no size limit bounded it."""
        sst = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1">'
            + '<si><t>x</t></si>' * (parser.MAX_SHARED_STRINGS + 1) + '</sst>'
        ).encode()
        data = rewrite_zip(standard_workbook(), add={'xl/sharedStrings.xml': sst})
        before = InventoryImport.objects.count()
        response = post_upload(client, data)
        assert response.status_code == 400
        assert response.json()['code'] == 'ARCHIVO_INVALIDO'
        assert InventoryImport.objects.count() == before

    def test_only_the_central_warehouse_can_be_imported(self, client):
        response = post_upload(client, standard_workbook(), warehouse='SITE')
        assert response.status_code == 400 and response.json()['code'] == 'warehouse_not_supported'
        assert post_upload(client, standard_workbook(), warehouse='NOPE').status_code == 400

    def test_update_existing_flag(self, client, make_item):
        make_item('EPP000001', 'GUANTE DE CUERO T: L', purchase_order_ref='9.99')
        body = post_upload(client, standard_workbook(), update_existing='true').json()
        assert body['options'] == {'update_existing': True}
        assert body['summary']['catalog']['to_update'] >= 1

    def test_file_already_imported(self, client):
        data = standard_workbook()
        first = post_upload(client, data).json()
        assert confirm(client, first).status_code == 200
        response = post_upload(client, data)
        assert response.status_code == 409
        assert response.json()['code'] == 'file_already_imported' and response.json()['import_id'] == first['id']

    def test_upload_is_rate_limited(self, client, monkeypatch):
        monkeypatch.setattr(WarehouseImportThrottle, 'THROTTLE_RATES', {'warehouse_import': '2/hour'})
        codes = [
            post_upload(client, build_workbook(catalog=[catalog_row(f'EPP{n}', 'X')])).status_code for n in range(3)
        ]
        assert codes == [201, 201, 429]


class TestConfirmEndpoint:
    def test_confirm_applies_the_import_and_replays_idempotently(self, client):
        preview = post_upload(client, standard_workbook()).json()
        response = confirm(client, preview)
        assert response.status_code == 200 and 'Idempotent-Replay' not in response
        body = response.json()
        assert body['status'] == 'CONFIRMED' and body['result']['reconciled'] is True
        assert body['result']['catalog'] == {'created': 3, 'updated': 0}
        assert body['unmatched_destinations'] == []

        again = confirm(client, preview)
        assert again.status_code == 200 and again['Idempotent-Replay'] == 'true'
        assert again.json()['result'] == body['result']
        assert InventoryMovement.objects.count() == 5

    def test_stock_matches_the_books_summary_through_the_movement_summary_endpoint(self, client):
        preview = post_upload(client, standard_workbook()).json()
        confirm(client, preview)
        rows = client.get(
            '/api/v1/warehouse/inventory/movement-summary/', {'no_page': 'true'},
        ).json()['results']
        by_code = {r['product_code']: r for r in rows}
        assert (by_code['EPP000001']['entries'], by_code['EPP000001']['exits'], by_code['EPP000001']['closing_stock']) == (10, 5, 5)
        assert by_code['EPP000002']['closing_stock'] == 4 and by_code['EPP000003']['closing_stock'] == 0
        assert all(r['closing_stock'] == r['current_stock'] for r in rows)
        assert by_code['EPP000001']['location'] == 'PISO 2' and by_code['EPP000001']['location_status'] == 'FLOOR_ONLY'

    def test_missing_acceptance_is_a_400_with_the_missing_flags(self, client):
        preview = post_upload(client, standard_workbook()).json()
        response = confirm(client, preview, accept_warnings=False)
        assert response.status_code == 400
        assert response.json()['code'] == 'acceptance_required' and response.json()['required'] == ['accept_warnings']

    def test_missing_project_map(self, client):
        preview = post_upload(client, standard_workbook()).json()
        response = confirm(client, preview, project_map={})
        assert response.status_code == 400 and response.json()['code'] == 'project_map_required'

    def test_blockers_are_listed(self, client):
        preview = post_upload(client, build_workbook(entries=[entry_row('EPP1', 1)])).json()
        response = confirm(client, preview)
        assert response.status_code == 400 and response.json()['code'] == 'blockers'
        assert response.json()['problems'][0]['code'] == 'HOJA_FALTANTE'

    def test_changed_data_is_a_409(self, client, make_item):
        preview = post_upload(client, standard_workbook()).json()
        make_item('EPP000002', 'CASCO BLANCO')
        response = confirm(client, preview)
        assert response.status_code == 409 and response.json()['code'] == 'plan_changed'

    def test_body_validation_uses_the_error_envelope(self, client):
        preview = post_upload(client, standard_workbook()).json()
        response = client.post(f"{API}{preview['id']}/confirm/", {}, format='json')
        assert response.status_code == 400 and response.json()['error'] is True

    def test_project_map_accepts_a_real_project(self, client, project):
        preview = post_upload(client, standard_workbook()).json()
        assert confirm(client, preview, project_map={'ON PLANT': project.pk}).status_code == 200
        assert set(Inventory.objects.values_list('project_id', flat=True)) == {project.pk}

    def test_destination_map_links_workers(self, client, make_personal):
        person = make_personal('Otro Nombre Distinto')
        preview = post_upload(client, standard_workbook()).json()
        key = next(d['key'] for d in preview['unmatched_destinations'] if d['display'] == WORKER_B)
        assert confirm(client, preview, destination_map={key: person.pk}).status_code == 200
        assert InventoryMovement.objects.filter(destination_personal=person).count() == 1

    # Fixed UUID (not uuid4()): test IDs must be identical in every pytest-xdist worker.
    @pytest.mark.parametrize('bad_id', ['00000000-0000-4000-8000-000000000000', 'z' * 36, '1234'])
    def test_unknown_or_malformed_ids_are_404(self, client, bad_id):
        assert client.get(f'{API}{bad_id}/').status_code == 404
        assert client.post(f'{API}{bad_id}/confirm/', {'expected_plan_digest': 'x'}, format='json').status_code == 404
        assert client.delete(f'{API}{bad_id}/').status_code == 404


class TestCancel:
    def test_cancel_a_preview(self, client):
        preview = post_upload(client, standard_workbook()).json()
        response = client.delete(f"{API}{preview['id']}/")
        assert response.status_code == 200 and response.json() == {'id': preview['id'], 'status': 'CANCELLED'}
        assert bytes(InventoryImport.objects.get(pk=preview['id']).file_bytes) == b''
        assert confirm(client, preview).status_code == 409

    def test_a_confirmed_import_cannot_be_cancelled(self, client):
        preview = post_upload(client, standard_workbook()).json()
        confirm(client, preview)
        response = client.delete(f"{API}{preview['id']}/")
        assert response.status_code == 409 and response.json()['code'] == 'import_not_cancellable'


class TestProblems:
    @pytest.fixture
    def preview(self, client):
        data = build_workbook(
            catalog=[catalog_row('EPP1', 'X')],
            entries=[entry_row('EPP1', None), entry_row('EPP1', 3, when=None), entry_row('EPP999', 1)],
            exits=[exit_row('EPP1', 1, WORKER_A)],
        )
        return post_upload(client, data).json()

    def test_problems_are_paginated_and_most_severe_first(self, client, preview):
        body = client.get(f"{API}{preview['id']}/problems/", {'page_size': 2}).json()
        assert body['count'] >= 4 and len(body['results']) == 2 and body['next']
        assert body['results'][0]['severity'] == 'blocker'
        assert body['meta']['counts']['blocker'] == 1 and body['meta']['truncated'] is False
        assert set(body['results'][0]) == {'code', 'severity', 'sheet', 'row', 'message'}

    def test_filters(self, client, preview):
        url = f"{API}{preview['id']}/problems/"
        only_errors = client.get(url, {'severity': 'error'}).json()['results']
        assert {p['severity'] for p in only_errors} == {'error'}
        both = client.get(url, {'severity': 'blocker,error'}).json()['count']
        assert both == len(only_errors) + 1
        by_sheet = client.get(url, {'sheet': 'egresos'}).json()['results']
        assert by_sheet and {p['sheet'] for p in by_sheet} == {'EGRESOS'}
        by_code = client.get(url, {'code': 'cantidad_vacia'}).json()['results']
        assert [p['code'] for p in by_code] == ['CANTIDAD_VACIA']
        assert client.get(url, {'severity': 'nonsense'}).json()['count'] >= 4          # ignored, not an error

    def test_export_has_no_worker_names_and_is_a_real_workbook(self, client, preview):
        response = client.get(f"{API}{preview['id']}/problems/export/")
        assert response.status_code == 200 and response['Content-Type'] == XLSX
        assert 'PROBLEMAS_IMPORTACION_' in response['Content-Disposition'] and response['Cache-Control'] == 'no-store'
        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ('Gravedad', 'Hoja', 'Fila', 'Tipo de problema', 'Mensaje') and len(rows) > 3
        text = ' '.join(str(c) for row in rows for c in row)
        assert not any(name in text for name in NAMES)

    def test_export_respects_the_filters(self, client, preview):
        response = client.get(f"{API}{preview['id']}/problems/export/", {'severity': 'blocker'})
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        assert len(rows) == 2 and rows[1][3] == 'Ítem que no existe'


class TestReads:
    def test_history_is_paginated_and_carries_no_workbook_nor_names(self, client):
        post_upload(client, standard_workbook())
        body = client.get(API).json()
        assert body['count'] == 1
        row = body['results'][0]
        assert set(row) == {
            'id', 'status', 'warehouse', 'file_name', 'summary', 'result', 'uploaded_by_name', 'created_at',
            'confirmed_at', 'expires_at',
        }
        text = str(body)
        assert 'unmatched_destinations' not in text and not any(name in text for name in NAMES)
        assert 'file_bytes' not in text

    def test_detail_after_confirmation_has_no_destinations(self, client):
        preview = post_upload(client, standard_workbook()).json()
        confirm(client, preview)
        detail = client.get(f"{API}{preview['id']}/").json()
        assert detail['status'] == 'CONFIRMED' and detail['can_confirm'] is False
        assert detail['unmatched_destinations'] == [] and not any(name in str(detail) for name in NAMES)

    def test_template_has_the_four_sheets_with_headers(self, client):
        response = client.get(f'{API}template/')
        assert response.status_code == 200 and response['Content-Type'] == XLSX
        wb = load_workbook(io.BytesIO(response.content))
        assert wb.sheetnames == list(parser.SHEET_ORDER)
        assert wb['EGRESOS']['I1'].value == 'USUARIO/ DESTINO' and wb['MATRIZ DE RECURSOS']['B1'].value == 'Item'

    def test_the_template_is_a_valid_importable_book_with_no_rows(self, client):
        wb = client.get(f'{API}template/').content
        body = post_upload(client, wb).json()
        assert body['can_confirm'] is True and body['summary']['catalog']['new'] == 0


ALL_ROLES = [r for r in RoleChoices]


@pytest.mark.parametrize('role,can_write,can_read', [
    (RoleChoices.CENTRAL_WAREHOUSE, True, True),
    (RoleChoices.LOGISTICS_CHIEF, True, True),
    (RoleChoices.SITE_WAREHOUSE, False, True),
    (RoleChoices.LOGISTICS_SUPERVISOR, False, True),
    (RoleChoices.LOGISTICS_COORDINATOR, False, True),
    (RoleChoices.REQUESTER, False, False),
    (RoleChoices.PROJECT_RESIDENT, False, False),
    (RoleChoices.GENERAL_MANAGER, False, False),
    (RoleChoices.HR_MANAGER, False, False),
    (RoleChoices.ADMIN_MANAGER, False, False),
])
class TestPermissionMatrix:
    def test_matrix(self, role, can_write, can_read, central_warehouse):
        owner = client_for(central_warehouse)
        existing = post_upload(owner, standard_workbook()).json()

        client = client_for(make_user(f'u_{role.lower()}', role))
        read_status = 200 if can_read else 403
        assert client.get(API).status_code == read_status
        assert client.get(f"{API}{existing['id']}/").status_code == read_status
        assert client.get(f"{API}{existing['id']}/problems/").status_code == read_status
        assert client.get(f"{API}{existing['id']}/problems/export/").status_code == read_status
        assert client.get(f'{API}template/').status_code == read_status

        write = post_upload(client, build_workbook(catalog=[catalog_row('EPP9', 'Y')]))
        assert write.status_code == (201 if can_write else 403)
        confirm_status = confirm(client, existing).status_code
        assert (confirm_status == 200) if can_write else (confirm_status == 403)
        delete_status = client.delete(f"{API}{existing['id']}/").status_code
        assert delete_status == (409 if can_write else 403)           # confirmed above -> conflict, not 403


class TestNoBypass:
    def test_superuser_without_a_warehouse_role_is_rejected(self, superuser):
        client = client_for(superuser)
        assert client.get(API).status_code == 403
        assert post_upload(client, standard_workbook()).status_code == 403

    def test_anonymous(self, api_client):
        assert api_client.get(API).status_code in (401, 403)
        assert post_upload(api_client, standard_workbook()).status_code in (401, 403)


class TestRoundTripWithTheSyspccExport:
    """The workbook SYSPCC exports can be imported again (empty or other database)."""

    def test_exported_book_reimports_to_the_same_stock(self, client, central_warehouse, make_item):
        glove = make_item('EPP000001', 'GUANTE DE CUERO T: L', stock=None, location='F3-2', floor=2)
        helmet = make_item('EPP000002', 'CASCO BLANCO', stock=None, floor=2)
        movement_svc.register_entry_batch(
            [{'inventory_id': glove.pk, 'quantity': 7}], 'CENTRAL', None, 'INITIAL_LOAD', '', '', '', central_warehouse,
            movement_date=date(2026, 9, 1),
        )
        movement_svc.register_entry_batch(
            [{'inventory_id': glove.pk, 'quantity': 10, 'tag': 'LOTE-1'}, {'inventory_id': helmet.pk, 'quantity': 6}],
            'CENTRAL', None, 'PURCHASE', '', '', '', central_warehouse,
            movement_date=date(2026, 10, 1), po_number='PO-5', guide_number='GR-5',
        )
        movement_svc.register_exit_batch(
            [{'inventory_id': glove.pk, 'quantity': 4}, {'inventory_id': helmet.pk, 'quantity': 1}],
            'CENTRAL', None, 'EMPLOYEE', WORKER_A, None, None, '', central_warehouse,
            movement_date=date(2026, 10, 2), voucher_number='501',
        )
        expected = {i.product_code: InventoryStock.objects.get(inventory=i).quantity for i in (glove, helmet)}

        export = client.get('/api/v1/warehouse/inventory/movement-summary/export/', {'include_movements': 'true'})
        assert export.status_code == 200
        InventoryMovement.objects.all().delete()
        MovementGroup.objects.all().delete()
        InventoryStock.objects.all().delete()
        Inventory.objects.all().delete()

        preview = post_upload(client, export.content).json()
        assert preview['can_confirm'] is True and preview['summary']['catalog_source'] == parser.SHEET_SUMMARY
        assert preview['summary']['catalog']['new'] == 2
        assert preview['summary']['stock']['mismatched_items'] == 0
        done = confirm(client, preview, project_map={})
        assert done.status_code == 200 and done.json()['result']['reconciled'] is True

        restored = {i.product_code: InventoryStock.objects.get(inventory=i).quantity for i in Inventory.objects.all()}
        assert restored == expected
        item = Inventory.objects.get(product_code='EPP000001')
        assert (item.location, item.floor, item.size) == ('F3-2', 2, 'L')
        out = MovementGroup.objects.get(movement_type='EXIT')
        assert (out.voucher_number, out.destination_detail, out.movement_date) == ('501', WORKER_A, date(2026, 10, 2))
        entry = InventoryMovement.objects.get(movement_type='ENTRY', tag='LOTE-1')
        assert (entry.po_number, entry.guide_number, entry.movement_date) == ('PO-5', 'GR-5', date(2026, 10, 1))
        assert InventoryMovement.objects.get(source_type='INITIAL_LOAD').quantity == 7

    def test_a_workbook_with_a_truncation_warning_row_is_still_importable(self, client):
        data = build_workbook(
            catalog=[catalog_row('EPP1', 'X')], entries=[entry_row('EPP1', 2, location='F1-1')], header_offset=1,
        )
        body = post_upload(client, data).json()
        assert body['can_confirm'] is True and body['summary']['entries']['to_import'] == 1
