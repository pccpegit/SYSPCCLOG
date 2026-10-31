"""
SYSPCC-025 · QA ronda 1 — el importador frente a libros problemáticos.

Libros sintéticos (nombres y códigos inventados): números fuera de rango, celdas
infinitas, hojas vacías, duplicados, ítems desconocidos, fechas absurdas, destinos raros y
reimportaciones. Lo esperado siempre es una respuesta 4xx/2xx con el problema reportado,
nunca un 500, y que lo rechazado no deje rastro en la base de datos.

Los tests marcados "DEFECTO QA-n" describen un comportamiento incorrecto de producción que
esta ronda encontró: fallan a propósito hasta que se corrija.
"""
import pytest

from apps.warehouse.models import (
    Inventory, InventoryImport, InventoryImportStatus, InventoryMovement, InventoryStock,
)

from .conftest import client_for
from .import_workbooks import (
    D1, D2, D3, WORKER_A, build_workbook, catalog_row, entry_row, exit_row, standard_workbook, summary_row,
)
from .qa_support import (  # noqa: F401 (import_throttle_isolated is a fixture)
    IMPORTS, INVENTORY, confirm, import_throttle_isolated, put_number_in_sheets, snapshot_database, stock_of,
    tolerant_client, upload,
)

pytestmark = pytest.mark.django_db

ITEM = 'EPP000001'


@pytest.fixture(autouse=True)
def _isolated_throttle(import_throttle_isolated):
    yield


@pytest.fixture
def client(central_warehouse):
    return tolerant_client(central_warehouse)


def book(*, entries=None, exits=None, summary=None, catalog=None):
    return build_workbook(
        catalog=catalog if catalog is not None else [catalog_row(ITEM, 'GUANTE DE CUERO T: L')],
        entries=entries, exits=exits, summary=summary,
    )


def codes_of(response) -> dict:
    return response.json()['summary']['by_code']


# ---------------------------------------------------------------------------
# Números fuera de rango y celdas no finitas
# ---------------------------------------------------------------------------

HUGE = [1e30, 1e300, 10 ** 40, '9' * 30, '1' + '0' * 28]


class TestNumbersOutOfRange:
    """
    DEFECTO QA-1: `QuantityField` hace `value.quantize(...)` FUERA de su try/except; un número de
    28 o más dígitos (un `1E+30` tecleado por error en Excel) lanza `decimal.InvalidOperation`,
    que nadie captura: la subida responde 500 en lugar de reportar CANTIDAD_INVALIDA.
    """

    @pytest.mark.parametrize('value', HUGE, ids=lambda v: repr(v)[:14])
    def test_entry_quantity_is_a_skipped_row_not_a_server_error(self, client, value):
        response = upload(client, book(entries=[entry_row(ITEM, value)]))
        assert response.status_code == 201, response.status_code
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    @pytest.mark.parametrize('value', HUGE, ids=lambda v: repr(v)[:14])
    def test_exit_quantity_is_a_skipped_row_not_a_server_error(self, client, value):
        response = upload(client, book(exits=[exit_row(ITEM, value, WORKER_A, voucher=1)]))
        assert response.status_code == 201, response.status_code
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    @pytest.mark.parametrize('column', ['initial', 'stock'])
    def test_summary_figures_out_of_range_are_reported(self, client, column):
        row = summary_row(ITEM, 'GUANTE DE CUERO T: L', **{column: 1e30})
        response = upload(client, book(summary=[row]))
        assert response.status_code == 201, response.status_code
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    def test_the_largest_quantity_the_column_holds_is_still_accepted(self, client):
        data = build_workbook(
            catalog=[catalog_row('LYPCC0001', 'LIQUIDO', unit='LT')],
            entries=[entry_row('LYPCC0001', '999999999.999')],
        )
        response = upload(client, data)
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' not in codes_of(response)

    def test_the_largest_whole_quantity_of_a_piece_item_is_still_accepted(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, '999999999')]))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' not in codes_of(response)

    def test_one_more_unit_than_the_column_holds_is_rejected(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, '1000000000')]))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' in codes_of(response)


class TestNonFiniteCells:
    """
    DEFECTO QA-2: un campo de texto numérico (OC, PO, guía) con `1E999` (infinito para openpyxl)
    rompe `number_as_text` con OverflowError -> 500. Un libro fabricado a mano basta para provocarlo.
    """

    MARK = 424242

    def test_infinite_po_cell_is_reported_not_a_crash(self, client):
        data = book(entries=[entry_row(ITEM, 5, po=self.MARK)])
        response = upload(client, put_number_in_sheets(data, self.MARK, '1E999'))
        assert response.status_code in (201, 400), response.status_code

    def test_infinite_oc_cell_in_the_catalog_is_reported_not_a_crash(self, client):
        data = build_workbook(catalog=[catalog_row(ITEM, 'GUANTE', oc=self.MARK)])
        response = upload(client, put_number_in_sheets(data, self.MARK, '1E999'))
        assert response.status_code in (201, 400), response.status_code

    def test_infinite_quantity_cell_is_an_invalid_quantity(self, client):
        data = book(entries=[entry_row(ITEM, self.MARK)])
        response = upload(client, put_number_in_sheets(data, self.MARK, '1E999'))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' in codes_of(response)


class TestQuantityTextVariants:
    @pytest.mark.parametrize('text', ['1,5', '-3', '0', '1e3', '3 unidades', '١٢٣x', 'NaN', 'inf', '1.2.3'])
    def test_malformed_text_quantities_are_invalid(self, client, text):
        response = upload(client, book(entries=[entry_row(ITEM, text)]))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    @pytest.mark.parametrize('text', ['10', ' 10 ', '10.0'])
    def test_plain_text_quantities_are_accepted(self, client, text):
        response = upload(client, book(entries=[entry_row(ITEM, text)]))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' not in codes_of(response)
        assert 'CANTIDAD_VACIA' not in codes_of(response)

    def test_a_plain_text_decimal_is_accepted_for_a_measured_unit(self, client):
        data = build_workbook(
            catalog=[catalog_row('LYPCC0001', 'LIQUIDO', unit='LT')], entries=[entry_row('LYPCC0001', '0.5')],
        )
        response = upload(client, data)
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' not in codes_of(response)

    def test_a_blank_text_quantity_is_empty_not_invalid(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, '   ')]))
        assert response.status_code == 201
        assert 'CANTIDAD_VACIA' in codes_of(response)

    def test_more_than_three_decimals_is_invalid(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, 1.0004)]))
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    def test_boolean_cells_are_not_numbers(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, True)]))
        assert 'CANTIDAD_INVALIDA' in codes_of(response)


class TestFractionalPiecesInTheImporter:
    """
    DEFECTO QA-3: la API manual rechaza cantidades fraccionarias para EA/UND/PAR/JGO ("La cantidad
    de EA debe ser un número entero."), pero el importador no aplica esa regla: un 2.5 en un ítem EA
    entra al stock como 2.5 pares de guantes, que luego ninguna pantalla puede entregar.
    """

    def test_a_fractional_quantity_of_a_piece_unit_is_rejected_like_in_the_manual_entry(self, client):
        response = upload(client, book(entries=[entry_row(ITEM, 2.5)]))
        assert response.status_code == 201
        assert 'CANTIDAD_INVALIDA' in codes_of(response)

    def test_fractional_quantities_are_fine_for_measured_units(self, client):
        data = build_workbook(
            catalog=[catalog_row('LYPCC0001', 'LIQUIDO', unit='LT')],
            entries=[entry_row('LYPCC0001', 2.5)],
        )
        response = upload(client, data)
        assert 'CANTIDAD_INVALIDA' not in codes_of(response)


# ---------------------------------------------------------------------------
# Libros vacíos o casi vacíos
# ---------------------------------------------------------------------------

class TestEmptyAndDegenerateBooks:
    def empty_book(self):
        return build_workbook(catalog=[], entries=[], exits=[], summary=[])

    def test_a_book_with_only_headers_can_be_confirmed_and_writes_nothing(self, client):
        preview = upload(client, self.empty_book()).json()
        assert preview['can_confirm'] is True
        response = confirm(client, preview)
        assert response.status_code == 200, response.content
        assert response.json()['status'] == InventoryImportStatus.CONFIRMED
        assert (Inventory.objects.count(), InventoryMovement.objects.count(), InventoryStock.objects.count()) == (0, 0, 0)

    def test_that_empty_book_cannot_be_uploaded_again_after_confirming(self, client):
        preview = upload(client, self.empty_book()).json()
        confirm(client, preview)
        again = upload(client, self.empty_book())
        assert again.status_code == 409
        assert again.json()['code'] == 'file_already_imported'

    def test_a_catalog_without_movements_creates_items_with_zero_stock(self, client, central_warehouse):
        data = build_workbook(catalog=[catalog_row('EPP000001', 'A'), catalog_row('EPP000002', 'B')])
        preview = upload(client, data).json()
        assert confirm(client, preview).status_code == 200
        assert Inventory.objects.count() == 2
        assert InventoryMovement.objects.count() == 0
        listed = client_for(central_warehouse).get(INVENTORY, {'no_page': 'true'}).json()
        assert [row['total_stock'] for row in listed] == [0, 0]

    def test_zero_stock_items_are_warned_about(self, client):
        data = build_workbook(catalog=[catalog_row('EPP000001', 'A')])
        assert 'PRODUCTO_SIN_STOCK_CERO' in codes_of(upload(client, data))

    def test_rows_with_only_the_row_number_are_blank_and_row_numbers_stay_true(self, client):
        entries = [entry_row(ITEM, 5), [None] * 12, [None] * 12, entry_row(ITEM, 'abc')]
        preview = upload(client, book(entries=entries)).json()
        listing = client.get(f"{IMPORTS}{preview['id']}/problems/", {'code': 'CANTIDAD_INVALIDA'}).json()
        # header = row 1, so the bad quantity is the 4th data row = Excel row 5
        assert [p['row'] for p in listing['results']] == [5]
        assert preview['summary']['sheets']['INGRESOS']['rows'] == 2

    def test_a_sheet_with_a_header_and_a_thousand_blank_rows_is_fast_and_empty(self, client):
        entries = [[None] * 12 for _ in range(1000)]
        response = upload(client, book(entries=entries))
        assert response.status_code == 201
        assert response.json()['summary']['sheets']['INGRESOS']['rows'] == 0


# ---------------------------------------------------------------------------
# Duplicados, ítems desconocidos, conflictos
# ---------------------------------------------------------------------------

class TestDuplicatesAndUnknownItems:
    def test_codes_that_differ_only_in_case_and_spaces_are_one_item(self, client):
        data = build_workbook(
            catalog=[catalog_row('EPP000001', 'CASCO'), catalog_row(' epp000001 ', 'casco')],
            entries=[entry_row(' epp000001', 4)],
        )
        preview = upload(client, data).json()
        assert preview['summary']['by_code']['CODIGO_DUPLICADO_IDENTICO'] == 1
        assert confirm(client, preview).status_code == 200
        assert Inventory.objects.count() == 1
        assert stock_of('EPP000001') == 4

    def test_the_same_code_with_another_unit_blocks_the_import_and_leaves_no_trace(self, client):
        data = build_workbook(
            catalog=[catalog_row('EPP000001', 'CASCO', unit='EA'), catalog_row('EPP000001', 'CASCO', unit='PAR')],
            entries=[entry_row('EPP000001', 4)],
        )
        preview = upload(client, data).json()
        assert preview['can_confirm'] is False
        before = snapshot_database()
        response = confirm(client, preview)
        assert response.status_code == 400
        assert response.json()['code'] == 'blockers'
        assert snapshot_database() == before

    def test_movements_of_an_unknown_item_block_the_import_even_with_every_acceptance(self, client):
        data = book(entries=[entry_row('ZZZ999', 3)], exits=[exit_row('ZZZ998', 1, WORKER_A, voucher=9)])
        preview = upload(client, data).json()
        assert preview['summary']['by_code']['ITEM_DESCONOCIDO'] == 2
        before = snapshot_database()
        response = confirm(client, preview, accept_warnings=True, accept_skipped_rows=True)
        assert response.status_code == 400 and response.json()['code'] == 'blockers'
        assert snapshot_database() == before
        assert InventoryImport.objects.get().status == InventoryImportStatus.PREVIEWED

    def test_an_item_known_only_to_the_system_is_accepted(self, client, make_item):
        make_item('EPP777', 'EXISTENTE', stock=3)
        data = build_workbook(catalog=[catalog_row(ITEM, 'X')], entries=[entry_row('EPP777', 2)])
        preview = upload(client, data).json()
        assert 'ITEM_DESCONOCIDO' not in preview['summary']['by_code']
        assert confirm(client, preview).status_code == 200
        assert stock_of('EPP777') == 5

    def test_exiting_more_than_exists_blocks_and_never_leaves_negative_stock(self, client):
        data = book(entries=[entry_row(ITEM, 3)], exits=[exit_row(ITEM, 4, WORKER_A, voucher=1)])
        preview = upload(client, data).json()
        assert preview['summary']['by_code']['STOCK_NEGATIVO'] == 1
        before = snapshot_database()
        response = confirm(client, preview)
        assert response.status_code == 400 and response.json()['code'] == 'blockers'
        assert snapshot_database() == before
        assert not InventoryStock.objects.filter(quantity__lt=0).exists()

    def test_exiting_exactly_what_exists_leaves_zero_not_negative(self, client):
        data = book(entries=[entry_row(ITEM, 3)], exits=[exit_row(ITEM, 3, WORKER_A, voucher=1)])
        preview = upload(client, data).json()
        assert confirm(client, preview).status_code == 200
        assert stock_of(ITEM) == 0

    def test_an_exit_dated_before_its_entry_still_nets_to_the_right_stock_and_shows_in_the_kardex(
        self, client, central_warehouse,
    ):
        data = book(
            entries=[entry_row(ITEM, 10, D3)],
            exits=[exit_row(ITEM, 4, WORKER_A, voucher=1, when=D1)],        # exit BEFORE the entry
        )
        preview = upload(client, data).json()
        assert confirm(client, preview).status_code == 200
        item = Inventory.objects.get(product_code=ITEM)
        assert stock_of(ITEM) == 6
        kardex = client_for(central_warehouse).get(f'{INVENTORY}{item.pk}/kardex/').json()
        balances = [row['balance_after'] for row in kardex['movements']]
        assert balances == [-4, 6]                              # the chronological dip is visible to the user
        assert kardex['closing_stock'] == 6 == balances[-1]

    def test_a_row_without_the_item_code_is_skipped_but_does_not_block(self, client):
        data = book(entries=[entry_row(None, 5), entry_row(ITEM, 2)])
        preview = upload(client, data).json()
        assert preview['summary']['by_code']['ITEM_VACIO'] == 1
        assert preview['can_confirm'] is True

    def test_an_over_long_description_skips_the_item_and_blocks_its_movements_without_crashing(self, client):
        data = build_workbook(
            catalog=[catalog_row(ITEM, 'X' * 301)], entries=[entry_row(ITEM, 2)],
        )
        response = upload(client, data)
        assert response.status_code == 201
        assert response.json()['can_confirm'] is False


# ---------------------------------------------------------------------------
# Fechas
# ---------------------------------------------------------------------------

class TestDates:
    @pytest.mark.parametrize('value', ['16/10/2026', '2026-10-16', '16-10-2026', 46311, 46311.75])
    def test_valid_date_formats_are_read(self, client, value):
        response = upload(client, book(entries=[entry_row(ITEM, 1, value)]))
        assert response.status_code == 201
        assert not {'FECHA_INVALIDA', 'FECHA_FUTURA', 'FECHA_VACIA'} & set(codes_of(response))

    @pytest.mark.parametrize('value', ['31/02/2026', '2026-13-01', 'mañana', '0000-00-00', -1, 0, 1e30, '99/99/9999'])
    def test_impossible_dates_are_invalid_rows(self, client, value):
        response = upload(client, book(entries=[entry_row(ITEM, 1, value)]))
        assert response.status_code == 201
        assert 'FECHA_INVALIDA' in codes_of(response)

    def test_a_date_in_the_future_is_rejected(self, client):
        from datetime import date, timedelta
        from django.utils import timezone
        future = timezone.localdate() + timedelta(days=2)
        assert isinstance(future, date)
        response = upload(client, book(entries=[entry_row(ITEM, 1, future)]))
        assert 'FECHA_FUTURA' in codes_of(response)

    def test_a_date_before_2000_is_invalid(self, client):
        from datetime import date
        response = upload(client, book(entries=[entry_row(ITEM, 1, date(1999, 12, 31))]))
        assert 'FECHA_INVALIDA' in codes_of(response)

    def test_an_undated_row_is_imported_without_date_and_warned(self, client):
        preview = upload(client, book(entries=[entry_row(ITEM, 7, None)])).json()
        assert preview['summary']['by_code']['FECHA_VACIA'] == 1
        assert confirm(client, preview).status_code == 200
        assert InventoryMovement.objects.get().movement_date is None


# ---------------------------------------------------------------------------
# Destinos extraños
# ---------------------------------------------------------------------------

class TestOddDestinations:
    @pytest.mark.parametrize('destination', [
        'Ñandú Pérez 🙂', 'محمد علي', 'Zero​Width Test', 'Tab\tName', "O'Brien; DROP TABLE x;--",
    ])
    def test_unusual_text_destinations_import_without_errors(self, client, destination):
        data = book(entries=[entry_row(ITEM, 5)], exits=[exit_row(ITEM, 1, destination, voucher=1)])
        preview = upload(client, data)
        assert preview.status_code == 201, preview.status_code
        assert confirm(client, preview.json()).status_code == 200
        assert InventoryMovement.objects.get(movement_type='EXIT').destination_detail

    def test_an_over_long_destination_is_a_skipped_row(self, client):
        data = book(entries=[entry_row(ITEM, 5)], exits=[exit_row(ITEM, 1, 'N' * 201, voucher=1)])
        response = upload(client, data)
        assert response.status_code == 201
        assert 'FILA_INVALIDA' in codes_of(response)

    @pytest.mark.parametrize('destination', ['@SUM(1,1)', '+1+1', '-2+3'])
    def test_a_formula_looking_destination_is_stored_as_text_and_neutralised_in_the_export(
        self, client, central_warehouse, destination,
    ):
        # '=...' cannot be a text cell with openpyxl (it becomes a formula, which is never evaluated:
        # that case is covered by the parser tests); '@', '+' and '-' are the other formula triggers.
        import io
        import openpyxl
        data = book(entries=[entry_row(ITEM, 5)], exits=[exit_row(ITEM, 1, destination, voucher=1)])
        preview = upload(client, data).json()
        assert confirm(client, preview).status_code == 200
        response = client_for(central_warehouse).get('/api/v1/warehouse/movements/export/', {'layout': 'excel_pcc'})
        assert response.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(response.content))['EGRESOS']
        destinations = [c.value for c in ws['I'][1:] if c.value]
        assert destinations, 'the exit must have been imported and exported'
        assert not any(str(v)[0] in '=+-@' for v in destinations)


# ---------------------------------------------------------------------------
# Reimportaciones
# ---------------------------------------------------------------------------

class TestReimport:
    def test_importing_the_same_rows_in_a_second_book_changes_nothing(self, client):
        first = upload(client, standard_workbook()).json()
        assert confirm(client, first).status_code == 200
        snapshot = snapshot_database()
        # same business rows, another file (an extra row of an unrelated item changes the hash)
        data = standard_workbook(catalog=[
            catalog_row('EPP000001', 'GUANTE DE CUERO T: L'), catalog_row('EPP000002', 'CASCO BLANCO'),
            catalog_row('EPP000003', 'LENTE CLARO'), catalog_row('EPP000004', 'OTRO'),
        ])
        second = upload(client, data).json()
        assert second['summary']['entries']['already_imported'] == 2
        assert second['summary']['exits']['already_imported'] == 3
        assert confirm(client, second).status_code == 200
        after = snapshot_database()
        assert after['movements'] == snapshot['movements']
        assert after['items'] == snapshot['items'] + 1
        assert after['stock'] == snapshot['stock']

    def test_manual_movements_made_after_an_import_survive_a_longer_book(self, client, central_warehouse):
        base_entries = [entry_row(ITEM, 10, D1)]
        first = upload(client, book(entries=base_entries)).json()
        assert confirm(client, first).status_code == 200
        item = Inventory.objects.get(product_code=ITEM)
        manual = client_for(central_warehouse).post('/api/v1/warehouse/movements/exit/', {
            'inventory': item.pk, 'quantity': '4', 'warehouse': 'CENTRAL',
            'destination_type': 'EMPLOYEE', 'destination_detail': 'Persona Manual Test',
        }, format='json')
        assert manual.status_code == 201, manual.content
        longer = upload(client, book(entries=base_entries + [entry_row(ITEM, 5, D2)])).json()
        assert confirm(client, longer).status_code == 200
        assert stock_of(ITEM) == 10 - 4 + 5

    def test_a_corrected_quantity_is_imported_as_new_and_the_old_row_is_reported_missing(self, client):
        first = upload(client, book(entries=[entry_row(ITEM, 10, D1)])).json()
        confirm(client, first)
        second = upload(client, book(entries=[entry_row(ITEM, 12, D1)])).json()
        assert second['summary']['by_code'].get('FILAS_IMPORTADAS_AUSENTES') == 1
        confirm(client, second)
        assert InventoryMovement.objects.filter(movement_type='ENTRY').count() == 2      # nothing is ever deleted

    def test_cancelling_a_preview_then_uploading_it_again_works(self, client):
        data = standard_workbook()
        first = upload(client, data).json()
        assert client.delete(f"{IMPORTS}{first['id']}/").status_code == 200
        second = upload(client, data)
        assert second.status_code == 201
        assert confirm(client, second.json()).status_code == 200

    def test_two_previews_of_one_book_only_confirm_once(self, client):
        data = standard_workbook()
        first, second = upload(client, data).json(), upload(client, data).json()
        assert confirm(client, first).status_code == 200
        loser = confirm(client, second)
        assert loser.status_code == 409, loser.content
        assert InventoryMovement.objects.count() == 5


# ---------------------------------------------------------------------------
# Bordes de la API del importador
# ---------------------------------------------------------------------------

NO_SUCH = '11111111-2222-3333-4444-555555555555'


class TestImportApiEdges:
    @pytest.mark.parametrize('suffix', ['', 'problems/', 'problems/export/'])
    def test_unknown_or_malformed_ids_are_a_404_for_reads(self, client, suffix):
        assert client.get(f'{IMPORTS}{NO_SUCH}/{suffix}').status_code == 404
        assert client.get(f'{IMPORTS}no-es-un-uuid/{suffix}').status_code == 404

    def test_unknown_or_malformed_ids_are_a_404_for_writes(self, client):
        body = {'expected_plan_digest': 'x'}
        assert client.post(f'{IMPORTS}{NO_SUCH}/confirm/', body, format='json').status_code == 404
        assert client.post(f'{IMPORTS}no-es-un-uuid/confirm/', body, format='json').status_code == 404
        assert client.delete(f'{IMPORTS}{NO_SUCH}/').status_code == 404

    @pytest.mark.parametrize('body', [
        {}, {'expected_plan_digest': ''}, {'expected_plan_digest': 'x' * 65},
        {'expected_plan_digest': 'x', 'project_map': {'ON PLANT': 'abc'}},
        {'expected_plan_digest': 'x', 'project_map': ['ON PLANT']},
        {'expected_plan_digest': 'x', 'destination_map': {'k': 0}},
        {'expected_plan_digest': 'x', 'destination_map': {'k': -3}},
        {'expected_plan_digest': 'x', 'accept_warnings': 'tal vez'},
    ])
    def test_malformed_confirmation_bodies_are_a_400_and_write_nothing(self, client, body):
        preview = upload(client, standard_workbook()).json()
        before = snapshot_database()
        response = client.post(f"{IMPORTS}{preview['id']}/confirm/", body, format='json')
        assert response.status_code == 400, (body, response.status_code)
        assert snapshot_database() == before

    def test_a_wrong_digest_is_refused_without_writing(self, client):
        preview = upload(client, standard_workbook()).json()
        before = snapshot_database()
        response = confirm(client, preview, expected_plan_digest='0' * 64)
        assert response.status_code == 409 and response.json()['code'] == 'plan_changed'
        assert snapshot_database() == before
        assert InventoryImport.objects.get().status == InventoryImportStatus.PREVIEWED

    @pytest.mark.parametrize('extra', [{'warehouse': 'SITE'}, {'warehouse': 'OFFICE'}, {'warehouse': 'MARTE'}])
    def test_only_the_central_warehouse_can_be_imported(self, client, extra):
        response = upload(client, standard_workbook(), **extra)
        assert response.status_code == 400
        assert InventoryImport.objects.count() == 0

    def test_a_body_without_a_file_or_with_the_file_as_text_is_a_400(self, client):
        assert client.post(IMPORTS, {}, format='multipart').status_code == 400
        assert client.post(IMPORTS, {'file': 'no soy un archivo'}, format='multipart').status_code == 400
        assert client.post(IMPORTS, {'file': 'x'}, format='json').status_code in (400, 415)

    def test_problem_pages_never_exceed_one_hundred_rows_and_out_of_range_pages_are_a_404(self, client):
        entries = [entry_row(ITEM, 'abc') for _ in range(150)]
        preview = upload(client, book(entries=entries)).json()
        listing = client.get(f"{IMPORTS}{preview['id']}/problems/", {'page_size': 100000}).json()
        assert len(listing['results']) == 100 and listing['count'] >= 150
        assert client.get(f"{IMPORTS}{preview['id']}/problems/", {'page': 999}).status_code == 404

    def test_unknown_filters_are_ignored_and_the_export_still_opens(self, client):
        import io
        import openpyxl
        preview = upload(client, book(entries=[entry_row(ITEM, 'abc')])).json()
        response = client.get(f"{IMPORTS}{preview['id']}/problems/export/", {'severity': 'zzz', 'sheet': 'nope'})
        assert response.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(response.content))['PROBLEMAS']
        assert ws.max_row == 1                                  # 'nope' is not a sheet: nothing matches

    def test_every_import_response_is_never_cached(self, client):
        preview = upload(client, standard_workbook()).json()
        urls = [IMPORTS, f"{IMPORTS}{preview['id']}/", f"{IMPORTS}{preview['id']}/problems/", f'{IMPORTS}template/']
        for url in urls:
            assert client.get(url).headers['Cache-Control'] == 'no-store', url
        assert confirm(client, preview).headers['Cache-Control'] == 'no-store'


# ---------------------------------------------------------------------------
# El comando comparte el pipeline: mismos libros, mismas respuestas
# ---------------------------------------------------------------------------

class TestCommandWithHostileBooks:
    def run(self, tmp_path, data, *args):
        from io import StringIO

        from django.core.management import call_command

        path = tmp_path / 'libro.xlsx'
        path.write_bytes(data)
        out = StringIO()
        call_command('import_warehouse_excel', str(path), *args, stdout=out)
        return out.getvalue()

    def test_dry_run_of_an_out_of_range_quantity_is_reported_not_a_traceback(self, tmp_path):
        """DEFECTO QA-1 (mismo origen): el comando también revienta con `decimal.InvalidOperation`."""
        output = self.run(tmp_path, book(entries=[entry_row(ITEM, 1e30)]), '--dry-run')
        assert 'CANTIDAD_INVALIDA' in output

    def test_dry_run_of_an_empty_book_says_there_is_nothing_to_import(self, tmp_path):
        output = self.run(tmp_path, build_workbook(catalog=[], entries=[], exits=[], summary=[]), '--dry-run')
        assert 'Se puede confirmar: sí' in output
        assert Inventory.objects.count() == 0
